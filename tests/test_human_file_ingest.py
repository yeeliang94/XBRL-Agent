"""Reading a human-filled mTool file against one run (eval/human_file.py).

Figures are proven with a real fill -> read round trip on a repository
template: whatever the mTool fill writes, the reader must read back exactly,
including concepts that share a label on one sheet (the Step 1 failure of the
old label reader).
"""
from __future__ import annotations

import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from db.schema import init_db
from mtool.offline_fill import wrap_footnote_html
from eval.human_file import (
    HumanFileError,
    ingest_human_file,
    load_human_file,
    read_human_file,
)
from tests._human_file_fixture import (
    SOFP,
    add_run,
    filled_file,
    import_template_file,
    leaf_facts,
    make_sofp_db,
    sofp_config,
)


@pytest.fixture
def sofp_db(tmp_path):
    return make_sofp_db(tmp_path)


def _typed(read) -> dict[tuple[str, str], float]:
    return {(f["concept_uuid"], f["period"]): f["value"]
            for f in read.facts if f["value"] is not None}


def test_round_trip_reads_every_filled_value_including_same_label_fields(
    sofp_db, tmp_path,
):
    run_id = add_run(sofp_db, sofp_config())
    facts = leaf_facts(sofp_db)
    path = filled_file(sofp_db, run_id, facts, tmp_path)

    with sqlite3.connect(sofp_db) as conn:
        read = read_human_file(conn, run_id, path, "units")

    assert _typed(read) == {(u, p): v for u, p, v in facts}
    assert read.unmatched == []
    assert read.not_compared == []
    assert read.summary["typed_values"] == len(facts)
    assert read.summary["magnitude_warning"] is None


def test_unaddressed_typed_row_is_unmatched_and_formula_cell_is_calculated(
    sofp_db, tmp_path,
):
    run_id = add_run(sofp_db, sofp_config())
    facts = leaf_facts(sofp_db)
    path = filled_file(sofp_db, run_id, facts, tmp_path)
    with sqlite3.connect(sofp_db) as conn:
        header_row, header_label = conn.execute(
            "SELECT render_row, canonical_label FROM concept_nodes "
            "WHERE template_id = ? AND kind = 'ABSTRACT' "
            "AND render_sheet = 'SOFP-CuNonCu' ORDER BY render_row LIMIT 1",
            (SOFP,),
        ).fetchone()
        formula_uuid, formula_row = conn.execute(
            "SELECT n.concept_uuid, t.target_row FROM concept_nodes n "
            "JOIN concept_targets t USING(concept_uuid) "
            "WHERE n.concept_uuid = ? AND t.period = 'CY'", (facts[-1][0],),
        ).fetchone()
    wb = load_workbook(path)
    wb["SOFP-CuNonCu"][f"B{header_row}"] = 77
    wb["SOFP-CuNonCu"][f"B{formula_row}"] = "=1+1"
    wb.save(path)

    with sqlite3.connect(sofp_db) as conn:
        read = read_human_file(conn, run_id, path, "units")

    assert [(u["sheet"], u["row"], u["label"], u["values"]) for u in read.unmatched] == [
        ("SOFP-CuNonCu", header_row, header_label, {"B": 77.0})]
    calculated = [f for f in read.facts if f["calculated"]]
    assert [(f["concept_uuid"], f["period"], f["value"]) for f in calculated] == [
        (formula_uuid, "CY", None)]
    assert (formula_uuid, "CY") not in _typed(read)


def test_statement_filled_in_another_variant_is_not_compared(sofp_db, tmp_path):
    import_template_file(sofp_db, tmp_path, "02-SOFP-OrderOfLiquidity.xlsx")
    source_run = add_run(sofp_db, sofp_config())
    path = filled_file(sofp_db, source_run, leaf_facts(sofp_db), tmp_path)
    run_id = add_run(sofp_db, sofp_config(variant="OrderOfLiquidity"))

    with sqlite3.connect(sofp_db) as conn, pytest.raises(
        HumanFileError, match=r"different statement layout: SOFP \(CuNonCu\)"
    ):
        read_human_file(conn, run_id, path, "units")


def test_unit_converts_to_run_denomination_and_warns_on_1000x(sofp_db, tmp_path):
    run_id = add_run(sofp_db, sofp_config(denomination="thousands"))
    facts = leaf_facts(sofp_db)
    path = filled_file(sofp_db, run_id, facts, tmp_path)

    with sqlite3.connect(sofp_db) as conn:
        read = read_human_file(conn, run_id, path, "units")

    assert _typed(read) == {(u, p): v / 1000 for u, p, v in facts}
    assert "1,000 times smaller" in read.summary["magnitude_warning"]


def test_unit_conversion_leaves_share_counts_unchanged(sofp_db, tmp_path, monkeypatch):
    from mtool.units import MONETARY, SHARES

    run_id = add_run(sofp_db, sofp_config(denomination="thousands"))
    facts = leaf_facts(sofp_db)
    path = filled_file(sofp_db, run_id, facts, tmp_path)
    with sqlite3.connect(sofp_db) as conn:
        share_label = conn.execute(
            "SELECT canonical_label FROM concept_nodes WHERE concept_uuid = ?",
            (facts[0][0],),
        ).fetchone()[0]
    monkeypatch.setattr(
        "eval.human_file.unit_class_for_label",
        lambda label, standard: SHARES if label == share_label else MONETARY,
    )

    with sqlite3.connect(sofp_db) as conn:
        read = read_human_file(conn, run_id, path, "units")

    typed = _typed(read)
    assert typed[(facts[0][0], "CY")] == facts[0][2]
    assert typed[(facts[-1][0], "PY")] == facts[-1][2] / 1000


def test_unmatched_number_alone_does_not_compare_statement(sofp_db, tmp_path):
    run_id = add_run(sofp_db, sofp_config())
    facts = leaf_facts(sofp_db)
    path = filled_file(sofp_db, run_id, facts, tmp_path)
    wb = load_workbook(path)
    with sqlite3.connect(sofp_db) as conn:
        targets = conn.execute(
            "SELECT target_sheet, target_row, target_col FROM concept_targets "
            f"WHERE concept_uuid IN ({','.join('?' * len(facts))})",
            [uuid for uuid, _, _ in facts],
        ).fetchall()
        header_row = conn.execute(
            "SELECT render_row FROM concept_nodes WHERE template_id = ? "
            "AND kind = 'ABSTRACT' AND render_sheet = 'SOFP-CuNonCu' "
            "ORDER BY render_row LIMIT 1", (SOFP,),
        ).fetchone()[0]
    for sheet, row, col in targets:
        wb[sheet][f"{col}{row}"] = None
    wb["SOFP-CuNonCu"][f"B{header_row}"] = 77
    wb.save(path)

    with sqlite3.connect(sofp_db) as conn, pytest.raises(
        HumanFileError, match="No figures or notes in this file match"
    ):
        read_human_file(conn, run_id, path, "units")


def _notes_workbook(tmp_path: Path) -> Path:
    wb = Workbook()
    visible = wb.active
    visible.title = "Notes-CI"
    visible["A12"] = ("ssmt-mfrs-cor_2022-12-31.xsd#ssmt-mfrs_"
                      "DisclosureOfCorporateInformationExplanatory@http://x/role")
    visible["D12"] = "*Disclosure of corporate information"
    visible["E12"] = "[Text block added]"
    visible["A13"] = "ssmt-mfrs-cor_2022-12-31.xsd#ssmt-mfrs_UnknownExplanatory"
    visible["D13"] = "Unknown note"
    visible["E13"] = "[Text block added]"
    footnotes = wb.create_sheet("+FootnoteTexts")
    for row, key in ((1, "fn_1"), (2, "fn_2")):
        footnotes[f"A{row}"] = key
        footnotes[f"B{row}"] = "Notes-CI"
        # The mTool XHTML text-block shell, with Excel's _x000D_ line breaks.
        footnotes[f"C{row}"] = wrap_footnote_html(
            f'<p onclick="x()">Human note {row}</p><script>alert(1)</script>')
    wb.defined_names["fn_1"] = DefinedName("fn_1", attr_text="'Notes-CI'!$E$12")
    wb.defined_names["fn_2"] = DefinedName("fn_2", attr_text="'Notes-CI'!$E$13")
    path = tmp_path / "notes.xlsx"
    wb.save(path)
    return path


def test_note_maps_to_field_by_taxonomy_element_id_and_is_unwrapped(tmp_path):
    db = tmp_path / "audit.db"
    init_db(db)
    with sqlite3.connect(db) as conn:
        conn.execute(
            "INSERT INTO template_slots(target_id, canonical_target_id, "
            "template_id, sheet, row, col, label, slot_role, value_kind, "
            "taxonomy_element_id, mapping_source, manifest_version, "
            "workbook_fingerprint, validation_status) VALUES ('t1', 'ci-field', "
            "'mfrs-company-notes-corporateinfo-v1', 'Notes-CI', 5, 'B', "
            "'*Disclosure of corporate information', 'INPUT', 'html', "
            "'ssmt-mfrs_DisclosureOfCorporateInformationExplanatory', "
            "'presentation_linkbase', 'v1', 'fp', 'writable')")
    run_id = add_run(db, {"filing_standard": "mfrs", "filing_level": "company",
                       "statements": [], "notes_to_run": ["CORP_INFO"]})

    with sqlite3.connect(db) as conn:
        read = read_human_file(conn, run_id, _notes_workbook(tmp_path), "units")

    assert read.notes == [{"concept_uuid": "ci-field", "note_key": "fn_1",
                           "html": "<p>Human note 1</p>"}]
    assert [(u["kind"], u["row"], u["label"]) for u in read.unmatched] == [
        ("note", 13, "Unknown note")]


def test_empty_note_markup_is_not_a_filled_human_field(tmp_path):
    db = tmp_path / "audit.db"
    init_db(db)
    with sqlite3.connect(db) as conn:
        conn.execute(
            "INSERT INTO template_slots(target_id, canonical_target_id, "
            "template_id, sheet, row, col, label, slot_role, value_kind, "
            "taxonomy_element_id, mapping_source, manifest_version, "
            "workbook_fingerprint, validation_status) VALUES ('t1', 'ci-field', "
            "'mfrs-company-notes-corporateinfo-v1', 'Notes-CI', 5, 'B', "
            "'*Disclosure of corporate information', 'INPUT', 'html', "
            "'ssmt-mfrs_DisclosureOfCorporateInformationExplanatory', "
            "'presentation_linkbase', 'v1', 'fp', 'writable')")
    run_id = add_run(db, {"filing_standard": "mfrs", "filing_level": "company",
                           "statements": [], "notes_to_run": ["CORP_INFO"]})
    path = _notes_workbook(tmp_path)
    wb = load_workbook(path)
    wb["+FootnoteTexts"]["C1"] = wrap_footnote_html("<p></p>")
    wb.save(path)

    with sqlite3.connect(db) as conn, pytest.raises(
        HumanFileError, match="No figures or notes in this file match"
    ):
        read_human_file(conn, run_id, path, "units")


def test_replacing_a_file_keeps_one_record_and_refuses_unfinished_runs(
    sofp_db, tmp_path,
):
    run_id = add_run(sofp_db, sofp_config())
    facts = leaf_facts(sofp_db)
    path = filled_file(sofp_db, run_id, facts, tmp_path)

    with sqlite3.connect(sofp_db) as conn:
        for name in ("first.xlsx", "second.xlsx"):
            record = ingest_human_file(conn, run_id, path, filename=name,
                                       unit="units", uploaded_by="a@b.c")
        assert record["filename"] == "second.xlsx"
        assert conn.execute(
            "SELECT COUNT(*) FROM human_file_facts WHERE run_id = ? "
            "AND value IS NOT NULL", (run_id,)).fetchone()[0] == len(facts)
        assert load_human_file(conn, run_id)["summary"]["typed_values"] == len(facts)

        draft = add_run(sofp_db, sofp_config(), status="draft")
        with pytest.raises(HumanFileError, match="completed run"):
            read_human_file(conn, draft, path, "units")
