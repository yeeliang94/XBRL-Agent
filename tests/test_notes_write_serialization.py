"""Pin same-workbook notes writes and the source DB/file commit boundary."""
from __future__ import annotations

import asyncio
import sqlite3
import threading
import time
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
from pydantic_ai.models.test import TestModel

from concept_model.filing_targets import persist_template_manifest, writable_rows
from concept_model.notes_importer import import_notes_template
from concept_model.notes_parser import parse_notes_template
from db import repository as repo
from db.schema import init_db
from notes import agent as notes_agent
from notes import source_repository as srepo
from notes.source_models import SourceBlock, SourceNote
from notes_types import NOTES_REGISTRY, NotesTemplateType, notes_template_path


def _tool(agent, name):
    return next(
        tool.function
        for toolset in agent.toolsets
        for tool_name, tool in getattr(toolset, "tools", {}).items()
        if tool_name == name
    )


def test_parallel_tool_calls_serialize_one_workbook_read_modify_save(
    tmp_path, monkeypatch,
):
    kind = NotesTemplateType.CORP_INFO
    template = notes_template_path(kind, level="company")
    sheet = NOTES_REGISTRY[kind].sheet_name
    allowed = writable_rows(template, sheet)
    book = load_workbook(template, read_only=True)
    try:
        labels = [
            (row, str(book[sheet].cell(row, 1).value or "").strip())
            for row in sorted(allowed or [])
            if str(book[sheet].cell(row, 1).value or "").strip()
        ][:4]
    finally:
        book.close()
    assert len(labels) == 4

    agent, deps = notes_agent.create_notes_agent(
        template_type=kind,
        pdf_path=str(tmp_path / "source.pdf"),
        inventory=[],
        filing_level="company",
        model=TestModel(),
        output_dir=str(tmp_path),
    )
    real_writer = notes_agent.write_notes_workbook
    counter_lock = threading.Lock()
    active = 0
    peak = 0

    def observed_writer(*args, **kwargs):
        nonlocal active, peak
        with counter_lock:
            active += 1
            peak = max(peak, active)
        try:
            # Widen the overlap enough that the pre-fix implementation
            # deterministically entered several read-modify-save operations.
            time.sleep(0.03)
            return real_writer(*args, **kwargs)
        finally:
            with counter_lock:
                active -= 1

    monkeypatch.setattr(notes_agent, "write_notes_workbook", observed_writer)
    write_notes = _tool(agent, "write_notes")
    ctx = SimpleNamespace(deps=deps)

    async def run_all():
        return await asyncio.gather(*[
            write_notes(ctx, [{
                "chosen_row_label": label,
                "content": f"<p>parallel payload {index}</p>",
                "evidence": f"Page {index}",
                "source_pages": [index],
                "parent_note": {"number": str(index), "title": f"Title {index}"},
            }])
            for index, (_row, label) in enumerate(labels, start=1)
        ])

    messages = asyncio.run(run_all())
    assert peak == 1
    assert all(message.startswith("Wrote 1 row") for message in messages)
    assert len(deps.cells_written) == 4

    filled = load_workbook(deps.filled_path, read_only=True)
    try:
        for index, (row, _label) in enumerate(labels, start=1):
            assert f"parallel payload {index}" in filled[sheet].cell(row, 2).value
    finally:
        filled.close()


def test_source_projection_failure_rolls_back_cell_and_lineage(
    tmp_path, monkeypatch,
):
    kind = NotesTemplateType.CORP_INFO
    template = notes_template_path(kind, level="company")
    sheet = NOTES_REGISTRY[kind].sheet_name
    db = tmp_path / "audit.sqlite"
    init_db(db)
    template_id, nodes = parse_notes_template(str(template), sheet)
    import_notes_template(db, template_id, nodes)
    persist_template_manifest(db, template)
    with repo.db_session(db) as conn:
        run_id = repo.create_run(
            conn, "source.pdf", session_id="s", output_dir=str(tmp_path),
        )
        generation_id = srepo.begin_generation(
            conn, run_id, input_kind="prepared_document",
        )
        srepo.write_blocks(conn, generation_id, [
            SourceBlock(
                "b1", "paragraph", 1, "<p>Source disclosure.</p>",
                source_note_id="n1",
            ),
        ])
        srepo.write_notes(conn, generation_id, [
            SourceNote("n1", "1", "Disclosure", ["b1"]),
        ])
        srepo.activate_generation(conn, generation_id)

    target_row = min(node.row for node in nodes if node.kind == "LEAF")
    agent, deps = notes_agent.create_notes_agent(
        template_type=kind,
        pdf_path=str(tmp_path / "source.pdf"),
        inventory=[],
        filing_level="company",
        model=TestModel(),
        output_dir=str(tmp_path),
        run_id=run_id,
        db_path=str(db),
        source_generation_id=generation_id,
    )

    def fail_projection(*args, **kwargs):
        raise PermissionError("injected workbook replace failure")

    monkeypatch.setattr(notes_agent, "write_notes_workbook", fail_projection)
    result = asyncio.run(_tool(agent, "write_note_from_source")(
        SimpleNamespace(deps=deps),
        sheet=sheet,
        row=target_row,
        block_ids=["b1"],
        source_pages=[1],
        evidence="Page 1",
    ))

    assert result.startswith("rejected: workbook projection failed")
    assert not deps.wrote_once
    assert deps.cells_written == []
    with repo.db_session(db) as conn:
        assert conn.execute(
            "SELECT 1 FROM notes_cells WHERE run_id=? AND sheet=? AND row=?",
            (run_id, sheet, target_row),
        ).fetchone() is None
        assert srepo.fetch_usages(conn, generation_id) == []
        assert srepo.active_placements(conn, generation_id) == []


def test_source_projection_does_not_hold_sqlite_writer_lock(tmp_path, monkeypatch):
    """The workbook load/save phase must leave the DB writer slot available
    to sibling sheet agents and the SSE recorder."""
    kind = NotesTemplateType.CORP_INFO
    template = notes_template_path(kind, level="company")
    sheet = NOTES_REGISTRY[kind].sheet_name
    db = tmp_path / "audit.sqlite"
    init_db(db)
    template_id, nodes = parse_notes_template(str(template), sheet)
    import_notes_template(db, template_id, nodes)
    persist_template_manifest(db, template)
    with repo.db_session(db) as conn:
        run_id = repo.create_run(
            conn, "source.pdf", session_id="s", output_dir=str(tmp_path),
        )
        generation_id = srepo.begin_generation(
            conn, run_id, input_kind="prepared_document",
        )
        srepo.write_blocks(conn, generation_id, [SourceBlock(
            "b1", "paragraph", 1, "<p>Source disclosure.</p>",
            source_note_id="n1",
        )])
        srepo.write_notes(conn, generation_id, [
            SourceNote("n1", "1", "Disclosure", ["b1"]),
        ])
        srepo.activate_generation(conn, generation_id)

    target_row = min(node.row for node in nodes if node.kind == "LEAF")
    agent, deps = notes_agent.create_notes_agent(
        template_type=kind, pdf_path=str(tmp_path / "source.pdf"),
        inventory=[], filing_level="company", model=TestModel(),
        output_dir=str(tmp_path), run_id=run_id, db_path=str(db),
        source_generation_id=generation_id,
    )
    real_writer = notes_agent.write_notes_workbook

    def writer_while_recorder_writes(*args, **kwargs):
        with sqlite3.connect(str(db)) as competing:
            competing.execute("PRAGMA busy_timeout = 50")
            competing.execute("BEGIN IMMEDIATE")
            competing.rollback()
        return real_writer(*args, **kwargs)

    monkeypatch.setattr(
        notes_agent, "write_notes_workbook", writer_while_recorder_writes,
    )
    result = asyncio.run(_tool(agent, "write_note_from_source")(
        SimpleNamespace(deps=deps), sheet=sheet, row=target_row,
        block_ids=["b1"], source_pages=[1], evidence="Page 1",
    ))

    assert result.startswith("ok:"), result


def test_failed_final_promotion_restores_previous_artifacts_and_db(
    tmp_path, monkeypatch,
):
    from notes.writer import payload_sidecar_path

    kind = NotesTemplateType.CORP_INFO
    template = notes_template_path(kind, level="company")
    sheet = NOTES_REGISTRY[kind].sheet_name
    db = tmp_path / "audit.sqlite"
    init_db(db)
    template_id, nodes = parse_notes_template(str(template), sheet)
    import_notes_template(db, template_id, nodes)
    persist_template_manifest(db, template)
    with repo.db_session(db) as conn:
        run_id = repo.create_run(
            conn, "source.pdf", session_id="s", output_dir=str(tmp_path),
        )
        generation_id = srepo.begin_generation(
            conn, run_id, input_kind="prepared_document",
        )
        srepo.write_blocks(conn, generation_id, [
            SourceBlock("b1", "paragraph", 1, "<p>First version.</p>",
                        source_note_id="n1"),
            SourceBlock("b2", "paragraph", 2, "<p>Rejected version.</p>",
                        source_note_id="n1"),
        ])
        srepo.write_notes(conn, generation_id, [
            SourceNote("n1", "1", "Disclosure", ["b1", "b2"]),
        ])
        srepo.activate_generation(conn, generation_id)

    target_row = min(node.row for node in nodes if node.kind == "LEAF")
    agent, deps = notes_agent.create_notes_agent(
        template_type=kind, pdf_path=str(tmp_path / "source.pdf"),
        inventory=[], filing_level="company", model=TestModel(),
        output_dir=str(tmp_path), run_id=run_id, db_path=str(db),
        source_generation_id=generation_id,
    )
    tool = _tool(agent, "write_note_from_source")
    first = asyncio.run(tool(
        SimpleNamespace(deps=deps), sheet=sheet, row=target_row,
        block_ids=["b1"], source_pages=[1], evidence="Page 1",
    ))
    assert first.startswith("ok:"), first
    output = Path(deps.filled_path)
    sidecar = payload_sidecar_path(str(output))
    before_workbook = output.read_bytes()
    before_sidecar = sidecar.read_bytes()

    real_replace = notes_agent.os.replace

    def fail_staged_workbook(source, destination):
        if (str(source).endswith(".stage.xlsx")
                and Path(destination) == output):
            raise PermissionError("injected final promotion failure")
        return real_replace(source, destination)

    monkeypatch.setattr(notes_agent.os, "replace", fail_staged_workbook)
    rejected = asyncio.run(tool(
        SimpleNamespace(deps=deps), sheet=sheet, row=target_row,
        block_ids=["b2"], source_pages=[2], evidence="Page 2",
    ))

    assert rejected.startswith("rejected: workbook projection failed"), rejected
    assert output.read_bytes() == before_workbook
    assert sidecar.read_bytes() == before_sidecar
    with repo.db_session(db) as conn:
        cell = conn.execute(
            "SELECT html FROM notes_cells WHERE run_id=? AND sheet=? AND row=?",
            (run_id, sheet, target_row),
        ).fetchone()
        assert "First version" in cell["html"]
        assert [row["block_id"] for row in srepo.active_placements(
            conn, generation_id)] == ["b1"]
