"""Operator resolution must remain bound to the current fact and workbook."""
import copy

import pytest
from openpyxl import load_workbook

from test_mtool_template_map import _save_semantic_marker_workbook
from mtool.template_map import resolve_filing_doc


def _doc():
    return {"sheets": {"SOCIE": {"columns": {}}}, "writes": [{
        "concept_uuid": "fact-1", "sheet": "SOCIE", "label": "Profit or loss",
        "value": 100, "period": "CY", "entity_scope": "Company",
        "column_role": "current_year", "semantic_address": {
            "primary_concept": "ifrs-full_ProfitLoss", "dimensions": {},
        },
    }]}


def test_missing_category_can_be_explicitly_resolved_and_audited(tmp_path):
    path = tmp_path / "mtool.xlsx"
    _save_semantic_marker_workbook(path, two_periods=True)
    doc = _doc()
    _, report = resolve_filing_doc(str(path), doc)
    issue = report["unresolved_writes"][0]
    assert issue["period"] == "CY"
    assert issue["entity_scope"] == "Company"
    assert [o["cell"] for o in issue["resolution_options"]] == ["SOCIE!E5"]
    selection = {issue["resolution_key"]: "SOCIE!E5"}
    ready, report = resolve_filing_doc(str(path), doc, filing_targets=selection)
    assert ready["writes"][0]["cell"] == "E5"
    assert report["status"] == "attention"
    assert report["operator_resolutions"][0]["dimensions"] == {
        "ifrs-full_ComponentsOfEquityAxis": "ifrs-full_IssuedCapitalMember",
    }


def test_arbitrary_or_wrong_period_destination_is_rejected(tmp_path):
    path = tmp_path / "mtool.xlsx"
    _save_semantic_marker_workbook(path, two_periods=True)
    doc = _doc()
    _, report = resolve_filing_doc(str(path), doc)
    key = report["unresolved_writes"][0]["resolution_key"]
    for cell in ("SOCIE!E13", "SOCIE!Z99"):
        with pytest.raises(ValueError, match="destination"):
            resolve_filing_doc(str(path), doc, filing_targets={key: cell})


def test_selection_is_invalidated_when_fact_changes(tmp_path):
    path = tmp_path / "mtool.xlsx"
    _save_semantic_marker_workbook(path)
    doc = _doc()
    _, report = resolve_filing_doc(str(path), doc)
    key = report["unresolved_writes"][0]["resolution_key"]
    changed = copy.deepcopy(doc)
    changed["writes"][0]["value"] = 200
    with pytest.raises(ValueError, match="stale"):
        resolve_filing_doc(str(path), changed, filing_targets={key: "SOCIE!E5"})


def test_selection_is_invalidated_when_workbook_changes(tmp_path):
    path = tmp_path / "mtool.xlsx"
    _save_semantic_marker_workbook(path)
    doc = _doc()
    _, report = resolve_filing_doc(str(path), doc)
    key = report["unresolved_writes"][0]["resolution_key"]
    wb = load_workbook(path)
    wb.active["D5"] = "Revised row label"
    wb.save(path)
    with pytest.raises(ValueError, match="stale"):
        resolve_filing_doc(str(path), doc, filing_targets={key: "SOCIE!E5"})


def test_formula_destination_is_never_offered(tmp_path):
    path = tmp_path / "mtool.xlsx"
    _save_semantic_marker_workbook(path)
    wb = load_workbook(path)
    wb.active["E5"] = "=SUM(E6:E10)"
    wb.save(path)
    _, report = resolve_filing_doc(str(path), _doc())
    assert report["status"] == "blocked"
    assert report["unresolved_writes"][0]["resolution_options"] == []


def test_two_facts_cannot_be_placed_in_one_cell(tmp_path):
    path = tmp_path / "mtool.xlsx"
    _save_semantic_marker_workbook(path)
    doc = _doc()
    other = copy.deepcopy(doc["writes"][0])
    other["concept_uuid"] = "fact-2"
    doc["writes"].append(other)
    _, report = resolve_filing_doc(str(path), doc)
    selections = {i["resolution_key"]: "SOCIE!E5" for i in report["unresolved_writes"]}
    with pytest.raises(ValueError, match="same destination"):
        resolve_filing_doc(str(path), doc, filing_targets=selections)


def test_repeated_primary_is_resolved_by_exact_row_label_within_taxonomy_candidates(tmp_path):
    path = tmp_path / "mtool.xlsx"
    _save_semantic_marker_workbook(path)
    wb = load_workbook(path)
    ws = wb.active
    ws["A6"] = ws["A5"].value
    ws["D5"] = "Opening balance"
    ws["D6"] = "Closing balance"
    wb.save(path)
    doc = _doc()
    doc["writes"][0]["label"] = "*Opening balance"
    doc["writes"][0]["semantic_address"]["dimensions"] = {
        "ifrs-full_ComponentsOfEquityAxis": "ifrs-full_IssuedCapitalMember",
    }
    ready, report = resolve_filing_doc(str(path), doc)
    assert report["ambiguous"] == 0
    assert ready["writes"][0]["cell"] == "E5"


def test_duplicate_labels_still_require_operator_resolution(tmp_path):
    path = tmp_path / "mtool.xlsx"
    _save_semantic_marker_workbook(path)
    wb = load_workbook(path)
    ws = wb.active
    ws["A6"] = ws["A5"].value
    ws["D5"] = ws["D6"] = "Profit or loss"
    wb.save(path)
    doc = _doc()
    doc["writes"][0]["semantic_address"]["dimensions"] = {
        "ifrs-full_ComponentsOfEquityAxis": "ifrs-full_IssuedCapitalMember",
    }
    _, report = resolve_filing_doc(str(path), doc)
    assert report["ambiguous"] == 1
    issue = report["ambiguous_writes"][0]
    ready, report = resolve_filing_doc(str(path), doc,
        filing_targets={issue["resolution_key"]: "SOCIE!E6"})
    assert ready["writes"][0]["cell"] == "E6"
    assert report["operator_resolutions"]
