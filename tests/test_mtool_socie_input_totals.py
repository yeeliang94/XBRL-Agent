"""mTool has unlocked horizontal totals and locked vertical formulas."""
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection

from concept_model.cell_resolver import project_writes
from concept_model.cascade import recompute_after_turn
from db.schema import init_db
from mtool.exporter import build_fill_doc
from mtool.template_map import resolve_filing_doc
from mtool.offline_fill import fill_workbook
from tools.verifier import _resolve_cell_value
from test_mtool_exporter import _import, _init_run


def test_nested_horizontal_reserves_total_is_a_native_input(tmp_path):
    from openpyxl.utils import get_column_letter
    from mtool.offline_fill import verify_numeric_snapshot
    db = tmp_path / 'audit.db'
    init_db(db)
    run = _init_run(db)
    source = Path(__file__).resolve().parents[1] / 'XBRL-template-MFRS/Company/09-SOCIE.xlsx'
    tid = _import(db, source)
    project_writes(db, run, tid, [{'sheet': 'SOCIE', 'row': 6, 'col': 8, 'value': 50}])
    recompute_after_turn(db, run)
    doc = build_fill_doc(db, run, filing_standard='mfrs', filing_level='company')
    doc['writes'] = [w for w in doc['writes'] if w['label'] == '*Equity at beginning of period']
    doc['checks'] = []
    reserve = next(w for w in doc['writes'] if 'ifrs-full_OtherReservesMember' in w['semantic_address']['dimensions'].values())
    assert reserve['value_origin'] == 'canonical_calculation'
    assert reserve['value'] == 50
    book = Workbook(); sheet = book.active; sheet.title = 'SOCIE'
    prefix = 'synthetic.xsd#'
    sheet['C3'] = '#DOM#'; sheet['D3'] = '#PRIM#'; sheet['C5'] = '#ENDT#'
    sheet['A7'] = prefix + 'ifrs-full_Equity@periodStartLabel'
    sheet['D7'] = '*Equity at beginning of period'
    reserve_cell = None
    for index, write in enumerate(doc['writes'], 5):
        col = get_column_letter(index)
        member = next(iter(write['semantic_address']['dimensions'].values()))
        sheet[f'{col}2'] = '::'.join(prefix + item for item in [
            'ifrs-full_StatementOfChangesInEquityTable', 'ifrs-full_ComponentsOfEquityAxis', member])
        sheet[f'{col}5'] = '31/12/2026'
        if write is reserve:
            reserve_cell = f'{col}7'
    template = tmp_path / 'native.xlsx'; book.save(template); book.close()
    ready, coverage = resolve_filing_doc(str(template), doc)
    assert coverage['unmapped'] == coverage['ambiguous'] == 0, coverage
    output = tmp_path / 'filled.xlsx'
    assert fill_workbook(str(template), ready, str(output))['status'] == 'ok'
    assert verify_numeric_snapshot(str(output), ready)['status'] == 'ok'
    filled = load_workbook(output)
    assert filled.active[reserve_cell].value == 50
    filled.close()


@pytest.mark.parametrize("lock_total", [False, True])
@pytest.mark.parametrize("total_marker", [0, "0:::abc::abc::abc"])
def test_horizontal_totals_reconcile_without_overwriting_formulas(tmp_path, lock_total, total_marker):
    db = tmp_path / "audit.db"
    init_db(db)
    run_id = _init_run(db)
    source = Path(__file__).resolve().parents[1] / "XBRL-template-MFRS/Company/09-SOCIE.xlsx"
    tid = _import(db, source)
    result = project_writes(db, run_id, tid, [
        {"sheet": "SOCIE", "row": row, "col": 3, "value": value}
        for row, value in [(6, 1000), (11, 250), (17, 50)]
    ])
    assert result.projected == 3
    recompute_after_turn(db, run_id)
    doc = build_fill_doc(db, run_id, filing_standard="mfrs", filing_level="company")
    assert len(doc["writes"]) == 17
    assert sum(w.get("value_origin") == "canonical_calculation" for w in doc["writes"]) == 14

    wb = Workbook(); ws = wb.active; ws.title = "SOCIE"
    prefix = "full_ifrs-cor_2022-03-24.xsd#"
    for col, member in [("E", "RetainedEarningsMember"), ("F", "EquityAttributableToOwnersOfParentMember")]:
        ws[f"{col}2"] = "::".join(prefix + item for item in (
            "ifrs-full_StatementOfChangesInEquityTable", "ifrs-full_ComponentsOfEquityAxis", "ifrs-full_" + member))
    ws["G2"] = total_marker; ws["G4"] = "Total"
    ws["C3"] = "#DOM#"; ws["D3"] = "#PRIM#"; ws["C5"] = "#ENDT#"
    for col in "EFG":
        ws[f"{col}5"] = "31/12/2026"
        for row in (7, 8, 9):
            ws[f"{col}{row}"].protection = Protection(locked=lock_total and col == "G")
        ws[f"{col}10"] = f"={col}7+{col}8-{col}9"
        ws[f"{col}11"] = f"={col}8-{col}9"
        ws[f"{col}12"] = f"={col}7"
        ws[f"{col}13"] = f"={col}8"
    for row, primary, label in [
        (7, "Equity", "*Equity at beginning of period"),
        (8, "ProfitLoss", "*Profit (loss)"),
        (9, "DividendsPaid", "Dividends paid"),
        (10, "Equity", "*Equity at end of period"),
        (11, "ChangesInEquity", "*Total increase (decrease) in equity"),
        (13, "ComprehensiveIncome", "*Total comprehensive income"),
    ]:
        ws[f"A{row}"] = prefix + "ifrs-full_" + primary
        ws[f"D{row}"] = label
    ws['A12'] = 'ssmt-mfrs-cor_2022-12-31.xsd#ssmt-mfrs_EquityBalanceRestated'
    ws['D12'] = '*Equity at beginning of period, restated'
    ws.protection.sheet = True
    template, output = tmp_path / "mtool.xlsx", tmp_path / "draft.xlsx"
    wb.save(template)
    ready, coverage = resolve_filing_doc(str(template), doc)
    assert coverage["mapped"] == 17, coverage
    assert coverage['unmapped'] == coverage['ambiguous'] == 0, coverage
    report = fill_workbook(str(template), ready, str(output))
    filled = load_workbook(output)
    for col in "EFG":
        assert filled["SOCIE"][f"{col}10"].value == f"={col}7+{col}8-{col}9"
    assert _resolve_cell_value(filled, "SOCIE", "E10") == 1200
    assert _resolve_cell_value(filled, "SOCIE", "F10") == 1200
    if lock_total:
        assert len(report["unresolved"]) >= 3
        assert report["status"] == "degraded"
        assert filled["SOCIE"]["G7"].value is None
    else:
        assert report["status"] == "ok", report
        assert _resolve_cell_value(filled, "SOCIE", "G10") == 1200
    filled.close()
