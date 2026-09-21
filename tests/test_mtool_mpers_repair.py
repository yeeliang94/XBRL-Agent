"""MPERS native layout regressions; synthetic values, no model calls."""
import json
import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection

from mtool.column_detect import detect_column_map
from mtool.offline_fill import fill_workbook, resolve_label_to_note_cell
from mtool.template_map import resolve_filing_doc
from test_mtool_routes import client
from test_mtool_offline_fill import footnote_template

ROOT = Path(__file__).resolve().parents[1]


def sore_template(path):
    wb = Workbook(); ws = wb.active; ws.title = "StatementofRetainedEarnings"
    ws['A1'] = 'ssmt.xsd#ssmt-mpers_DisclosureOfRetainedEarningsAbstract'
    ws['D2'] = '#PRIM#'; ws['C3'] = '#ENDT#'
    ws['E3'] = '01/12/2025'; ws['F3'] = '31/12/2024'
    for row, primary, label, role in [
        (7, 'ifrs-smes_RetainedEarnings', 'Retained earnings at beginning of period', 'periodStartLabel'),
        (8, 'ssmt-mpers_ImpactOfChangesInAccountingPolicies', 'Impact of changes in accounting policies', ''),
        (9, 'ssmt-mpers_RetainedEarningsRestated', 'Retained earnings at beginning of period, restated', 'periodStartLabel'),
        (11, 'ifrs-smes_ProfitLoss', 'Profit (loss)', ''),
        (14, 'ifrs-smes_DividendsPaid', 'Dividends paid', ''),
        (16, 'ifrs-smes_RetainedEarnings', 'Retained earnings at end of period', 'periodEndLabel'),
    ]:
        ws[f'A{row}'] = f'sme.xsd#{primary}' + (f'@http://www.xbrl.org/2003/role/{role}' if role else '')
        ws[f'D{row}'] = '*' + label
        for col in 'EF':
            ws[f'{col}{row}'].protection = Protection(locked=False)
    for col in 'EF':
        ws[f'{col}9'] = f'={col}7+{col}8'
        ws[f'{col}16'] = f'={col}9+{col}11-{col}14'
        ws[f'{col}9'].protection = Protection(locked=True)
        ws[f'{col}16'].protection = Protection(locked=True)
    ws.protection.sheet = True
    wb.save(path)


def sore_doc():
    rows = [(7,'ifrs-smes_RetainedEarnings','Retained earnings at beginning of period',1000),
            (8,'ssmt-mpers_ImpactOfChangesInAccountingPolicies','Impact of changes in accounting policies',20),
            (9,'ssmt-mpers_RetainedEarningsRestated','Retained earnings at beginning of period, restated',1020),
            (11,'ifrs-smes_ProfitLoss','Profit (loss)',250),
            (14,'ifrs-smes_DividendsPaid','Dividends paid',50),
            (16,'ifrs-smes_RetainedEarnings','Retained earnings at end of period',1220)]
    return {'meta': {'filing_standard':'mpers','filing_level':'company'},
            'sheets': {'SoRE': {'columns': {'current_year':'B','prior_year':'C'}}},
            'writes': [{'sheet':'SoRE','label':label,'value':value,'period':period,
                        'entity_scope':'Company','column_role':role,'concept_uuid':str(row),
                        'semantic_address': {'primary_concept':primary,'dimensions':{}}}
                       for row,primary,label,value in rows
                       for period,role in [('CY','current_year'),('PY','prior_year')]]}


def test_mpers_sheet_alias_and_calculated_closing_are_distinct(tmp_path):
    source, output = tmp_path/'source.xlsx', tmp_path/'filled.xlsx'
    sore_template(source); original = source.read_bytes()
    doc = sore_doc()
    columns = detect_column_map(str(source), doc)
    assert columns['SoRE']['columns'] == {'current_year':'E','prior_year':'F'}
    ready, coverage = resolve_filing_doc(str(source), doc)
    assert coverage['mapped'] == 12, coverage
    assert not coverage['unmapped'] and not coverage['ambiguous']
    report = fill_workbook(str(source), ready, str(output))
    assert report['status'] == 'ok', report
    assert len(report['written']) == 8
    assert len(report['reconciled_formula']) == 4
    assert not report['skipped_formula']
    wb = load_workbook(output)
    assert wb.active['E7'].value == 1000
    assert wb.active['E16'].value == '=E9+E11-E14'
    wb.close()
    assert source.read_bytes() == original


@pytest.mark.parametrize('formula', ['=E9+E11-E14', '=UNSUPPORTED(E7)', '=E16+1'])
def test_formula_disagreement_or_unsupported_calculation_never_clean(tmp_path, formula):
    source, output = tmp_path/'source.xlsx', tmp_path/'filled.xlsx'
    sore_template(source)
    wb = load_workbook(source); wb.active['E16'] = formula; wb.save(source); wb.close()
    doc = sore_doc(); doc['writes'][-2]['value'] = 9999
    ready, coverage = resolve_filing_doc(str(source),doc)
    assert coverage['mapped'] == 12
    report = fill_workbook(str(source),ready,str(output))
    assert report['status'] == 'degraded'
    assert report['mismatches'] or report['unresolved']


def test_alias_requires_mpers_taxonomy_marker(tmp_path):
    source = tmp_path/'source.xlsx'; sore_template(source)
    wb = load_workbook(source)
    for row in wb.active:
        for cell in row:
            if isinstance(cell.value,str):
                cell.value = cell.value.replace('ssmt-mpers_', 'ssmt_').replace('ifrs-smes_', 'ifrs-full_')
    wb.save(source); wb.close()
    assert detect_column_map(str(source),sore_doc())['SoRE']['basis'] == 'missing'


@pytest.mark.parametrize('profit,expected_status', [(250, 'ok'), (300, 'degraded')])
def test_calculation_uses_current_cross_sheet_inputs(tmp_path, profit, expected_status):
    source = tmp_path / 'source.xlsx'
    sore_template(source)
    wb = load_workbook(source)
    wb.active['E16'] = "='SOPL-Function'!$E$11+E9-E14"
    wb.create_sheet('SOPL-Function')['E11'] = profit
    wb.save(source)
    wb.close()
    ready, _ = resolve_filing_doc(str(source), sore_doc())
    report = fill_workbook(str(source), ready, str(tmp_path / 'filled.xlsx'))
    assert report['status'] == expected_status
    if profit == 300:
        assert any(x['cell'] == 'E16' and x['found'] == '1270' for x in report['mismatches'])
    else:
        assert any(x['cell'] == 'E16' and x['found'] == '1220' for x in report['reconciled_formula'])


def test_repeated_exact_label_prefers_unique_input_for_leaf_fact(tmp_path):
    source = tmp_path / 'source.xlsx'
    sore_template(source)
    wb = load_workbook(source)
    # SOCF has both an input closing cash balance and a calculated reconciliation
    # with the same concept and visible label. Reproduce that layout compactly.
    wb.active['A18'] = wb.active['A16'].value
    wb.active['D18'] = wb.active['D16'].value
    for col in 'EF':
        wb.active[f'{col}18'] = 0
        wb.active[f'{col}18'].protection = Protection(locked=False)
    wb.save(source)
    wb.close()
    doc = sore_doc()
    doc['writes'] = doc['writes'][-2:]
    for write in doc['writes']:
        write['kind'] = 'LEAF'
    ready, coverage = resolve_filing_doc(str(source), doc)
    assert coverage['mapped'] == 2
    assert {w['cell'] for w in ready['writes']} == {'E18', 'F18'}
    assert not any(w.get('reconcile_formula') for w in ready['writes'])


@pytest.mark.parametrize("with_checks", [False, True])
def test_duplicate_sheet_identities_and_wrong_standard_do_not_resolve(tmp_path, with_checks):
    source=tmp_path/'ambiguous.xlsx';sore_template(source)
    wrong=sore_doc();wrong['meta']['filing_standard']='mfrs'
    if with_checks:
        wrong['checks'] = list(wrong['writes'])
    ready,coverage=resolve_filing_doc(str(source),wrong)
    assert coverage['status']=='blocked'
    assert not ready['writes']
    wb=load_workbook(source);ws=wb.copy_worksheet(wb.active);ws.title='SoRE';wb.save(source);wb.close()
    ready,coverage=resolve_filing_doc(str(source),sore_doc())
    assert not ready['writes']
    assert coverage['unmapped']==12


def test_blank_formula_input_still_reports_disagreement_with_canonical_total(tmp_path):
    source=tmp_path/'source.xlsx';sore_template(source)
    doc=sore_doc()
    doc['writes']=[w for w in doc['writes'] if w['semantic_address']['primary_concept']!='ssmt-mpers_ImpactOfChangesInAccountingPolicies']
    ready,_=resolve_filing_doc(str(source),doc)
    report=fill_workbook(str(source),ready,str(tmp_path/'filled.xlsx'))
    assert report['status']=='degraded'
    assert report['mismatches']
    assert not report['reconciled_formula']


def test_abstract_note_heading_is_not_a_text_destination():
    sheets = {'Notes-CI': {'label_col':'D','cells': {
        11: {'A':('S','tax.xsd#ssmt-mpers_CorporateInformationAbstract'), 'D':('S','Corporate information')},
        12: {'A':('S','tax.xsd#ssmt-mpers_DisclosureOfCorporateInformationExplanatory'), 'D':('S','*Disclosure of corporate information')},
    }}}
    result = resolve_label_to_note_cell('Disclosure of corporate information',sheets,'Notes-CI')
    assert result['status'] == 'resolved'
    assert result['cell'] == 'E12'


def test_text_block_is_rejected_by_numeric_projection(tmp_path):
    from concept_model.bootstrap import _import_one
    from concept_model.cell_resolver import project_writes
    from db.schema import init_db
    db = tmp_path/'facts.db'; init_db(db)
    tid = _import_one(db,ROOT/'XBRL-template-MPERS/Company/14-Notes-IssuedCapital.xlsx','company')
    with sqlite3.connect(db) as c:
        rid = c.execute("INSERT INTO runs(created_at,pdf_filename,status) VALUES ('test','test','completed')").lastrowid
        row = c.execute("SELECT render_row,render_sheet FROM concept_nodes n JOIN concept_semantic_addresses a USING(concept_uuid) WHERE n.template_id=? AND a.primary_concept='ifrs-smes_DisclosureOfClassesOfShareCapitalExplanatory'",(tid,)).fetchone()
    result = project_writes(db,rid,tid,[{'sheet':row[1],'row':row[0],'col':2,'value':123}])
    assert result.projected == 0
    assert result.rejected


def test_previous_manifest_refresh_quarantines_numeric_text_disclosure(tmp_path):
    from concept_model.bootstrap import _import_one
    from concept_model.filing_targets import MANIFEST_VERSION, persist_template_manifest
    from db.schema import init_db
    db = tmp_path / 'historical.db'
    template = ROOT / 'XBRL-template-MPERS/Company/14-Notes-IssuedCapital.xlsx'
    init_db(db)
    tid = _import_one(db, template, 'company')
    with sqlite3.connect(db) as conn:
        uuid = conn.execute(
            "SELECT canonical_target_id FROM template_slots WHERE template_id=? "
            "AND taxonomy_element_id='ifrs-smes_DisclosureOfClassesOfShareCapitalExplanatory'",
            (tid,),
        ).fetchone()[0]
        rid = conn.execute(
            "INSERT INTO runs(created_at,pdf_filename,status) VALUES ('test','test','completed')"
        ).lastrowid
        conn.execute(
            "INSERT INTO run_concept_facts(run_id,concept_uuid,period,entity_scope,value,value_status) "
            "VALUES (?,?,'CY','Company',123,'reported')", (rid, uuid),
        )
        conn.execute(
            "UPDATE template_slots SET manifest_version='2022-v1-slot-semantics-2', "
            "value_kind='numeric' WHERE template_id=?", (tid,),
        )
    persist_template_manifest(db, template)
    with sqlite3.connect(db) as conn:
        assert conn.execute(
            "SELECT DISTINCT value_kind,manifest_version FROM template_slots "
            "WHERE canonical_target_id=?", (uuid,),
        ).fetchall() == [('html', MANIFEST_VERSION)]
        assert conn.execute(
            "SELECT value,invalid_target FROM run_concept_facts WHERE concept_uuid=?", (uuid,),
        ).fetchone() == (123, 1)


def test_missing_category_period_is_reported_before_category_choices(tmp_path):
    source = tmp_path/'category.xlsx'
    wb=Workbook();ws=wb.active;ws.title='Notes-IssuedCap'
    ws['D3']='#PRIM#';ws['C3']='#DOM#';ws['C4']='#ENDT#';ws['E4']='01/12/2025'
    ws['E2']='::'.join('sme.xsd#'+i for i in (
        'ifrs-smes_DisclosureOfShareCapitalTable','ifrs-smes_ClassesOfShareCapitalAxis','ssmt-mpers_OrdinarySharesMember'))
    ws['A7']='sme.xsd#ifrs-smes_NumberOfSharesIssuedAndFullyPaid';ws['D7']='Number of shares issued and fully paid'
    wb.save(source)
    doc={'sheets':{'Notes-Issuedcapital':{'columns':{'current_year':'B','prior_year':'C'}}},
         'writes':[{'sheet':'Notes-Issuedcapital','label':'Number of shares issued and fully paid','value':100,
                    'period':'CY','entity_scope':'Company','column_role':'current_year',
                    'semantic_address':{'primary_concept':'ifrs-smes_NumberOfSharesIssuedAndFullyPaid','dimensions':{}}}]}
    _,report=resolve_filing_doc(str(source),doc)
    issue=report['unresolved_writes'][0]
    assert issue['reason_code']=='missing_category_dimensions'
    assert [x['cell'] for x in issue['resolution_options']]==['Notes-IssuedCap!E7']
    ready,report=resolve_filing_doc(str(source),doc,filing_targets={issue['resolution_key']:'Notes-IssuedCap!E7'})
    assert report['mapped']==1
    assert ready['writes'][0]['canonical_sheet']=='Notes-Issuedcapital'
    doc['writes'][0].update(period='PY',column_role='prior_year')
    _,report=resolve_filing_doc(str(source),doc)
    assert report['unresolved_writes'][0]['reason_code']=='template_period_section_missing'
    assert not report['unresolved_writes'][0]['resolution_options']


def test_mpers_preparation_endpoint_preserves_calculations_and_receipt(client, tmp_path):
    import io
    from concept_model.bootstrap import _import_one
    from concept_model.cell_resolver import project_writes
    from concept_model.cascade import recompute_after_turn
    tc,db,_=client
    tid=_import_one(db,ROOT/'XBRL-template-MPERS/Company/10-SoRE.xlsx','company')
    with sqlite3.connect(db) as c:
        rid=c.execute("INSERT INTO runs(created_at,pdf_filename,status,run_config_json) VALUES ('test','test','completed',?)",
                      (json.dumps({'filing_standard':'mpers','filing_level':'company'}),)).lastrowid
    result=project_writes(db,rid,tid,[{'sheet':'SoRE','row':r,'col':2,'value':v}
        for r,v in [(7,1000),(8,20),(9,1020),(11,250),(14,50),(16,1220)]])
    assert result.projected==6
    recompute_after_turn(db,rid)
    source=tmp_path/'sore.xlsx';sore_template(source)
    response=tc.post(f'/api/runs/{rid}/mtool-fill/patch',
        files={'template':('sore.xlsx',source.read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
        data={'fill_notes':'false','selected_sheets':json.dumps(['SoRE'])})
    assert response.status_code==200,response.text
    body=response.json()
    assert body['counts']['reconciled_formula']==2
    assert body['counts']['written']==4
    assert body['filing_coverage']['mapped']==6
    with sqlite3.connect(db) as c:
        report=json.loads(c.execute('SELECT report_json FROM mtool_fill_receipts WHERE id=?',(body['receipt_id'],)).fetchone()[0])
    assert len(report['reconciled_formula'])==2
    downloaded=tc.get(body['download_url'],params={'acknowledge_degraded':'reviewed'})
    assert downloaded.status_code==200
    wb=load_workbook(io.BytesIO(downloaded.content))
    assert wb.active['E16'].value=='=E9+E11-E14'
    assert wb.active['E7'].value==1000
    wb.close()


def test_supplied_mpers_workbook_balanced_sore(tmp_path):
    source=ROOT/'data/FS-MPERS-Test_Sdn_Bhd-12345678910-01122025.xlsx'
    if not source.exists(): pytest.skip('supplied MPERS workbook unavailable')
    ready,coverage=resolve_filing_doc(str(source),sore_doc())
    assert coverage['mapped']==12,coverage
    output=tmp_path/'filled.xlsx'
    report=fill_workbook(str(source),ready,str(output))
    assert report['status']=='ok',report
    assert {x['cell']:x['found'] for x in report['reconciled_formula']}=={
        'E41':'1020','F41':'1020','E48':'1220','F48':'1220'}


@pytest.mark.parametrize('target', [
    {'sheet': 'Notes-CI', 'cell': 'E20'},
    {'key': 'fn_14'},
    {'key': 'fn_999'},
    {'key': 'fn_99'},  # Payload exists, but no visible taxonomy-linked destination.
])
def test_semantic_prose_target_and_explicit_override_share_identity_guard(tmp_path, footnote_template, target):
    from zipfile import ZipFile
    from mtool.offline_fill import fill_footnotes, write_patched_zip
    source=tmp_path/'semantic.xlsx'
    primary='ssmt-mpers_DisclosureOfCorporateInformationExplanatory'
    with ZipFile(footnote_template) as z:
        xml=z.read('xl/worksheets/sheet2.xml').decode()
    xml=xml.replace('<row r="14">', '<row r="14"><c r="A14" t="inlineStr"><is><t>tax.xsd#'+primary+'</t></is></c>')
    write_patched_zip(str(footnote_template),str(source),{'xl/worksheets/sheet2.xml':xml.encode()})
    note={'source_sheet':'Notes-CI','label':'A different visible label','primary_concept':primary,'html':'<p>Test</p>'}
    report=fill_footnotes(str(source),{'footnotes':[note]},str(tmp_path/'out.xlsx'),create_missing=True)
    assert report['status']=='ok',report
    assert report['footnotes_written'][0]['visible_cell']=='Notes-CI!E14'
    note.update(target)
    report=fill_footnotes(str(source),{'footnotes':[note]},str(tmp_path/'wrong.xlsx'),create_missing=True)
    assert report['status']=='degraded'
    assert not report['footnotes_written']
    issue = report['unresolved'][0]
    assert issue['reason'] == 'identity_mismatch'
    assert 'Nothing was written' in issue['detail']
    assert issue['candidates'] == [{'sheet': 'Notes-CI', 'cell': 'E14', 'label_cell': 'D14'}]
    from api.mtool import _unresolved_entry, _apply_notes_targets
    assert _unresolved_entry(issue)['candidates'] == issue['candidates']
    retry = {'footnotes': [{k: v for k, v in note.items() if k not in ('key', 'cell', 'sheet')}]}
    _apply_notes_targets(retry, json.dumps({'0': {'sheet': 'Notes-CI', 'cell': 'E14'}}))
    report = fill_footnotes(str(source), retry, str(tmp_path/'retry.xlsx'), create_missing=True)
    assert report['status'] == 'ok', report
    assert report['footnotes_written'][0]['cell'] == 'E14'


def test_note_identity_rejection_without_matching_destinations(tmp_path, footnote_template):
    from mtool.offline_fill import fill_footnotes
    output = tmp_path / 'unresolved.xlsx'
    report = fill_footnotes(footnote_template, {'footnotes': [{
        'source_sheet': 'Notes-CI', 'sheet': 'Notes-CI', 'cell': 'E14',
        'primary_concept': 'ssmt-mpers_DisclosureOfCorporateInformationExplanatory',
        'html': '<p>Corporate information</p>',
    }]}, str(output), create_missing=True)
    assert report['status'] == 'degraded'
    assert not report['footnotes_written']
    issue = report['unresolved'][0]
    assert issue['reason'] == 'identity_mismatch'
    assert issue['candidates'] == []
    assert 'no verified destination' in issue['detail']
    from zipfile import ZipFile
    from mtool.offline_fill import get_shared_strings
    with ZipFile(footnote_template) as before, ZipFile(output) as after:
        for name in before.namelist():
            if name == 'xl/sharedStrings.xml':
                # The normal fill path repairs string-reference counters even
                # on a rejected placement; stored note content must not change.
                assert get_shared_strings({name: before.read(name)}) == get_shared_strings({name: after.read(name)})
            else:
                assert before.read(name) == after.read(name)


def test_supplied_mpers_all_prose_slots_keep_their_taxonomy_and_content(tmp_path):
    from concept_model.notes_parser import parse_notes_template
    from concept_model.notes_importer import import_notes_template
    from concept_model.filing_targets import persist_template_manifest
    from mtool.notes_exporter import build_notes_fill_doc
    from mtool.offline_fill import fill_footnotes
    from db.schema import init_db
    source=ROOT/'data/FS-MPERS-Test_Sdn_Bhd-12345678910-01122025.xlsx'
    if not source.exists(): pytest.skip('supplied MPERS workbook unavailable')
    db=tmp_path/'notes.db';init_db(db)
    with sqlite3.connect(db) as c:
        rid=c.execute("INSERT INTO runs(created_at,pdf_filename,status,run_config_json) VALUES ('test','test','completed',?)",
                      (json.dumps({'filing_standard':'mpers','filing_level':'company'}),)).lastrowid
    count=0
    for file,sheet in [('11-Notes-CorporateInfo.xlsx','Notes-CI'),('12-Notes-AccountingPolicies.xlsx','Notes-SummaryofAccPol'),('13-Notes-ListOfNotes.xlsx','Notes-Listofnotes')]:
        path=ROOT/'XBRL-template-MPERS/Company'/file
        tid,nodes=parse_notes_template(str(path),sheet)
        import_notes_template(db,tid,nodes);persist_template_manifest(db,path)
        with sqlite3.connect(db) as c:
            for n in nodes:
                if n.slot_role!='INPUT':continue
                count+=1
                c.execute("INSERT INTO notes_cells(run_id,sheet,row,label,html,updated_at,concept_uuid) VALUES (?,?,?,?,?,'test',?)",
                          (rid,n.sheet,n.row,n.label,f'<p>TEST {count}</p>',n.node_uuid))
    doc=build_notes_fill_doc(db,rid)
    assert len(doc['footnotes'])==count==138
    assert all(n.get('primary_concept') for n in doc['footnotes'])
    original=source.read_bytes()
    report=fill_footnotes(str(source),doc,str(tmp_path/'notes.xlsx'),create_missing=True)
    assert report['status']=='ok',report
    assert len(report['footnotes_written'])==138
    assert not report['footnote_mismatches']
    assert source.read_bytes()==original


def test_supplied_mpers_api_category_retry_notes_receipt_and_download(client):
    import io
    from zipfile import ZipFile
    from xml.etree import ElementTree as ET
    from concept_model.bootstrap import _import_one
    from concept_model.cell_resolver import project_writes
    from concept_model.cascade import recompute_after_turn
    from concept_model.notes_parser import parse_notes_template
    from concept_model.notes_importer import import_notes_template
    from concept_model.filing_targets import persist_template_manifest
    source=ROOT/'data/FS-MPERS-Test_Sdn_Bhd-12345678910-01122025.xlsx'
    if not source.exists(): pytest.skip('supplied MPERS workbook unavailable')
    tc,db,_=client
    with sqlite3.connect(db) as c:
        rid=c.execute("INSERT INTO runs(created_at,pdf_filename,status,run_config_json) VALUES ('test','test','completed',?)",
                      (json.dumps({'filing_standard':'mpers','filing_level':'company'}),)).lastrowid
    tid=_import_one(db,ROOT/'XBRL-template-MPERS/Company/10-SoRE.xlsx','company')
    project_writes(db,rid,tid,[{'sheet':'SoRE','row':r,'col':2,'value':v}
        for r,v in [(7,1000),(8,20),(9,1020),(11,250),(14,50),(16,1220)]])
    tid=_import_one(db,ROOT/'XBRL-template-MPERS/Company/14-Notes-IssuedCapital.xlsx','company')
    with sqlite3.connect(db) as c:
        row=c.execute("SELECT n.render_row FROM concept_nodes n JOIN concept_semantic_addresses a USING(concept_uuid) WHERE n.template_id=? AND a.primary_concept='ifrs-smes_NumberOfSharesIssuedAndFullyPaid'",(tid,)).fetchone()[0]
    project_writes(db,rid,tid,[{'sheet':'Notes-Issuedcapital','row':row,'col':2,'value':100}])
    recompute_after_turn(db,rid)
    sheets=['SoRE','Notes-Issuedcapital']
    for filename,sheet in [('11-Notes-CorporateInfo.xlsx','Notes-CI'),('12-Notes-AccountingPolicies.xlsx','Notes-SummaryofAccPol'),('13-Notes-ListOfNotes.xlsx','Notes-Listofnotes')]:
        path=ROOT/'XBRL-template-MPERS/Company'/filename
        tid,nodes=parse_notes_template(str(path),sheet)
        import_notes_template(db,tid,nodes);persist_template_manifest(db,path)
        n=next(n for n in nodes if n.slot_role=='INPUT')
        with sqlite3.connect(db) as c:
            c.execute("INSERT INTO notes_cells(run_id,sheet,row,label,html,updated_at,concept_uuid) VALUES (?,?,?,?,?,'test',?)",
                      (rid,n.sheet,n.row,n.label,'<p>TEST ONLY</p>',n.node_uuid))
        sheets.append(sheet)
    raw=source.read_bytes()
    files={'template':('mpers.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    data={'selected_sheets':json.dumps(sheets)}
    response=tc.post(f'/api/runs/{rid}/mtool-fill/patch',files=files,data=data)
    assert response.status_code==200,response.text
    partial=response.json()
    assert partial['status']=='degraded'
    assert partial['counts']['written']==4
    assert partial['notes']['counts']['written']==3
    issue=partial['filing_coverage']['unresolved_writes'][0]
    option=next(o for o in issue['resolution_options'] if o['dimensions']=={
        'ifrs-smes_ClassesOfShareCapitalAxis':'ssmt-mpers_OrdinarySharesMember'})
    data['filing_targets']=json.dumps({issue['resolution_key']:option['cell']})
    response=tc.post(f'/api/runs/{rid}/mtool-fill/patch',files=files,data=data)
    assert response.status_code==200,response.text
    body=response.json()
    assert body['counts']['written']==5
    # Verify the intermediate profit/movement totals as well as opening/closing
    # retained earnings; each native formula must agree with the snapshot.
    reconciled=body['reconciled_formula']
    assert body['counts']['reconciled_formula']==len(reconciled)==4
    assert {(r['sheet'],r['cell'],r['value'],r['found']) for r in reconciled}=={
        ('StatementofRetainedEarnings','E41',1020,'1020'),
        ('StatementofRetainedEarnings','E44',250,'250'),
        ('StatementofRetainedEarnings','E47',200,'200'),
        ('StatementofRetainedEarnings','E48',1220,'1220'),
    }
    assert body['notes']['counts']['written']==3
    assert not body['notes']['unresolved']
    with sqlite3.connect(db) as c:
        receipt=c.execute('SELECT snapshot_notes_count,report_json FROM mtool_fill_receipts WHERE id=?',(body['receipt_id'],)).fetchone()
    assert receipt[0]==3
    assert json.loads(receipt[1])['filing_coverage']['operator_resolutions'][0]['cell']==option['cell']
    response=tc.get(body['download_url'],params={'acknowledge_degraded':'reviewed'})
    assert response.status_code==200
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(io.BytesIO(raw)) as before,ZipFile(io.BytesIO(response.content)) as after:
        assert before.read('xl/styles.xml')==after.read('xl/styles.xml')
        count=0
        for name in before.namelist():
            if not name.startswith('xl/worksheets/') or not name.endswith('.xml'):continue
            def formulas(blob):
                return {c.get('r'):ET.tostring(c.find('m:f',ns))
                        for c in ET.fromstring(blob).findall('.//m:c',ns) if c.find('m:f',ns) is not None}
            a,b=formulas(before.read(name)),formulas(after.read(name))
            assert a==b,name
            count+=len(a)
        assert count==402
    assert source.read_bytes()==raw
