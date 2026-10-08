"""Read-only source tools on the notes agent — plan Phase 6, Step 6.1.

Three properties the plan calls out, each tested here:

* the tools are registered ONLY when the run has a frozen source reading, so
  an `off`-mode or PDF run sees exactly the agent it saw before;
* every response is capped in bytes — an uncapped `view_source_blocks`
  recreates the context problem the 60,000-char snippet cap was built to
  solve, one call at a time;
* document content keeps the untrusted-content framing `read_source_note`
  already uses.
"""
from __future__ import annotations

import pytest

from db import repository as repo
from db.schema import init_db
from notes import agent as notes_agent
from notes import source_repository as srepo
from notes_types import NotesTemplateType
from notes.source_models import SourceBlock, SourceNote

BLOCKS = [
    SourceBlock(block_id="b1", block_kind="heading", reading_order=0,
                canonical_html="<h3>5. Receivables</h3>", source_note_id="n5"),
    SourceBlock(block_id="b2", block_kind="paragraph", reading_order=1,
                canonical_html="<p>Stated at cost.</p>", source_note_id="n5"),
    SourceBlock(block_id="b3", block_kind="table", reading_order=2,
                canonical_html="<table><tr><td>a</td></tr></table>",
                source_note_id="n6"),
]


@pytest.fixture()
def seeded(tmp_path):
    db = tmp_path / "audit.sqlite"
    init_db(db)
    with repo.db_session(db) as conn:
        run_id = repo.create_run(
            conn, "x.docx", session_id="s", output_dir=str(tmp_path / "s")
        )
        gen = srepo.begin_generation(conn, run_id, input_kind="docx_html")
        srepo.write_blocks(conn, gen, BLOCKS)
        srepo.write_notes(conn, gen, [
            SourceNote(source_note_id="n5", top_note_num="5",
                       title="Receivables", block_ids=["b1", "b2"]),
            SourceNote(source_note_id="n6", top_note_num="6",
                       title="Cash", block_ids=["b3"]),
        ])
        srepo.activate_generation(conn, gen)
    return str(db), run_id, gen


# --------------------------------------------------------------------------
# registration
# --------------------------------------------------------------------------

def _tool_names(agent) -> set[str]:
    """Same introspection the existing agent-tool tests use."""
    names: set[str] = set()
    for ts in getattr(agent, "toolsets", []) or []:
        tools = getattr(ts, "tools", {}) or {}
        if isinstance(tools, dict):
            names.update(tools.keys())
    return names


@pytest.fixture()
def assigned_agent(tmp_path, seeded):
    """One Sheet-12 worker with a real live label and a frozen assigned note."""
    from pydantic_ai.models.test import TestModel
    db, run_id, gen = seeded
    with repo.db_session(db) as conn:
        conn.execute("UPDATE notes_source_generations SET input_kind='prepared_document' WHERE id=?", (gen,))
        conn.execute(
            "INSERT INTO notes_nodes(node_uuid,template_id,sheet,row,label,kind) "
            "VALUES ('managed-target','mfrs-company-notes-list-v1',"
            "'Notes-Listofnotes',140,'Disclosure of trade and other receivables','LEAF')"
        )
    agent, deps = notes_agent.create_notes_agent(
        template_type=NotesTemplateType.LIST_OF_NOTES, pdf_path="synthetic.pdf",
        inventory=[], filing_level="company", model=TestModel(), output_dir=str(tmp_path),
        run_id=run_id, db_path=db, source_generation_id=gen, batch_note_nums=[5],
    )
    deps.payload_sink = []
    tools = {name: tool for ts in agent.toolsets
             for name, tool in getattr(ts, "tools", {}).items()}
    return deps, tools


def test_managed_whole_note_tools_expose_only_meaningful_choices(assigned_agent):
    deps, tools = assigned_agent
    read_schema = tools["read_assigned_note"].function_schema.json_schema
    write_schema = tools["write_assigned_note"].function_schema.json_schema
    assert set(read_schema["properties"]) == {"note_num", "offset"}
    assert set(write_schema["properties"]) == {"note_num", "destination_label"}
    assert "write_note_from_source" in tools  # Mixed notes retain section routing.


@pytest.mark.asyncio
async def test_assigned_note_read_is_complete_bounded_and_has_no_copyable_ids(assigned_agent, monkeypatch):
    from types import SimpleNamespace
    deps, tools = assigned_agent
    ctx = SimpleNamespace(deps=deps)
    read = tools["read_assigned_note"].function
    full = await read(ctx, 5)
    body = lambda text: text.split("<<<SOURCE>>>\n", 1)[1].split("\n<<<END_SOURCE>>>", 1)[0]
    assert "UNTRUSTED" in full
    assert "<h3>5. Receivables</h3>" in full and "<p>Stated at cost.</p>" in full
    assert "b1" not in full and "b2" not in full and "b3" not in full
    monkeypatch.setattr(notes_agent, "SOURCE_TOOL_RESPONSE_CAP", 40)
    offset, chunks = 0, []
    while True:
        response = await read(ctx, 5, offset)
        assert len(body(response)) <= 40
        chunks.append(body(response))
        if "next_offset=" not in response:
            break
        offset = int(response.split("next_offset=", 1)[1].split(".", 1)[0])
    assert "".join(chunks) == body(full)
    assert "Invalid offset" in await read(ctx, 5, -1)


@pytest.mark.asyncio
async def test_managed_whole_note_persists_same_canonical_content_as_source_write(assigned_agent):
    from types import SimpleNamespace
    deps, tools = assigned_agent
    ctx = SimpleNamespace(deps=deps)
    direct = tools["write_note_from_source"].function
    managed = tools["write_assigned_note"].function
    label = "Disclosure of trade and other receivables"
    await tools["read_assigned_note"].function(ctx, 5)
    assert (await direct(ctx, deps.sheet_name, 140, label, ["b1", "b2"])).startswith("ok:")
    with repo.db_session(deps.db_path) as conn:
        before = conn.execute("SELECT html,concept_uuid,source_pages,content_origin FROM notes_cells WHERE run_id=?", (deps.run_id,)).fetchone()
        before = tuple(before)
        assert before[3] == "source_exact"
    result = await managed(ctx, 5, label)
    assert result.startswith("ok:") and "Placed complete Note 5" in result
    assert "saved source parts" not in result and "b1" not in result and "b2" not in result
    with repo.db_session(deps.db_path) as conn:
        after = conn.execute("SELECT html,concept_uuid,source_pages,content_origin FROM notes_cells WHERE run_id=?", (deps.run_id,)).fetchone()
        assert tuple(after) == before
        assert {p["block_id"] for p in srepo.active_placements(conn, deps.source_generation_id)} == {"b1", "b2"}
    assert len(deps.payload_sink) == 1
    assert deps.payload_sink[0].source_built and deps.payload_sink[0].note_num == 5
    # A Sheet-12 worker retains payloads; its coordinator projects the workbook.
    assert deps.payload_sink[0].content == before[0]

    # A worker corrects its own whole-note destination by moving the canonical
    # cell, rather than proposing a duplicate placement through another write.
    corrected_label = "Disclosure of deferred income"
    with repo.db_session(deps.db_path) as conn:
        conn.execute("INSERT INTO notes_nodes(node_uuid,template_id,sheet,row,label,kind) "
                     "VALUES ('corrected-target','mfrs-company-notes-list-v1',"
                     "'Notes-Listofnotes',32,?,'LEAF')", (corrected_label,))
        conn.execute("INSERT INTO template_slots("
                     "target_id,canonical_target_id,template_id,sheet,row,col,label,"
                     "slot_role,value_kind,mapping_source,manifest_version,"
                     "workbook_fingerprint,validation_status) "
                     "VALUES ('corrected-target','corrected-target','mfrs-company-notes-list-v1',"
                     "'Notes-Listofnotes',32,'B',?,'INPUT','html','test','test-v1','fixture','writable')",
                     (corrected_label,))
    deps.coverage_receipt = object()
    moved = await tools["move_own_source_cell"].function(
        ctx, 140, 32, corrected_label, "Correct my original destination after reading the note.",
    )
    assert moved.startswith("Moved")
    assert deps.coverage_receipt is None
    assert [payload.chosen_row_label for payload in deps.payload_sink] == [corrected_label]
    with repo.db_session(deps.db_path) as conn:
        cells = conn.execute("SELECT row,html FROM notes_cells WHERE run_id=?", (deps.run_id,)).fetchall()
        assert [(cell["row"], cell["html"]) for cell in cells] == [(32, before[0])]
        assert {p["row"] for p in srepo.active_placements(conn, deps.source_generation_id)} == {32}
        assert repo.fetch_notes_review_flags(conn, deps.run_id) == []
    receipt = [{"note_num": 5, "action": "written", "row_labels": [corrected_label]}]
    assert "accepted" in notes_agent._submit_coverage_entries_impl(deps, receipt)


@pytest.mark.asyncio
@pytest.mark.parametrize("failure", ["unassigned", "missing_note", "ambiguous_note", "duplicate_label", "unknown_label", "stale", "foreign_run"])
async def test_managed_whole_note_rejects_invalid_context_without_writing(assigned_agent, failure):
    from types import SimpleNamespace
    deps, tools = assigned_agent
    note_num, label = 5, "Disclosure of trade and other receivables"
    if failure == "unassigned":
        note_num = 6
    elif failure == "missing_note":
        deps.batch_note_nums.append(7)
        note_num = 7
    elif failure == "foreign_run":
        deps.run_id += 100
    elif failure == "unknown_label":
        label = "Receivables"
    else:
        with repo.db_session(deps.db_path) as conn:
            if failure == "stale":
                conn.execute("UPDATE notes_source_generations SET status='superseded' WHERE id=?", (deps.source_generation_id,))
            elif failure == "duplicate_label":
                conn.execute("INSERT INTO notes_nodes(node_uuid,template_id,sheet,row,label,kind) VALUES ('duplicate-target','mfrs-company-notes-list-v1','Notes-Listofnotes',141,?,'LEAF')", (label,))
            elif failure == "ambiguous_note":
                conn.execute("UPDATE notes_source_notes SET top_note_num='5' WHERE generation_id=? AND source_note_id='n6'", (deps.source_generation_id,))
    ctx = SimpleNamespace(deps=deps)
    await tools["read_assigned_note"].function(ctx, note_num)
    response = await tools["write_assigned_note"].function(ctx, note_num, label)
    assert response.startswith("rejected:")
    if failure in {"unassigned", "missing_note", "ambiguous_note", "stale", "foreign_run"}:
        assert (await tools["read_assigned_note"].function(ctx, note_num)).startswith("rejected:")
    assert not deps.payload_sink and not deps.wrote_once
    with repo.db_session(deps.db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM notes_cells").fetchone()[0] == 0


@pytest.mark.asyncio
async def test_assigned_note_preserves_capture_uncertainty(assigned_agent):
    from types import SimpleNamespace
    deps, tools = assigned_agent
    with repo.db_session(deps.db_path) as conn:
        conn.execute("UPDATE notes_source_blocks SET locator_json=? WHERE generation_id=? AND block_id='b2'", ('{"capture_uncertain":true}', deps.source_generation_id))
    ctx = SimpleNamespace(deps=deps)
    assert "original wording is uncertain" in await tools["read_assigned_note"].function(ctx, 5)
    assert (await tools["write_assigned_note"].function(ctx, 5, "Disclosure of trade and other receivables")).startswith("ok:")
    with repo.db_session(deps.db_path) as conn:
        assert conn.execute("SELECT content_origin FROM notes_cells WHERE run_id=?", (deps.run_id,)).fetchone()[0] == "vision_transcribed"


@pytest.mark.asyncio
async def test_assigned_note_conflict_uses_existing_durable_review_path(assigned_agent):
    from types import SimpleNamespace
    deps, tools = assigned_agent
    ctx = SimpleNamespace(deps=deps)
    await tools["read_assigned_note"].function(ctx, 5)
    first = await tools["write_assigned_note"].function(ctx, 5, "Disclosure of trade and other receivables")
    assert first.startswith("ok:")
    with repo.db_session(deps.db_path) as conn:
        conn.execute("INSERT INTO notes_nodes(node_uuid,template_id,sheet,row,label,kind) VALUES ('other-managed-target','mfrs-company-notes-list-v1','Notes-Listofnotes',32,'Disclosure of deferred income','LEAF')")
    conflict = await tools["write_assigned_note"].function(ctx, 5, "Disclosure of deferred income")
    assert conflict.startswith("conflict recorded for review:")
    with repo.db_session(deps.db_path) as conn:
        assert len(repo.fetch_notes_review_flags(conn, deps.run_id)) == 1
        assert [row["row"] for row in conn.execute("SELECT row FROM notes_cells WHERE run_id=?", (deps.run_id,))] == [140]
    assert len(deps.payload_sink) == 1


@pytest.mark.asyncio
async def test_assigned_whole_note_requires_contiguous_complete_read(assigned_agent, monkeypatch):
    from types import SimpleNamespace
    deps, tools = assigned_agent
    ctx = SimpleNamespace(deps=deps)
    read = tools["read_assigned_note"].function
    write = tools["write_assigned_note"].function
    label = "Disclosure of trade and other receivables"
    assert (await write(ctx, 5, label)).startswith("rejected: read the complete")
    monkeypatch.setattr(notes_agent, "SOURCE_TOOL_RESPONSE_CAP", 40)
    # Invalid or tail-first reads do not account for omitted earlier content.
    assert "Invalid offset" in await read(ctx, 5, -1)
    await read(ctx, 5, 80)
    assert (await write(ctx, 5, label)).startswith("rejected: read the complete")
    response = await read(ctx, 5)
    assert (await write(ctx, 5, label)).startswith("rejected: read the complete")
    while "next_offset=" in response:
        offset = int(response.split("next_offset=", 1)[1].split(".", 1)[0])
        response = await read(ctx, 5, offset)
    assert (await write(ctx, 5, label)).startswith("ok:")


@pytest.mark.parametrize("template,batch,source", [
    (NotesTemplateType.LIST_OF_NOTES, None, True),
    (NotesTemplateType.LIST_OF_NOTES, [5], False),
    (NotesTemplateType.CORP_INFO, [5], True),
])
def test_managed_source_tools_are_only_registered_for_assigned_source_batches(tmp_path, seeded, template, batch, source):
    from pydantic_ai.models.test import TestModel
    db, run_id, gen = seeded
    agent, _ = notes_agent.create_notes_agent(
        template_type=template, pdf_path="synthetic.pdf", inventory=[], filing_level="company",
        model=TestModel(), output_dir=str(tmp_path), run_id=run_id, db_path=db,
        source_generation_id=gen if source else None, batch_note_nums=batch,
    )
    assert {"read_assigned_note", "write_assigned_note"}.isdisjoint(_tool_names(agent))


def _seed_mixed_source_sections(deps):
    """Complete source sections sharing only their verified parent heading."""
    blocks = [SourceBlock("b1", "heading", 0, "<h2>5 Basis of preparation</h2>", source_note_id="n5")]
    targets = [
        (133, "Disclosure of statement of compliance", "Statement of compliance", "The statements comply with MFRS."),
        (10, "Disclosure of basis of preparation of financial statements", "Basis of measurement", "Historical cost and going concern; liabilities exceed assets by RM60,606,000."),
        (31, "Disclosure of critical accounting estimates and judgements", "Estimates and judgements", "Estimates are reviewed on an ongoing basis."),
    ]
    with repo.db_session(deps.db_path) as conn:
        for index, (row, label, title, prose) in enumerate(targets, 1):
            heading = f"s{index}-heading"
            blocks.extend([
                SourceBlock(heading, "heading", index * 2 - 1, f"<h3>5.{index} {title}</h3>", source_note_id="n5", locator={"heading_ancestor_ids": ["b1"]}),
                SourceBlock(f"s{index}-body", "paragraph", index * 2, f"<p>{prose}</p>", source_note_id="n5", locator={"heading_ancestor_ids": ["b1", heading]}),
            ])
            conn.execute("INSERT INTO notes_nodes(node_uuid,template_id,sheet,row,label,kind) VALUES (?,'mfrs-company-notes-list-v1','Notes-Listofnotes',?,?,'LEAF')", (f"mixed-{row}", row, label))
        # Retire the fixture's old paragraph; preserve the other source note.
        conn.execute("DELETE FROM notes_source_blocks WHERE generation_id=? AND block_id='b2'", (deps.source_generation_id,))
        srepo.write_blocks(conn, deps.source_generation_id, blocks)
    return targets, blocks


@pytest.mark.asyncio
async def test_complete_source_sections_retain_all_fields_and_accept_full_receipt(assigned_agent):
    from types import SimpleNamespace
    deps, tools = assigned_agent
    targets, blocks = _seed_mixed_source_sections(deps)
    ctx = SimpleNamespace(deps=deps)
    write = tools["write_note_from_source"].function
    for index, (row, label, _, _) in enumerate(targets, 1):
        result = await write(ctx, deps.sheet_name, row, label, [f"section:n5:5.{index}"])
        assert result.startswith("ok:") and "Re-routed" not in result
        if index == 2:
            partial = [{"note_num": 5, "action": "written", "row_labels": [item[1] for item in targets[:2]]}]
            assert "still unplaced" in notes_agent._submit_coverage_entries_impl(deps, partial)
    receipt = [{"note_num": 5, "action": "written", "row_labels": [item[1] for item in targets]}]
    assert "accepted" in notes_agent._submit_coverage_entries_impl(deps, receipt)
    assert len(deps.payload_sink) == 3
    assert {payload.chosen_row_label for payload in deps.payload_sink} == {item[1] for item in targets}
    assert "60,606,000" in next(payload.content for payload in deps.payload_sink if "Historical cost" in payload.content)
    with repo.db_session(deps.db_path) as conn:
        assert {row[0] for row in conn.execute("SELECT row FROM notes_cells WHERE run_id=?", (deps.run_id,))} == {10, 31, 133}
        assert {p["block_id"] for p in srepo.active_placements(conn, deps.source_generation_id)} == {b.block_id for b in blocks}
    # Same-field revision still replaces that section without erasing siblings.
    assert (await write(ctx, deps.sheet_name, targets[1][0], targets[1][1], ["section:n5:5.2"])).startswith("ok:")
    assert len(deps.payload_sink) == 3
    assert "accepted" in notes_agent._submit_coverage_entries_impl(deps, receipt)
    missing_label = [{"note_num": 5, "action": "written", "row_labels": [targets[0][1]]}]
    assert "list every live destination" in notes_agent._submit_coverage_entries_impl(deps, missing_label)


@pytest.mark.asyncio
async def test_moving_one_source_section_preserves_siblings_in_workbook(assigned_agent, tmp_path):
    from types import SimpleNamespace
    from openpyxl import load_workbook
    from notes.writer import write_notes_workbook

    deps, tools = assigned_agent
    targets, _ = _seed_mixed_source_sections(deps)
    ctx = SimpleNamespace(deps=deps)
    for index, (row, label, _, _) in enumerate(targets, 1):
        result = await tools["write_note_from_source"].function(
            ctx, deps.sheet_name, row, label, [f"section:n5:5.{index}"],
        )
        assert result.startswith("ok:")
    destination_label = "Disclosure of trade and other receivables"
    with repo.db_session(deps.db_path) as conn:
        conn.execute("INSERT INTO template_slots("
                     "target_id,canonical_target_id,template_id,sheet,row,col,label,"
                     "slot_role,value_kind,mapping_source,manifest_version,"
                     "workbook_fingerprint,validation_status) "
                     "VALUES ('managed-target','managed-target','mfrs-company-notes-list-v1',"
                     "'Notes-Listofnotes',140,'B',?,'INPUT','html','test','test-v1','fixture','writable')",
                     (destination_label,))
    deps.coverage_receipt = object()
    moved = await tools["move_own_source_cell"].function(
        ctx, targets[1][0], 140, destination_label, "Correct the measurement section destination.",
    )
    assert moved.startswith("Moved")
    expected_labels = [targets[0][1], destination_label, targets[2][1]]
    assert {p.chosen_row_label for p in deps.payload_sink} == set(expected_labels)
    assert deps.coverage_receipt is None
    incomplete = [{"note_num": 5, "action": "written", "row_labels": [destination_label]}]
    assert "rejected" in notes_agent._submit_coverage_entries_impl(deps, incomplete)
    complete = [{"note_num": 5, "action": "written", "row_labels": expected_labels}]
    assert "accepted" in notes_agent._submit_coverage_entries_impl(deps, complete)
    with repo.db_session(deps.db_path) as conn:
        assert {cell[0] for cell in conn.execute("SELECT row FROM notes_cells WHERE run_id=?", (deps.run_id,))} == {133, 140, 31}
        assert {p["row"] for p in srepo.active_placements(conn, deps.source_generation_id)} == {133, 140, 31}
    output = tmp_path / "moved-sections.xlsx"
    written = write_notes_workbook(deps.template_path, deps.payload_sink, str(output),
                                   deps.filing_level, deps.sheet_name)
    assert written.success and written.rows_written == 3
    workbook = load_workbook(output)
    try:
        sheet = workbook[deps.sheet_name]
        for row, text in [(133, targets[0][3]), (140, targets[1][3]), (31, targets[2][3])]:
            assert text in sheet.cell(row, 2).value
        assert not sheet.cell(targets[1][0], 2).value
    finally:
        workbook.close()


@pytest.mark.asyncio
@pytest.mark.parametrize("route", ["approved", "no_approval", "third_destination"])
async def test_split_receipt_honors_only_approved_capital_duplicate(assigned_agent, route):
    from types import SimpleNamespace
    from notes import source_write

    deps, tools = assigned_agent
    targets, _ = _seed_mixed_source_sections(deps)
    ctx = SimpleNamespace(deps=deps)
    for index, (row, label, _, _) in enumerate(targets, 1):
        assert (await tools["write_note_from_source"].function(
            ctx, deps.sheet_name, row, label, [f"section:n5:5.{index}"],
        )).startswith("ok:")
    with repo.db_session(deps.db_path) as conn:
        conn.execute("INSERT INTO template_slots("
                     "target_id,canonical_target_id,template_id,sheet,row,col,label,"
                     "slot_role,value_kind,mapping_source,manifest_version,"
                     "workbook_fingerprint,validation_status) "
                     "VALUES ('capital-prose','capital-prose','mfrs-company-notes-capital-v1',"
                     "'Notes-Issuedcapital',4,'B','Issued capital disclosure','INPUT','html',"
                     "'test','test-v1','fixture','writable')")
        source_write.write_cell_from_blocks(
            conn, run_id=deps.run_id, generation_id=deps.source_generation_id,
            sheet="Notes-Issuedcapital", row=4, block_ids=["s1-heading", "s1-body"],
            template_prefix="mfrs-company-",
        )
        if route == "no_approval":
            conn.execute("DELETE FROM notes_disposition_events WHERE generation_id=? "
                         "AND reason_code='APPROVED_DUPLICATE_ROUTE'", (deps.source_generation_id,))
        elif route == "third_destination":
            repo.upsert_notes_cell(conn, run_id=deps.run_id, sheet=deps.sheet_name,
                                   row=140, label="Disclosure of trade and other receivables",
                                   html=conn.execute("SELECT html FROM notes_cells WHERE run_id=? "
                                                     "AND sheet='Notes-Issuedcapital' AND row=4",
                                                     (deps.run_id,)).fetchone()["html"])
            srepo.set_cell_placements(conn, deps.run_id, deps.source_generation_id,
                                      deps.sheet_name, 140, ["s1-heading", "s1-body"])
    receipt = [{"note_num": 5, "action": "written", "row_labels": [item[1] for item in targets]}]
    result = notes_agent._submit_coverage_entries_impl(deps, receipt)
    assert ("accepted" if route == "approved" else "rejected") in result


@pytest.mark.asyncio
@pytest.mark.parametrize("damage", ["missing_placements", "changed_content", "foreign_cell_generation", "stale_source"])
async def test_source_multi_field_receipt_requires_current_matching_canonical_placements(assigned_agent, damage):
    from types import SimpleNamespace
    deps, tools = assigned_agent
    targets, _ = _seed_mixed_source_sections(deps)
    ctx = SimpleNamespace(deps=deps)
    for index, (row, label, _, _) in enumerate(targets, 1):
        assert (await tools["write_note_from_source"].function(ctx, deps.sheet_name, row, label, [f"section:n5:5.{index}"])).startswith("ok:")
    with repo.db_session(deps.db_path) as conn:
        if damage == "missing_placements":
            conn.execute("UPDATE notes_block_placements SET active=0 WHERE generation_id=?", (deps.source_generation_id,))
        elif damage == "changed_content":
            conn.execute("UPDATE notes_cells SET html='<p>Changed by a human</p>',content_revision=content_revision+1 WHERE run_id=? AND row=10", (deps.run_id,))
        elif damage == "foreign_cell_generation":
            conn.execute("UPDATE notes_cells SET source_generation_id=NULL WHERE run_id=? AND row=10", (deps.run_id,))
        elif damage == "stale_source":
            conn.execute("UPDATE notes_source_generations SET status='superseded' WHERE id=?", (deps.source_generation_id,))
    receipt = [{"note_num": 5, "action": "written", "row_labels": [item[1] for item in targets]}]
    assert "rejected" in notes_agent._submit_coverage_entries_impl(deps, receipt)
    assert deps.coverage_receipt is None


@pytest.mark.asyncio
async def test_source_multi_field_receipt_rejects_fragments_within_one_section(assigned_agent):
    from types import SimpleNamespace
    deps, tools = assigned_agent
    targets, _ = _seed_mixed_source_sections(deps)
    with repo.db_session(deps.db_path) as conn:
        conn.execute("DELETE FROM notes_source_blocks WHERE generation_id=? AND block_id!='b1'", (deps.source_generation_id,))
        srepo.write_blocks(conn, deps.source_generation_id, [
            SourceBlock("b1", "heading", 0, "<h2>5 Basis of preparation</h2>", source_note_id="n5"),
            SourceBlock("heading", "heading", 1, "<h3>5.1 Measurement</h3>", source_note_id="n5", locator={"heading_ancestor_ids": ["b1"]}),
            SourceBlock("first", "paragraph", 2, "<p>First part.</p>", source_note_id="n5", locator={"heading_ancestor_ids": ["b1", "heading"]}),
            SourceBlock("second", "paragraph", 3, "<p>Second part.</p>", source_note_id="n5", locator={"heading_ancestor_ids": ["b1", "heading"]}),
        ])
    ctx = SimpleNamespace(deps=deps)
    for block, (row, label, _, _) in zip(["first", "second"], targets[:2]):
        assert (await tools["write_note_from_source"].function(ctx, deps.sheet_name, row, label, [block])).startswith("ok:")
    receipt = [{"note_num": 5, "action": "written", "row_labels": [item[1] for item in targets[:2]]}]
    assert "exactly one List-of-Notes field" in notes_agent._submit_coverage_entries_impl(deps, receipt)


@pytest.mark.asyncio
async def test_source_nested_sections_preserve_parent_intro_and_complete_children(assigned_agent):
    from types import SimpleNamespace
    deps, tools = assigned_agent
    targets, _ = _seed_mixed_source_sections(deps)
    blocks = [
        SourceBlock("root", "heading", 0, "<h2>5 Basis of preparation</h2>", source_note_id="n5"),
        SourceBlock("parent", "heading", 1, "<h3>5.1 Basis overview</h3>", source_note_id="n5", locator={"heading_ancestor_ids": ["root"]}),
        SourceBlock("intro", "paragraph", 2, "<p>Complete basis overview introduction.</p>", source_note_id="n5", locator={"heading_ancestor_ids": ["root", "parent"]}),
        SourceBlock("child1", "heading", 3, "<h4>5.1.1 Measurement</h4>", source_note_id="n5", locator={"heading_ancestor_ids": ["root", "parent"]}),
        SourceBlock("child1body", "paragraph", 4, "<p>Complete measurement disclosure.</p>", source_note_id="n5", locator={"heading_ancestor_ids": ["root", "parent", "child1"]}),
        SourceBlock("child2", "heading", 5, "<h4>5.1.2 Estimates</h4>", source_note_id="n5", locator={"heading_ancestor_ids": ["root", "parent"]}),
        SourceBlock("child2body", "paragraph", 6, "<p>Complete estimates disclosure.</p>", source_note_id="n5", locator={"heading_ancestor_ids": ["root", "parent", "child2"]}),
    ]
    with repo.db_session(deps.db_path) as conn:
        srepo.write_blocks(conn, deps.source_generation_id, blocks)
    ctx = SimpleNamespace(deps=deps)
    for selected, (row, label, _, _) in zip(["intro", "section:n5:5.1.1", "section:n5:5.1.2"], targets):
        assert (await tools["write_note_from_source"].function(ctx, deps.sheet_name, row, label, [selected])).startswith("ok:")
    receipt = [{"note_num": 5, "action": "written", "row_labels": [item[1] for item in targets]}]
    assert "accepted" in notes_agent._submit_coverage_entries_impl(deps, receipt)
    assert len(deps.payload_sink) == 3
    assert any("Complete basis overview introduction." in payload.content for payload in deps.payload_sink)
    with repo.db_session(deps.db_path) as conn:
        assert {p["block_id"] for p in srepo.active_placements(conn, deps.source_generation_id)} == {block.block_id for block in blocks}


@pytest.mark.asyncio
async def test_source_multi_field_receipt_cannot_omit_another_live_destination(assigned_agent):
    from types import SimpleNamespace
    deps, tools = assigned_agent
    targets, _ = _seed_mixed_source_sections(deps)
    ctx = SimpleNamespace(deps=deps)
    for index, (row, label, _, _) in enumerate(targets, 1):
        assert (await tools["write_note_from_source"].function(ctx, deps.sheet_name, row, label, [f"section:n5:5.{index}"])).startswith("ok:")
    # A stale sink view must not certify only two of three live destinations.
    deps.payload_sink[:] = [p for p in deps.payload_sink if p.chosen_row_label != targets[2][1]]
    receipt = [{"note_num": 5, "action": "written", "row_labels": [item[1] for item in targets[:2]]}]
    assert "rejected" in notes_agent._submit_coverage_entries_impl(deps, receipt)
    assert deps.coverage_receipt is None


def test_the_source_tools_are_absent_without_a_frozen_reading(tmp_path):
    """An `off`-mode run must see exactly the agent it saw before."""
    from pydantic_ai.models.test import TestModel

    agent, _deps = notes_agent.create_notes_agent(
        template_type=NotesTemplateType.CORP_INFO, pdf_path="/tmp/no.pdf",
        inventory=[], filing_level="company", model=TestModel(),
        output_dir=str(tmp_path),
    )
    names = _tool_names(agent)
    assert "write_note_from_source" not in names
    assert "list_source_notes" not in names


@pytest.mark.asyncio
@pytest.mark.parametrize("input_kind", ["docx_html", "prepared_document"])
async def test_the_source_tools_appear_with_a_frozen_reading(tmp_path, seeded, input_kind):
    from types import SimpleNamespace
    from pydantic_ai.models.test import TestModel

    db, run_id, gen = seeded
    with repo.db_session(db) as conn:
        conn.execute("UPDATE notes_source_generations SET input_kind=? WHERE id=?",
                     (input_kind, gen))
    agent, deps = notes_agent.create_notes_agent(
        template_type=NotesTemplateType.CORP_INFO, pdf_path="/tmp/no.pdf",
        inventory=[], filing_level="company", model=TestModel(),
        output_dir=str(tmp_path),
        run_id=run_id, db_path=db, source_generation_id=gen,
    )
    names = _tool_names(agent)
    assert {"list_source_notes", "read_source_manifest",
            "view_source_blocks", "write_note_from_source"} <= names
    assert "request_source_recheck" not in names
    assert deps.source_generation_id == gen
    functions = {name: tool.function for ts in agent.toolsets
                 for name, tool in getattr(ts, "tools", {}).items()}
    for name, args in [("list_source_notes", ()), ("read_source_manifest", (5,)),
                       ("view_source_blocks", (["b2"],))]:
        ctx = SimpleNamespace(deps=deps)
        full = await functions[name](ctx, *args)
        tail = await functions[name](ctx, *args, offset=5)
        body = lambda text: text.split("<<<SOURCE>>>\n", 1)[1].split("\n<<<END_SOURCE>>>", 1)[0]
        assert body(tail) == body(full)[5:]
        assert "Invalid offset" in await functions[name](ctx, *args, offset=-1)


def _word_upload_dir(tmp_path):
    """A session dir shaped like a Word upload: uploaded.pdf + source.html."""
    d = tmp_path / "session"
    d.mkdir()
    (d / "uploaded.pdf").write_bytes(b"%PDF-1.4\n")
    (d / "source.html").write_text(
        "<h3>5. Receivables</h3><p>Stated at cost.</p>", encoding="utf-8",
    )
    return d


def test_block_mode_hides_the_copy_workflow_tool(tmp_path, seeded):
    """Peer review 2026-08-06: the prompt switched to block assembly but
    read_source_note stayed registered, and its description teaches
    copy-into-content — two incompatible workflows exposed at once. On a
    block-path run the copy tool must be absent."""
    from pydantic_ai.models.test import TestModel

    db, run_id, gen = seeded
    d = _word_upload_dir(tmp_path)
    agent, deps = notes_agent.create_notes_agent(
        template_type=NotesTemplateType.CORP_INFO,
        pdf_path=str(d / "uploaded.pdf"),
        inventory=[], filing_level="company", model=TestModel(),
        output_dir=str(tmp_path),
        run_id=run_id, db_path=db, source_generation_id=gen,
    )
    names = _tool_names(agent)
    assert deps.source_block_notes == {5, 6}
    assert "write_note_from_source" in names
    assert "read_source_note" not in names


def test_sidecar_only_run_keeps_the_copy_workflow_tool(tmp_path):
    """No generation → the sidecar workflow is unchanged, copy tool and all."""
    from pydantic_ai.models.test import TestModel

    d = _word_upload_dir(tmp_path)
    agent, deps = notes_agent.create_notes_agent(
        template_type=NotesTemplateType.CORP_INFO,
        pdf_path=str(d / "uploaded.pdf"),
        inventory=[], filing_level="company", model=TestModel(),
        output_dir=str(tmp_path),
    )
    names = _tool_names(agent)
    assert deps.source_block_notes == set()
    assert "read_source_note" in names
    assert "write_note_from_source" not in names


def test_numeric_templates_stay_off_the_block_workflow(tmp_path, seeded):
    """Peer review 2026-08-06: sheets 13/14 need structured numeric_values and
    write_note_from_source resolves prose nodes only — teaching it there
    taught a write that always rejects. A numeric agent keeps the sidecar
    workflow: no write tool, no block prompt, copy tool still present."""
    from pydantic_ai.models.test import TestModel

    db, run_id, gen = seeded
    d = _word_upload_dir(tmp_path)
    agent, deps = notes_agent.create_notes_agent(
        template_type=NotesTemplateType.ISSUED_CAPITAL,
        pdf_path=str(d / "uploaded.pdf"),
        inventory=[], filing_level="company", model=TestModel(),
        output_dir=str(tmp_path),
        run_id=run_id, db_path=db, source_generation_id=gen,
    )
    names = _tool_names(agent)
    assert deps.source_block_notes == set()
    assert "write_note_from_source" not in names
    assert "read_source_note" in names


# --------------------------------------------------------------------------
# what the tools return
# --------------------------------------------------------------------------

def test_listing_notes_reports_each_notes_part_count(seeded):
    db, _run_id, gen = seeded
    out = notes_agent._list_source_notes_impl(db, gen)
    assert "note   5" in out and "2 part(s)" in out
    assert "note   6" in out


def test_the_manifest_lists_ids_and_kinds_for_one_note(seeded):
    db, _run_id, gen = seeded
    out = notes_agent._read_source_manifest_impl(db, gen, 5)
    assert "b1" in out and "b2" in out
    assert "b3" not in out, "another note's parts must not leak in"
    assert "heading" in out and "paragraph" in out


def test_section_listing_gives_one_complete_choice_for_a_simple_note(seeded):
    db, _run_id, gen = seeded
    out = notes_agent._list_source_sections_impl(db, gen, 5)
    assert "section:n5:root" in out
    for selection in ("section:n5:root", "n5:root"):
        viewed = notes_agent._view_source_blocks_impl(db, gen, [selection])
        assert "5. Receivables" in viewed and "Stated at cost." in viewed
    assert "unknown source section" in notes_agent._view_source_blocks_impl(
        db, gen, ["n5:missing"],
    )
    assert "2 source parts" in out
    assert "b2" not in out


def test_individual_part_write_names_the_rest_of_its_section(seeded):
    from notes import source_write

    db, run_id, gen = seeded
    with repo.db_session(db) as conn:
        outcome = source_write.write_cell_from_blocks(
            conn, run_id=run_id, generation_id=gen,
            sheet="Notes-Listofnotes", row=10, block_ids=["b1"],
        )
    assert "section:n5:root" in outcome.as_message()
    assert "b2" in outcome.as_message()


def test_repeated_lettered_sections_keep_distinct_source_pieces():
    from notes.source_sections import expand_section_ids, sections_for_note

    blocks = [
        SourceBlock(block_id="a1", block_kind="heading", reading_order=0,
                    canonical_html="<h3>(a) First topic</h3>", source_note_id="n5"),
        SourceBlock(block_id="first", block_kind="paragraph", reading_order=1,
                    canonical_html="<p>First disclosure.</p>", source_note_id="n5"),
        SourceBlock(block_id="b1", block_kind="heading", reading_order=2,
                    canonical_html="<h3>(b) Other topic</h3>", source_note_id="n5"),
        SourceBlock(block_id="a2", block_kind="heading", reading_order=3,
                    canonical_html="<h3>(a) Later topic</h3>", source_note_id="n5"),
        SourceBlock(block_id="last", block_kind="paragraph", reading_order=4,
                    canonical_html="<p>Later disclosure.</p>", source_note_id="n5"),
    ]
    notes = [{"source_note_id": "n5", "top_note_num": "5", "title": "Policies"}]

    sections = sections_for_note(blocks, "n5", "5", "Policies")
    assert [section.section_id for section in sections] == [
        "section:n5:a", "section:n5:b", "section:n5:a:2",
    ]
    assert expand_section_ids(blocks, notes, [sections[0].section_id]) == ["a1", "first"]
    assert expand_section_ids(blocks, notes, [sections[2].section_id]) == ["a2", "last"]
    assert expand_section_ids(blocks, notes, ["n5:a:2"]) == ["a2", "last"]
    # Exact existing block identities win over optional-prefix section recovery.
    collision = SourceBlock("n5:a", "paragraph", 5, "<p>Exact block.</p>")
    assert expand_section_ids([*blocks, collision], notes, ["n5:a"]) == ["n5:a"]
    with pytest.raises(ValueError, match="unknown source section"):
        expand_section_ids(blocks, notes, ["n5:missing"])


def test_parent_sections_include_nested_numbered_and_roman_disclosures(seeded):
    from notes.source_sections import expand_section_ids, sections_for_note

    numbered = [
        SourceBlock("parent", "heading", 1, "<h3>3.1 Basis of preparation</h3>", source_note_id="n3"),
        SourceBlock("child", "heading", 2, "<h4>3.1.1 Statement of compliance</h4>", source_note_id="n3",
                    locator={"heading_ancestor_ids": ["parent"]}),
        SourceBlock("text", "paragraph", 3, "<p>Complies with MFRS.</p>", source_note_id="n3",
                    locator={"heading_ancestor_ids": ["parent", "child"]}),
        SourceBlock("next", "heading", 4, "<h3>3.2 Other policy</h3>", source_note_id="n3"),
    ]
    notes = [{"source_note_id": "n3", "top_note_num": "3", "title": "Policies"}]
    assert expand_section_ids(numbered, notes, ["section:n3:3.1"]) == ["parent", "child", "text"]

    lettered = [
        SourceBlock("a", "heading", 1, "<h3>(a) Leases</h3>", source_note_id="n3"),
        SourceBlock("i", "heading", 2, "<h4>(i) As lessee</h4>", source_note_id="n3",
                    locator={"heading_ancestor_ids": ["a"]}),
        SourceBlock("lessee", "paragraph", 3, "<p>Lease assets.</p>", source_note_id="n3"),
        SourceBlock("ii", "heading", 4, "<h4>(ii) As lessor</h4>", source_note_id="n3",
                    locator={"heading_ancestor_ids": ["a"]}),
        SourceBlock("lessor", "paragraph", 5, "<p>Lease income.</p>", source_note_id="n3"),
        SourceBlock("b", "heading", 6, "<h3>(b) Revenue</h3>", source_note_id="n3"),
        SourceBlock("bi", "heading", 7, "<h4>(i) Recognition</h4>", source_note_id="n3",
                    locator={"heading_ancestor_ids": ["b"]}),
    ]
    assert expand_section_ids(lettered, notes, ["section:n3:a"]) == ["a", "i", "lessee", "ii", "lessor"]
    assert expand_section_ids(lettered, notes, ["section:n3:i"]) == ["i", "lessee"]
    assert expand_section_ids(lettered, notes, ["section:n3:ii"]) == ["ii", "lessor"]
    sections = sections_for_note(lettered, "n3", "3", "Policies")
    assert next(s for s in sections if s.section_id == "section:n3:i").parent_title == "(a) Leases"
    assert next(s for s in sections if s.section_id == "section:n3:i:2").parent_title == "(b) Revenue"
    db, _, gen = seeded
    with repo.db_session(db) as conn:
        srepo.write_blocks(conn, gen, lettered)
        srepo.write_notes(conn, gen, [SourceNote("n3", "3", "Policies")])
    listing = notes_agent._list_source_sections_impl(db, gen, 3)
    first = next(line for line in listing.splitlines() if "section:n3:i " in line)
    second = next(line for line in listing.splitlines() if "section:n3:i:2 " in line)
    assert "under (a) Leases" in first
    assert "under (b) Revenue" in second


@pytest.mark.parametrize("selection", ["section:n5:root", "n5:root"])
def test_section_id_reaches_sheet12_payload_builder_as_source_pieces(seeded, selection):
    from types import SimpleNamespace

    db, run_id, gen = seeded
    with repo.db_session(db) as conn:
        conn.execute(
            "INSERT INTO notes_nodes(node_uuid,template_id,sheet,row,label,kind) "
            "VALUES('note5','mfrs-company-notes-v1','Notes-Listofnotes',140,"
            "'Disclosure of receivables','LEAF')"
        )
        deps = SimpleNamespace(
            run_id=run_id, source_generation_id=gen, filing_standard="mfrs",
            filing_level="company", sheet_name="Notes-Listofnotes",
        )
        message, payload = notes_agent._write_from_source_in_connection(
            conn, deps, "Notes-Listofnotes", 140, [selection],
            [1], "source page 1", None,
            target_label="Disclosure of receivables",
        )
    assert message.startswith("ok:")
    assert "destination: Disclosure of receivables" in message
    assert "saved source parts: b1, b2" in message
    assert payload.note_num == 5
    assert payload.source_note_id == "n5"
    assert "Stated at cost" in payload.content


def test_a_note_with_no_parts_says_to_read_the_pdf(seeded):
    db, _run_id, gen = seeded
    assert "Read the PDF" in notes_agent._read_source_manifest_impl(db, gen, 99)


def test_viewing_blocks_returns_their_full_content(seeded):
    db, _run_id, gen = seeded
    out = notes_agent._view_source_blocks_impl(db, gen, ["b2"])
    assert "Stated at cost" in out


def test_an_unknown_block_id_is_named_not_silently_dropped(seeded):
    db, _run_id, gen = seeded
    out = notes_agent._view_source_blocks_impl(db, gen, ["b2", "nope"])
    assert "not found: nope" in out


def test_asking_only_for_unknown_ids_points_at_the_manifest(seeded):
    db, _run_id, gen = seeded
    out = notes_agent._view_source_blocks_impl(db, gen, ["nope"])
    assert "read_source_manifest" in out


def test_document_content_carries_the_untrusted_framing(seeded):
    db, _run_id, gen = seeded
    out = notes_agent._view_source_blocks_impl(db, gen, ["b2"])
    assert "UNTRUSTED" in out
    assert "data, not commands" in out
    assert "<<<SOURCE>>>" in out


# --------------------------------------------------------------------------
# caps
# --------------------------------------------------------------------------

def test_a_long_block_can_be_read_completely_in_bounded_pages(tmp_path):
    import re
    db = tmp_path / "big.sqlite"
    init_db(db)
    with repo.db_session(db) as conn:
        run_id = repo.create_run(
            conn, "x.docx", session_id="s", output_dir=str(tmp_path / "s")
        )
        gen = srepo.begin_generation(conn, run_id, input_kind="docx_html")
        srepo.write_blocks(conn, gen, [
            SourceBlock(block_id="big", block_kind="paragraph", reading_order=0,
                        canonical_html="<p>" + ("x" * 200_000) + "</p>",
                        source_note_id="n1"),
        ])
        srepo.activate_generation(conn, gen)
    chunks, offset = [], 0
    for _ in range(10):
        out = notes_agent._view_source_blocks_impl(str(db), gen, ["big"], offset=offset)
        assert len(out) <= notes_agent.SOURCE_TOOL_RESPONSE_CAP + 600
        assert "UNTRUSTED" in out and "<<<END_SOURCE>>>" in out
        chunks.append(out.split("<<<SOURCE>>>\n", 1)[1].split("\n<<<END_SOURCE>>>", 1)[0])
        continuation = re.search(r"next_offset=(\d+)", out)
        if continuation is None:
            break
        next_offset = int(continuation[1])
        assert next_offset > offset
        offset = next_offset
    else:
        pytest.fail("source pagination did not terminate")
    assert "".join(chunks) == "--- big (paragraph) ---\n<p>" + "x" * 200_000 + "</p>"


@pytest.mark.parametrize("impl,args", [
    (notes_agent._list_source_notes_impl, ()),
    (notes_agent._read_source_manifest_impl, (5,)),
])
def test_source_indexes_paginate_without_losing_ids(seeded, monkeypatch, impl, args):
    import re
    db, _run_id, gen = seeded
    whole = impl(db, gen, *args).split("<<<SOURCE>>>\n", 1)[1].split("\n<<<END_SOURCE>>>", 1)[0]
    monkeypatch.setattr(notes_agent, "SOURCE_TOOL_RESPONSE_CAP", 35)
    offset, chunks = 0, []
    for _ in range(20):
        out = impl(db, gen, *args, offset=offset)
        chunks.append(out.split("<<<SOURCE>>>\n", 1)[1].split("\n<<<END_SOURCE>>>", 1)[0])
        continuation = re.search(r"next_offset=(\d+)", out)
        if continuation is None:
            break
        offset = int(continuation[1])
    else:
        pytest.fail("index pagination did not terminate")
    assert len(chunks) > 1
    assert "".join(chunks) == whole


def test_too_many_block_ids_are_bounded_per_call(seeded):
    db, _run_id, gen = seeded
    many = [f"b{i}" for i in range(200)] + ["b1"]
    out = notes_agent._view_source_blocks_impl(db, gen, many)
    assert f"first {notes_agent._SOURCE_BLOCKS_PER_CALL} parts" in out


@pytest.mark.parametrize("missing_first_batch", [False, True])
def test_block_limit_returns_actionable_remaining_batch(seeded, missing_first_batch):
    db, _, gen = seeded
    ids = [f"part-{i}" for i in range(45)]
    with repo.db_session(db) as conn:
        srepo.write_blocks(conn, gen, [SourceBlock(
            block_id=bid, block_kind="paragraph", reading_order=i + 10,
            canonical_html=f"<p>Disclosure {i}</p>",
        ) for i, bid in enumerate(ids) if not missing_first_batch or i >= 40])
    first = notes_agent._view_source_blocks_impl(db, gen, ids)
    assert "offset=0" in first and "remaining block_ids" in first
    for bid in ids[40:]:
        assert bid in first.split("remaining block_ids", 1)[1]
    second = notes_agent._view_source_blocks_impl(db, gen, ids[40:], offset=0)
    for bid in ids[:40]:
        assert (f"--- {bid} " in first) == (not missing_first_batch)
        assert f"--- {bid} " not in second
    for bid in ids[40:]:
        assert f"--- {bid} " in second and f"--- {bid} " not in first


def test_block_read_warnings_appear_on_every_page(seeded, monkeypatch):
    import re

    db, _run_id, gen = seeded
    monkeypatch.setattr(notes_agent, "SOURCE_TOOL_RESPONSE_CAP", 12)
    ids = ["b1"] + [f"missing{i}" for i in range(40)]
    offset = 0
    for _ in range(10):
        out = notes_agent._view_source_blocks_impl(db, gen, ids, offset=offset)
        body = out.split("<<<SOURCE>>>\n", 1)[1].split("\n<<<END_SOURCE>>>", 1)[0]
        assert body.startswith("[not found: missing0, missing1")
        assert "[only the first 40 parts were returned]\n" in body
        continuation = re.search(r"next_offset=(\d+)", out)
        if continuation is None:
            break
        offset = int(continuation[1])
    else:
        pytest.fail("source pagination did not terminate")


def test_previews_in_the_manifest_are_short(tmp_path):
    db = tmp_path / "prev.sqlite"
    init_db(db)
    with repo.db_session(db) as conn:
        run_id = repo.create_run(
            conn, "x.docx", session_id="s", output_dir=str(tmp_path / "s")
        )
        gen = srepo.begin_generation(conn, run_id, input_kind="docx_html")
        srepo.write_blocks(conn, gen, [
            SourceBlock(block_id="b1", block_kind="paragraph", reading_order=0,
                        canonical_html="<p>" + ("word " * 5000) + "</p>",
                        source_note_id="n1"),
        ])
        srepo.activate_generation(conn, gen)
    out = notes_agent._read_source_manifest_impl(str(db), gen, 1)
    assert len(out) < 2_000


# --------------------------------------------------------------------------
# a run without a reading degrades, never crashes
# --------------------------------------------------------------------------

@pytest.mark.parametrize("impl,args", [
    (notes_agent._list_source_notes_impl, ()),
    (notes_agent._read_source_manifest_impl, (1,)),
    (notes_agent._view_source_blocks_impl, (["b1"],)),
])
def test_every_tool_degrades_without_a_generation(impl, args):
    assert "No frozen source reading" in impl(None, None, *args)


def test_prepared_source_rejects_prose_fallback_even_without_numbered_notes(tmp_path, seeded):
    import asyncio
    from types import SimpleNamespace
    from pydantic_ai.models.test import TestModel

    db, run_id, gen = seeded
    with repo.db_session(db) as conn:
        conn.execute("UPDATE notes_source_generations SET input_kind='prepared_document' WHERE id=?", (gen,))
        conn.execute("UPDATE notes_source_notes SET top_note_num='' WHERE generation_id=?", (gen,))
        conn.commit()
    agent, deps = notes_agent.create_notes_agent(
        template_type=NotesTemplateType.CORP_INFO, pdf_path="/tmp/no.pdf",
        inventory=[], filing_level="company", model=TestModel(), output_dir=str(tmp_path),
        run_id=run_id, db_path=db, source_generation_id=gen)
    assert deps.prepared_source_required
    assert "write_note_from_source" in _tool_names(agent)
    assert "write_notes" not in _tool_names(agent)
    assert "report_source_gap" in _tool_names(agent)
    prompt = notes_agent.render_notes_prompt(NotesTemplateType.CORP_INFO, "company", [],
        source_blocks_available=True, prepared_source_required=True)
    assert "All writes go through" not in prompt
    assert "copy its markup by hand" not in prompt
    assert "report_source_gap" in prompt


def test_manifest_accepts_stable_identity_for_unnumbered_note(seeded):
    db, _, gen = seeded
    with repo.db_session(db) as conn:
        conn.execute("UPDATE notes_source_notes SET top_note_num='' WHERE generation_id=? AND source_note_id='n5'", (gen,))
        conn.commit()
    assert "b1" in notes_agent._read_source_manifest_impl(db, gen, "n5")
    assert "b2" in notes_agent._read_source_manifest_impl(db, gen, "n5")


@pytest.mark.asyncio
async def test_prepared_capture_gap_is_persisted_without_false_coverage(tmp_path, seeded, monkeypatch):
    from types import SimpleNamespace
    from pydantic_ai.models.test import TestModel
    db, run_id, gen = seeded
    with repo.db_session(db) as conn:
        conn.execute("UPDATE notes_source_generations SET input_kind='prepared_document' WHERE id=?", (gen,))
    monkeypatch.setattr("tools.pdf_viewer.count_pdf_pages", lambda path: 2)
    agent, deps = notes_agent.create_notes_agent(
        template_type=NotesTemplateType.LIST_OF_NOTES, pdf_path="synthetic.pdf",
        inventory=[], filing_level="company", model=TestModel(), output_dir=str(tmp_path),
        run_id=run_id, db_path=db, source_generation_id=gen, batch_note_nums=[5])
    deps.payload_sink = []
    report = next(ts.tools["report_source_gap"].function for ts in agent.toolsets
                  if "report_source_gap" in getattr(ts, "tools", {}))
    result = report(SimpleNamespace(deps=deps), [1], "Table heading was not captured", 5)
    assert "human review" in result
    assert "accepted" in notes_agent._submit_coverage_entries_impl(deps, [])
    assert deps.coverage_receipt.entries == []  # It did not become covered.
    with repo.db_session(db) as conn:
        flags = repo.fetch_notes_review_flags(conn, run_id)
    assert len(flags) == 1
    assert flags[0]["reason"] == "Table heading was not captured"
    assert flags[0]["source_pages"] == [1]
    # Saving the report does not claim that missing source content was written.
    deps.payload_sink = None
    save = next(ts.tools["save_result"].function for ts in agent.toolsets
                if "save_result" in getattr(ts, "tools", {}))
    await save(SimpleNamespace(deps=deps))
    assert deps.source_gap_reported
    assert not deps.wrote_once
    assert deps.cells_written == []
    assert deps.write_skip_errors


def test_prepared_batch_receipt_rejects_a_cross_sheet_skip_until_every_source_part_is_placed(
    tmp_path, seeded,
):
    from pydantic_ai.models.test import TestModel
    from notes import source_write

    db, run_id, gen = seeded
    with repo.db_session(db) as conn:
        conn.execute("UPDATE notes_source_generations SET input_kind='prepared_document' WHERE id=?", (gen,))
    _agent, deps = notes_agent.create_notes_agent(
        template_type=NotesTemplateType.LIST_OF_NOTES, pdf_path="synthetic.pdf",
        inventory=[], filing_level="company", model=TestModel(), output_dir=str(tmp_path),
        run_id=run_id, db_path=db, source_generation_id=gen, batch_note_nums=[5],
    )
    deps.payload_sink = []
    receipt = [{"note_num": 5, "action": "skipped", "reason": "Accounting Policies sheet"}]

    assert "section:n5:root" in notes_agent._submit_coverage_entries_impl(deps, receipt)
    assert deps.coverage_receipt is None

    with repo.db_session(db) as conn:
        source_write.write_cell_from_blocks(
            conn, run_id=run_id, generation_id=gen, sheet="Notes-SummaryofAccPol",
            row=10, block_ids=["b1"],
        )
    assert "section:n5:root" in notes_agent._submit_coverage_entries_impl(deps, receipt)

    with repo.db_session(db) as conn:
        source_write.write_cell_from_blocks(
            conn, run_id=run_id, generation_id=gen, sheet="Notes-SummaryofAccPol",
            row=10, block_ids=["b1", "b2"],
        )
    assert "accepted" in notes_agent._submit_coverage_entries_impl(deps, receipt)


def test_conflict_only_exempts_its_named_parts_from_batch_coverage(tmp_path, seeded):
    import json
    from pydantic_ai.models.test import TestModel
    from notes import source_write

    db, run_id, gen = seeded
    with repo.db_session(db) as conn:
        conn.execute("UPDATE notes_source_generations SET input_kind='prepared_document' WHERE id=?", (gen,))
        repo.insert_notes_review_flag(
            conn, run_id=run_id, kind="needs_human", reason="placement conflict",
            finding_id='["source_placement",1]',
            evidence=json.dumps({"generation_id": gen, "block_ids": ["b1"]}),
        )
    _agent, deps = notes_agent.create_notes_agent(
        template_type=NotesTemplateType.LIST_OF_NOTES, pdf_path="synthetic.pdf",
        inventory=[], filing_level="company", model=TestModel(), output_dir=str(tmp_path),
        run_id=run_id, db_path=db, source_generation_id=gen, batch_note_nums=[5],
    )
    deps.payload_sink = []
    deps.source_placement_conflict_notes.add(5)
    assert "b2" in notes_agent._submit_coverage_entries_impl(deps, [])

    with repo.db_session(db) as conn:
        source_write.write_cell_from_blocks(
            conn, run_id=run_id, generation_id=gen,
            sheet="Notes-Listofnotes", row=10, block_ids=["b2"],
        )
    assert "accepted" in notes_agent._submit_coverage_entries_impl(deps, [])


def test_source_collision_is_persisted_with_actionable_agent_feedback(seeded):
    from types import SimpleNamespace
    from notes import source_write

    db, run_id, gen = seeded
    with repo.db_session(db) as conn:
        for row, label in ((10, "Income tax"), (11, "Deferred tax")):
            conn.execute(
                "INSERT INTO notes_nodes(node_uuid,template_id,sheet,row,label,kind) "
                "VALUES(?,?,?,?,?,'LEAF')",
                (f"policy-{row}", "mfrs-company-policies-v1",
                 "Notes-SummaryofAccPol", row, label),
            )
        source_write.write_cell_from_blocks(
            conn,
            run_id=run_id,
            generation_id=gen,
            sheet="Notes-SummaryofAccPol",
            row=10,
            block_ids=["b1", "b2"],
            template_prefix="mfrs-company-",
        )
    deps = SimpleNamespace(
        db_path=db,
        run_id=run_id,
        source_generation_id=gen,
        filing_standard="mfrs",
        filing_level="company",
        sheet_name="Notes-SummaryofAccPol",
        source_placement_conflict_notes=set(),
        write_skip_errors=[],
    )

    result = notes_agent._write_from_source_impl(
        deps,
        "Notes-SummaryofAccPol",
        11,
        ["b1", "b2"],
        [1],
        "page 1",
        None,
    )

    assert "conflict recorded for review" in result
    assert "Income tax" in result and "Deferred tax" in result
    assert "provisional" in result
    assert deps.source_placement_conflict_notes == {5}
    with repo.db_session(db) as conn:
        flags = repo.fetch_notes_review_flags(conn, run_id)
    assert len(flags) == 1
    assert flags[0]["finding_id"].startswith('["source_placement",')
    assert flags[0]["source_pages"] == [1]


def test_source_tools_identify_reconstructed_content_without_requiring_repair(seeded):
    db, _, gen = seeded
    with repo.db_session(db) as conn:
        conn.execute("UPDATE notes_source_blocks SET locator_json=? WHERE generation_id=? AND block_id='b2'",
                     ('{"capture_uncertain":true,"capture_method":"reconstructed"}', gen))
        conn.commit()
    result = notes_agent._view_source_blocks_impl(db, gen, ["b2"])
    assert "best-effort reconstruction" in result
    assert "use captured content" in result
    assert "Stated at cost." in result
