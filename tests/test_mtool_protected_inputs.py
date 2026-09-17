"""The same draft/filing injector must respect the uploaded workbook."""
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection

from mtool.offline_fill import fill_workbook, write_patched_zip


@pytest.mark.parametrize("custom_format,blocked", [
    ("", False), (' customFormat="0"', False), (' customFormat="false"', False),
    (' customFormat="1"', True), (' customFormat="true"', True),
])
def test_row_style_requires_custom_format(tmp_path, custom_format, blocked):
    original, source = tmp_path / "original.xlsx", tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.protection.sheet = True
    ws.column_dimensions["D"].protection = Protection(locked=False)
    wb.save(original)
    wb.close()
    with ZipFile(original) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode()
    xml = xml.replace('<sheetData></sheetData>',
        '<sheetData><row r="5" s="0"' + custom_format + '/></sheetData>')
    write_patched_zip(str(original), str(source), {"xl/worksheets/sheet1.xml": xml.encode()})
    report = fill_workbook(str(source), {"writes": [
        {"sheet": "Sheet", "cell": "D5", "value": 42}
    ]}, dry_run=True)
    assert bool(report["unresolved"]) == blocked, report
    assert bool(report["written"]) != blocked


def test_only_unlocked_inputs_change_and_package_protection_survives(tmp_path):
    source, output = tmp_path / "template.xlsx", tmp_path / "draft.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "SOCIE"
    ws.protection.sheet = True
    ws.protection.password = "fixture"
    ws["B1"] = 1
    ws["B1"].protection = Protection(locked=False)
    ws["B2"] = "=B1+1"
    ws["B3"] = 8  # Locked literal, not a formula.
    ws["B4"].protection = Protection(locked=False)
    ws.column_dimensions["D"].protection = Protection(locked=False)
    wb.save(source)
    original = source.read_bytes()
    report = fill_workbook(str(source), {"writes": [
        {"sheet": "SOCIE", "cell": cell, "value": 42}
        for cell in ("B1", "B2", "B3", "B4", "C5", "D5")
    ]}, str(output))
    assert {x["cell"] for x in report["written"]} == {"B1", "B4", "D5"}
    assert {x["cell"] for x in report["skipped_formula"]} == {"B2"}
    assert {x["cell"] for x in report["unresolved"]} == {"B3", "C5"}
    assert all(x["reason"] == "protected_cell" for x in report["unresolved"])
    assert report["status"] == "degraded"
    assert source.read_bytes() == original
    result = load_workbook(output)
    assert result.active["B2"].value == "=B1+1"
    assert result.active["B3"].value == 8
    assert result.active.protection.sheet
    assert result.active.protection.password == ws.protection.password
    result.close()
    with ZipFile(source) as before, ZipFile(output) as after:
        for name in before.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert before.read(name) == after.read(name), name


def test_unprotected_sheet_allows_default_locked_style_but_not_formulas(tmp_path):
    source, output = tmp_path / "template.xlsx", tmp_path / "filled.xlsx"
    wb = Workbook()
    wb.active["A1"] = 1
    wb.active["A2"] = "=A1"
    wb.active["A2"].protection = Protection(locked=False)
    wb.save(source)
    report = fill_workbook(str(source), {"writes": [
        {"sheet": "Sheet", "cell": "A1", "value": 7},
        {"sheet": "Sheet", "cell": "A2", "value": 8},
    ]}, str(output))
    assert len(report["written"]) == 1
    assert len(report["skipped_formula"]) == 1
