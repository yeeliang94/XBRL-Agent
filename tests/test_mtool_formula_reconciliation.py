"""Reconcile native formula XML without trusting Excel's stored results."""
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from mtool.offline_fill import fill_workbook, write_patched_zip


@pytest.mark.parametrize("operand", ["", '<c r="B2" s="0"/>'])
@pytest.mark.parametrize("dry_run", [False, True])
def test_blank_arithmetic_operand_is_zero(tmp_path, operand, dry_run):
    source = _template(tmp_path,
        '<row r="1"><c r="B1"><v>5</v></c></row>'
        '<row r="2">' + operand + '</row>'
        '<row r="3"><c r="B3"><f>B1+B2</f><v>999</v></c></row>')
    report = fill_workbook(str(source), {"writes": [
        {"sheet": "SOCIE", "cell": "B3", "value": 5, "reconcile_formula": True}
    ]}, str(tmp_path / "filled.xlsx"), dry_run=dry_run)
    assert report["status"] == "ok", report
    assert report["reconciled_formula"][0]["found"] == "5"


@pytest.mark.parametrize("operand,formula", [
    ('<c r="B2" t="inlineStr"><is><t></t></is></c>', 'B1+B2'),
    ('<c r="B2" t="e"><v>#REF!</v></c>', 'B1+B2'),
    ('<c r="B2"><f/><v>0</v></c>', 'B1+B2'),
    ('', "B1+'Missing'!B2"),
])
def test_unverifiable_operand_does_not_become_blank_zero(tmp_path, operand, formula):
    source = _template(tmp_path,
        '<row r="1"><c r="B1"><v>5</v></c></row>'
        '<row r="2">' + operand + '</row>'
        '<row r="3"><c r="B3"><f>' + formula + '</f><v>5</v></c></row>')
    report = fill_workbook(str(source), {"writes": [
        {"sheet": "SOCIE", "cell": "B3", "value": 5, "reconcile_formula": True}
    ]}, dry_run=True)
    assert report["status"] == "degraded"
    assert not report["reconciled_formula"]


def _template(tmp_path, rows):
    original = tmp_path / "original.xlsx"
    source = tmp_path / "source.xlsx"
    wb = Workbook()
    wb.active.title = "SOCIE"
    wb.save(original)
    wb.close()
    xml = ('<worksheet xmlns="http://schemas.openxmlformats.org/'
           'spreadsheetml/2006/main"><sheetData>' + rows + '</sheetData></worksheet>')
    write_patched_zip(str(original), str(source), {"xl/worksheets/sheet1.xml": xml.encode()})
    return source


@pytest.mark.parametrize("dry_run", [False, True])
def test_shared_formulas_use_translated_current_inputs(tmp_path, dry_run):
    source = _template(tmp_path,
        '<row r="1"><c r="A1"><v>2</v></c><c r="B1"><v>3</v></c>'
        '<c r="C1"><v>4</v></c></row>'
        '<row r="2"><c r="A2"><v>5</v></c><c r="B2"><v>6</v></c>'
        '<c r="C2"><v>7</v></c></row>'
        '<row r="3"><c r="B3"><v>8</v></c></row>'
        '<row r="10"><c r="B10"><f t="shared" si="0"/><v>999</v></c>'
        '<c r="C10"><f t="shared" si="0" ref="B10:C11">'
        "C1+$A1+C$1+$A$1+'SOCIE'!C2</f><v>999</v></c></row>"
        '<row r="11"><c r="B11"><f t="shared" si="0"/><v>999</v></c></row>'
        '<row r="12"><c r="B12"><f>B10+B11</f><v>999</v></c></row>')
    output = tmp_path / "filled.xlsx"
    doc = {"writes": [{"sheet": "SOCIE", "cell": "B1", "value": 10}] + [
        {"sheet": "SOCIE", "cell": cell, "value": value, "reconcile_formula": True}
        for cell, value in [("B10", 30), ("C10", 19), ("B11", 31), ("B12", 61)]
    ]}
    with ZipFile(source) as z:
        xml = z.read("xl/worksheets/sheet1.xml")
    report = fill_workbook(str(source), doc, str(output), dry_run=dry_run)
    assert report["status"] == "ok", report
    assert len(report["reconciled_formula"]) == 4
    if not dry_run:
        with ZipFile(output) as z:
            assert z.read("xl/worksheets/sheet1.xml") == xml.replace(
                b'<c r="B1"><v>3</v></c>', b'<c r="B1"><v>10</v></c>')


@pytest.mark.parametrize("master", ["", '<c r="C2"><f t="shared" si="0" ref="B2:C2">IF(C1=7,7,0)</f><v>7</v></c>'])
def test_unresolvable_shared_formula_does_not_use_cached_value(tmp_path, master):
    source = _template(tmp_path,
        '<row r="1"><c r="B1"><v>7</v></c></row>'
        '<row r="2"><c r="B2"><f t="shared" si="0"/><v>7</v></c>' + master + '</row>')
    report = fill_workbook(str(source), {"writes": [
        {"sheet": "SOCIE", "cell": "B2", "value": 7, "reconcile_formula": True}
    ]}, dry_run=True)
    assert report["status"] == "degraded"
    assert report["unresolved"][0]["reason"] == "formula_not_verified"
    assert not report["reconciled_formula"]


def test_unaddressed_formula_cell_does_not_crash_reconciliation(tmp_path):
    source = _template(tmp_path,
        '<row r="1"><c r="B1"><v>7</v></c><c><f>B1+1</f><v>8</v></c></row>'
        '<row r="2"><c r="B2"><f>B1+1</f><v>8</v></c></row>')
    report = fill_workbook(str(source), {"writes": [
        {"sheet": "SOCIE", "cell": "B2", "value": 8, "reconcile_formula": True}
    ]}, dry_run=True)
    assert report["status"] == "ok", report
    assert report["reconciled_formula"][0]["found"] == "8"


def test_final_readback_uses_expressions_after_later_workbook_mutation(tmp_path):
    from mtool.offline_fill import verify_numeric_snapshot
    source = _template(tmp_path,
        '<row r="1"><c r="B1"><v>5</v></c><c r="C1"><v>-2</v></c></row>'
        '<row r="2"><c r="B2"><v>0</v></c></row>'
        '<row r="3"><c r="B3"><f>SUM(\'SOCIE\'!B1:C2)+1</f><v>999</v></c></row>')
    doc = {"writes": [{"sheet": "SOCIE", "cell": "B1", "value": 5}],
           "checks": [{"sheet": "SOCIE", "cell": "B3", "value": 4}]}
    report = verify_numeric_snapshot(str(source), doc)
    assert report["status"] == "ok", report
    assert report["native_recalculation_verified"] is False
    with ZipFile(source) as z:
        xml = z.read("xl/worksheets/sheet1.xml")
    output = tmp_path / "after-notes.xlsx"
    write_patched_zip(str(source), str(output), {
        "xl/worksheets/sheet1.xml": xml.replace(b'<v>5</v>', b'<v>6</v>')})
    report = verify_numeric_snapshot(str(output), doc)
    assert report["status"] == "degraded"
    assert len(report["mismatches"]) == 2


@pytest.mark.parametrize("cell", [
    '<c r="B3"><f>IF(B1=5,5,0)</f><v>5</v></c>',
    '<c r="B3"/>',
])
def test_final_readback_never_accepts_cached_or_blank_required_total(tmp_path, cell):
    from mtool.offline_fill import verify_numeric_snapshot
    source = _template(tmp_path, '<row r="1"><c r="B1"><v>5</v></c></row>'
                       '<row r="3">' + cell + '</row>')
    report = verify_numeric_snapshot(str(source), {"checks": [
        {"sheet": "SOCIE", "cell": "B3", "value": 5}]})
    assert report["status"] == "degraded"
    assert len(report["unverified"]) == 1
    assert not report["verified"]


@pytest.mark.parametrize("sheet,reference", [
    ("Other", "Other"),
    ("Other.Data", "Other.Data"),
    ("Other-Data", "'Other-Data'"),
    ("Other's Data", "'Other''s Data'"),
])
@pytest.mark.parametrize("shared", [False, True])
def test_subtraction_before_cross_sheet_reference_reconciles(tmp_path, sheet, reference, shared):
    from mtool.offline_fill import verify_numeric_snapshot

    source = tmp_path / "source.xlsx"
    wb = Workbook()
    wb.active.title = "Main"
    wb.create_sheet(sheet)
    # This legal sheet name must not steal the local A1 minus Other!A1 expression.
    wb.create_sheet("A1-Other")["A1"] = 99
    wb[sheet]["A1"] = 3
    wb[sheet]["A2"] = 4
    wb.save(source)
    wb.close()
    master = ' t="shared" si="0" ref="B1:B2"' if shared else ''
    follower = ('<f t="shared" si="0"/>' if shared
                else f'<f>A2-{reference}!A2</f>')
    xml = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
           '<sheetData><row r="1"><c r="A1"><v>10</v></c>'
           f'<c r="B1"><f{master}>A1-{reference}!A1</f><v>99</v></c></row>'
           '<row r="2"><c r="A2"><v>20</v></c>'
           f'<c r="B2">{follower}<v>99</v></c></row></sheetData></worksheet>')
    patched = tmp_path / "formulas.xlsx"
    write_patched_zip(str(source), str(patched), {"xl/worksheets/sheet1.xml": xml.encode()})
    checks = [{"sheet": "Main", "cell": cell, "value": value}
              for cell, value in [("B1", 7), ("B2", 16)]]
    output = tmp_path / "filled.xlsx"
    report = fill_workbook(str(patched), {"writes": [
        dict(check, reconcile_formula=True) for check in checks
    ]}, str(output))
    assert report["status"] == "ok", report
    assert len(report["reconciled_formula"]) == 2
    assert verify_numeric_snapshot(str(output), {"checks": checks})["status"] == "ok"
    # The incorrect value from the similarly named sheet must never pass.
    wrong = verify_numeric_snapshot(str(output), {"checks": [
        {"sheet": "Main", "cell": "B1", "value": 99}]})
    assert len(wrong["mismatches"]) == 1
    with ZipFile(output) as z:
        assert z.read("xl/worksheets/sheet1.xml") == xml.encode()


@pytest.mark.parametrize("expression,expected", [
    ("A1-Other!A1", "A2-Other!A2"),
    ("$A1-Other!A$1", "$A2-Other!A$1"),
    ("A1-Other.Data!A1", "A2-Other.Data!A2"),
    ("A1-'Other-Data'!A1", "A2-'Other-Data'!A2"),
])
def test_shared_formula_subtraction_translates_both_operands(expression, expected):
    from mtool.offline_fill import _translate_shared_formula
    assert _translate_shared_formula(expression, "B1", "B2") == expected


@pytest.mark.parametrize("sheet,reference", [
    ("Other", "Other"), ("Other.Data", "Other.Data"),
    ("Other-Data", "'Other-Data'"),
])
def test_subtraction_of_cross_sheet_sum_preserves_range_names(sheet, reference):
    from mtool.offline_fill import _calculated_value
    cells = {
        "Main": {1: {"A": ("N", "10"), "B": ("F", f"A1-SUM({reference}!A1:A2)")}},
        sheet: {1: {"A": ("N", "3")}, 2: {"A": ("N", "4")}},
    }
    assert _calculated_value(cells, "Main", "B1") == 3
