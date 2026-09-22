"""Prepared source prose on numeric templates must leave numeric cells intact."""
import asyncio
import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pydantic_ai.models.test import TestModel

from concept_model.importer import import_template
from concept_model.filing_targets import persist_template_manifest
from concept_model.parser import parse_template
from db import repository as repo
from db.schema import init_db
from notes import agent as notes_agent, source_repository as srepo, source_write
from notes.source_models import SourceBlock, SourceNote
from notes_types import NotesTemplateType, NOTES_REGISTRY, notes_template_path


@pytest.mark.parametrize("kind", [NotesTemplateType.ISSUED_CAPITAL, NotesTemplateType.RELATED_PARTY])
def test_prepared_group_prompt_explains_numeric_entity_columns(kind):
    prompt = notes_agent.render_notes_prompt(kind, "group", [], prepared_source_required=True)
    assert "group_cy, group_py, company_cy and" in prompt
    assert "company_py map to columns B, C, D and E respectively" in prompt
    assert "Never copy a Group amount" in prompt
    assert "omit undisclosed scopes" in prompt


@pytest.mark.parametrize("level", ["company", "group"])
@pytest.mark.parametrize("standard", ["mfrs", "mpers"])
@pytest.mark.parametrize("kind", [NotesTemplateType.ISSUED_CAPITAL, NotesTemplateType.RELATED_PARTY])
def test_source_disclosure_and_numeric_facts_use_distinct_slots(tmp_path, standard, kind, level):
    db = tmp_path / "audit.sqlite"
    init_db(db)
    template = notes_template_path(kind, standard=standard, level=level)
    tree = parse_template(str(template))
    manifest = tmp_path / "template.json"
    manifest.write_text(json.dumps(tree.to_json()))
    import_template(db, manifest)
    persist_template_manifest(db, template)
    sheet = NOTES_REGISTRY[kind].sheet_name
    with repo.db_session(db) as conn:
        run = repo.create_run(conn, "source.pdf", session_id="s", output_dir=str(tmp_path))
        gen = srepo.begin_generation(conn, run, input_kind="prepared_document")
        srepo.write_blocks(conn, gen, [SourceBlock("a", "paragraph", 1,
            "<p>Complete source disclosure with <strong>emphasis</strong>.</p>", source_note_id="n5")])
        srepo.write_notes(conn, gen, [SourceNote("n5", "5", "Disclosure", ["a"])])
        srepo.activate_generation(conn, gen)
        # Taxonomy-based target selection accepts prose and rejects money/share rows.
        prose = source_write.resolve_target(conn, sheet, 4, template_prefix=f"{standard}-{level}-")
        assert prose["numeric_note_prose"]
        numeric = conn.execute(
            "SELECT n.render_row,n.canonical_label FROM concept_nodes n "
            "JOIN concept_semantic_addresses sa ON sa.concept_uuid=n.concept_uuid "
            "JOIN taxonomy_concepts tc ON tc.source_element_id=sa.primary_concept "
            "WHERE n.render_sheet=? AND n.kind='LEAF' AND tc.data_type LIKE '%monetaryItemType' LIMIT 1",
            (sheet,),
        ).fetchone()
        assert numeric
        with pytest.raises(source_write.SourceWriteError):
            source_write.resolve_target(conn, sheet, numeric["render_row"], template_prefix=f"{standard}-{level}-")
    agent, deps = notes_agent.create_notes_agent(template_type=kind, pdf_path="/tmp/no.pdf", inventory=[],
        filing_level=level, filing_standard=standard, model=TestModel(),
        output_dir=str(tmp_path), run_id=run, db_path=str(db), source_generation_id=gen)
    funcs = {name: tool.function for ts in agent.toolsets for name, tool in getattr(ts, "tools", {}).items()}
    assert "write_note_from_source" in funcs
    ctx = SimpleNamespace(deps=deps)
    number = asyncio.run(funcs["write_notes"](ctx, [{"chosen_row_label": numeric["canonical_label"],
        "content": "", "numeric_values": {f"{level}_cy": 123}, "evidence": "Page 1",
        "source_pages": [1], "parent_note": {"number": "5", "title": "Disclosure"}}]))
    assert "Wrote 1" in number
    written = asyncio.run(funcs["write_note_from_source"](ctx, sheet=sheet, row=4,
        block_ids=["a"], source_pages=[1], evidence="Page 1"))
    assert "ok:" in written
    book = load_workbook(deps.filled_path)
    assert book[sheet].cell(numeric["render_row"], 2).value == 123
    assert "Complete source disclosure" in book[sheet].cell(4, 2).value
    book.close()
    assert deps.numeric_cells[0]["value"] == 123
    assert "<strong>emphasis</strong>" in deps.cells_written[0]["html"]


def test_numeric_prose_duplicate_is_explicit_and_limited_to_list_of_notes(tmp_path):
    from notes.integrity_runner import build_input
    from notes.integrity import check_approved_duplicates
    db = tmp_path / "audit.sqlite"
    init_db(db)
    template = notes_template_path(NotesTemplateType.ISSUED_CAPITAL)
    tree = parse_template(str(template))
    manifest = tmp_path / "template.json"
    manifest.write_text(json.dumps(tree.to_json()))
    import_template(db, manifest)
    persist_template_manifest(db, template)
    with repo.db_session(db) as conn:
        run = repo.create_run(conn, "source.pdf", session_id="s", output_dir=str(tmp_path))
        gen = srepo.begin_generation(conn, run, input_kind="prepared_document")
        srepo.write_blocks(conn, gen, [SourceBlock("a", "paragraph", 1, "<p>Full prose.</p>")])
        srepo.activate_generation(conn, gen)
        source_write.write_cell_from_blocks(conn, run_id=run, generation_id=gen,
            sheet="Notes-Issuedcapital", row=4, block_ids=["a"], template_prefix="mfrs-company-")
        source_write.write_cell_from_blocks(conn, run_id=run, generation_id=gen,
            sheet="Notes-Listofnotes", row=20, block_ids=["a"])
        assert check_approved_duplicates(build_input(conn, run, gen)) == []
        with pytest.raises(source_write.SourcePlacementConflict):
            source_write.write_cell_from_blocks(conn, run_id=run, generation_id=gen,
                sheet="Notes-SummaryofAccPol", row=20, block_ids=["a"])
        assert check_approved_duplicates(build_input(conn, run, gen)) == []
