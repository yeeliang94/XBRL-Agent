"""Category identities must survive writes, review and filing snapshots."""
import sqlite3

import pytest

from concept_model.dimensions import dimension_key
from concept_model.facts_api import FactWrite, write_fact
from concept_model.versioning import snapshot_facts, compute_review_diff, revert_to_original
from test_facts_api_inprocess import db_and_run


def test_numeric_note_payloads_reach_distinct_native_categories(tmp_path):
    from openpyxl import Workbook, load_workbook
    from notes.payload import NotesPayload
    from notes.writer import write_notes_workbook
    from notes_types import NotesTemplateType, notes_template_path
    from concept_model.cell_resolver import project_writes
    from db.schema import init_db
    from mtool.exporter import build_fill_doc
    from mtool.template_map import resolve_filing_doc
    from mtool.offline_fill import fill_workbook, verify_numeric_snapshot
    from test_mtool_exporter import _import, _init_run

    db = tmp_path / 'audit.db'
    init_db(db)
    template = notes_template_path(NotesTemplateType.ISSUED_CAPITAL, level='company')
    tid = _import(db, template)
    run = _init_run(db)
    axis = 'ifrs-full_ClassesOfShareCapitalAxis'
    members = ['ifrs-full_OrdinarySharesMember', 'ssmt-mfrs_RedeemablePreferenceSharesMember']
    payloads = [NotesPayload(
        chosen_row_label='*Number of shares issued and fully paid', content='',
        evidence='Page 4: synthetic share register, class ' + member,
        parent_note={'number': '4', 'title': 'Share capital'},
        numeric_values={'cy': value, 'py': value - 5}, dimensions={axis: member},
    ) for member, value in zip(members, [100, 20])]
    result = write_notes_workbook(str(template), payloads, str(tmp_path / 'diagnostic.xlsx'), 'company', 'Notes-Issuedcapital')
    projected = project_writes(db, run, tid, result.numeric_cells)
    assert projected.projected == 4
    from api.notes import _numeric_sheet_rows
    with sqlite3.connect(db) as conn:
        conn.row_factory = sqlite3.Row
        rows = _numeric_sheet_rows(conn, run, tid, 'Notes-Issuedcapital', 'company')
    row = next(r for r in rows if r.get('categories'))
    assert sorted(c['values']['cy'] for c in row['categories']) == [20, 100]
    assert sorted(c['values']['py'] for c in row['categories']) == [15, 95]
    doc = build_fill_doc(db, run, filing_standard='mfrs', filing_level='company')
    assert len(doc['writes']) == 4
    assert len({w['dimension_key'] for w in doc['writes']}) == 2

    wb = Workbook(); ws = wb.active; ws.title = 'Notes-Issuedcapital'
    prefix = 'synthetic.xsd#'
    ws['C3'] = '#DOM#'; ws['D3'] = '#PRIM#'; ws['C5'] = '#ENDT#'
    for col, member in zip('EF', members):
        ws[f'{col}2'] = '::'.join(prefix + item for item in ['ifrs-full_DisclosureOfClassesOfShareCapitalTable', axis, member])
        ws[f'{col}5'] = '31/12/2026'
    ws['A7'] = prefix + doc['writes'][0]['semantic_address']['primary_concept']
    ws['D7'] = payloads[0].chosen_row_label
    native = tmp_path / 'native.xlsx'; wb.save(native); wb.close()
    ready, coverage = resolve_filing_doc(str(native), doc)
    assert coverage['unmapped'] == 2  # Comparatives persist; no PY section.
    assert all(w['period'] == 'CY' for w in ready['writes'])
    output = tmp_path / 'filled.xlsx'
    assert fill_workbook(str(native), ready, str(output))['status'] == 'ok'
    assert verify_numeric_snapshot(str(output), ready)['status'] == 'ok'
    filled = load_workbook(output)
    assert [filled.active[f'{c}7'].value for c in 'EF'] == [100, 20]
    filled.close()
    repeated = tmp_path / 'repeated.xlsx'
    assert fill_workbook(str(output), ready, str(repeated))['status'] == 'ok'
    assert repeated.read_bytes() == output.read_bytes()
    # A template explicitly declaring a comparative block gets both periods;
    # the current-only template above never receives PY figures in CY cells.
    comparative = load_workbook(native)
    ws = comparative.active
    ws['C12'] = '#DOM#'; ws['D12'] = '#PRIM#'; ws['C14'] = '#ENDT#'
    for col in 'EF':
        ws[f'{col}11'] = ws[f'{col}2'].value
        ws[f'{col}14'] = '31/12/2025'
    ws['A16'] = ws['A7'].value
    ws['D16'] = ws['D7'].value
    comparative_path = tmp_path / 'comparative.xlsx'
    comparative.save(comparative_path); comparative.close()
    ready, coverage = resolve_filing_doc(str(comparative_path), doc)
    assert coverage['unmapped'] == coverage['ambiguous'] == 0, coverage
    complete = tmp_path / 'both-periods.xlsx'
    assert fill_workbook(str(comparative_path), ready, str(complete))['status'] == 'ok'
    assert verify_numeric_snapshot(str(complete), ready)['status'] == 'ok'
    filled = load_workbook(complete)
    assert [filled.active[ref].value for ref in ['E7', 'F7', 'E16', 'F16']] == [100, 20, 95, 15]
    filled.close()


def test_dimension_order_is_not_identity():
    assert dimension_key({'b': 'two', 'a': 'one'}) == dimension_key({'a': 'one', 'b': 'two'})
    assert dimension_key({}) == ''
    with pytest.raises(ValueError):
        dimension_key({'': 'unknown'})


def test_numeric_category_replacement_uses_submission_order_in_both_outputs(tmp_path):
    from notes.payload import NotesPayload
    from notes.writer import write_notes_workbook
    from notes_types import NotesTemplateType, notes_template_path
    from openpyxl import load_workbook
    payloads = [NotesPayload(
        chosen_row_label='*Number of shares issued and fully paid', content='',
        parent_note={'number': '4', 'title': 'Share capital'}, source_pages=[page],
        evidence=f'Page {page}: synthetic share count',
        numeric_values={'cy': value, **({'py': 40} if page == 9 else {})},
        dimensions={'ifrs-full_ClassesOfShareCapitalAxis': 'ifrs-full_OrdinarySharesMember'},
    ) for page, value in [(9, 50), (3, 60)]]
    path = tmp_path / 'diagnostic.xlsx'
    result = write_notes_workbook(str(notes_template_path(NotesTemplateType.ISSUED_CAPITAL, level='company')),
                                 payloads, str(path), 'company', 'Notes-Issuedcapital')
    assert len(result.numeric_cells) == 2
    assert sorted(c['value'] for c in result.numeric_cells) == [40, 60]
    book = load_workbook(path)
    for cell in result.numeric_cells:
        assert book[cell['sheet']].cell(cell['row'], cell['col']).value == cell['value']
    book.close()


@pytest.mark.parametrize('level', ['company', 'group'])
def test_categories_coexist_and_revert_independently(tmp_path, level):
    from notes_types import NotesTemplateType, notes_template_path
    from db.schema import init_db
    from test_mtool_exporter import _import, _init_run
    db = tmp_path / 'facts.db'
    init_db(db)
    tid = _import(db, notes_template_path(NotesTemplateType.RELATED_PARTY, level=level), group=level == 'group')
    run = _init_run(db)
    with sqlite3.connect(db) as conn:
        uuid = conn.execute("SELECT concept_uuid FROM concept_nodes WHERE template_id=? AND kind='LEAF' AND canonical_label NOT LIKE '%Disclosure%' ORDER BY render_row LIMIT 1", (tid,)).fetchone()[0]
    axis = 'ifrs-full_CategoriesOfRelatedPartiesAxis'
    members = ['ifrs-full_ParentMember', 'ifrs-full_AssociatesMember']
    scope = 'Group' if level == 'group' else 'Company'
    for category, value in zip(members, [100, 20]):
        write_fact(db, run, FactWrite(concept_uuid=uuid, dimensions={axis: category}, entity_scope=scope, value=value))
    snapshot_facts(db, run)
    write_fact(db, run, FactWrite(concept_uuid=uuid, dimensions={axis: members[0]}, entity_scope=scope, value=110))
    changes = compute_review_diff(db, run)
    assert len(changes) == 1
    assert changes[0]['original'] == 100
    assert changes[0]['current'] == 110
    revert_to_original(db, run)
    with sqlite3.connect(db) as conn:
        values = conn.execute('SELECT dimension_key,value FROM run_concept_facts WHERE concept_uuid=? ORDER BY dimension_key', (uuid,)).fetchall()
    assert values == sorted([(dimension_key({axis: member}), value) for member, value in zip(members, [100, 20])])


def test_face_fact_cannot_acquire_an_unsupported_category(db_and_run):
    from fastapi import HTTPException
    db, run, uuid = db_and_run
    with pytest.raises(HTTPException, match='additional category axis'):
        write_fact(db, run, FactWrite(concept_uuid=uuid, dimensions={'ClassAxis': 'Ordinary'}, value=1))
