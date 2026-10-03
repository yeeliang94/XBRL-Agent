"""CLBG is an explicit filing family with its own taxonomy and fund dimensions."""
from pathlib import Path

import pytest

from concept_model.parser import parse_template
from notes_types import NotesTemplateType, notes_template_path
from statement_types import StatementType, template_path, variants_for_standard


def test_clbg_registry_and_request_preserve_filing_family():
    from server import RunConfigRequest, RunConfigPatchRequest
    from scout.standard_detector import detect_filing_standard
    config = RunConfigRequest(statements=["SOFP", "SOPL", "SOCF", "SOCIE"], filing_standard="clbg")
    assert config.filing_standard == "clbg"
    assert RunConfigPatchRequest(filing_standard="clbg").model_dump(exclude_unset=True) == {"filing_standard": "clbg"}
    assert detect_filing_standard("Company limited by guarantee. MFRS financial statements.") == "clbg"
    assert [v.name for v in variants_for_standard(StatementType.SOCI, "clbg")] == ["NotPrepared"]
    assert [v.name for v in variants_for_standard(StatementType.SOCF, "clbg")] == ["Indirect"]
    for statement, variant in ((StatementType.SOFP, "OrderOfLiquidity"), (StatementType.SOCF, "Direct"), (StatementType.SOCIE, "SoRE")):
        with pytest.raises(ValueError):
            template_path(statement, variant, standard="clbg")
    with pytest.raises(ValueError, match="Company"):
        RunConfigRequest(statements=["SOFP"], filing_standard="clbg", filing_level="group")
    with pytest.raises(ValueError, match="issued-capital"):
        notes_template_path(NotesTemplateType.ISSUED_CAPITAL, standard="clbg")
    assert "XBRL-template-MFRS" in str(template_path(StatementType.SOFP, "CuNonCu"))
    assert "XBRL-template-MPERS" in str(template_path(StatementType.SOFP, "CuNonCu", standard="mpers"))


def test_clbg_generated_templates_have_exact_taxonomy_addresses_and_fund_targets(tmp_path):
    from scripts.generate_clbg_templates import TEMPLATES, build_template, fund_components
    from concept_model.taxonomy_semantics import semantic_addresses_for
    root = Path(__file__).resolve().parents[1] / "XBRL-template-CLBG/Company"
    for filename in TEMPLATES:
        tree = parse_template(str(root / filename))
        assert tree.template_id.startswith("clbg-company-")
        assert all(n.render_key.get("semantic_address") for n in tree.concepts if n.kind != "ABSTRACT")
        # Regeneration must reproduce semantic row/column/formula content,
        # without rewriting the maintained workbook or changing other standards.
        regenerated = build_template(filename, tmp_path / "XBRL-template-CLBG/Company")
        new_tree = parse_template(str(regenerated))
        assert tree.to_json() == new_tree.to_json()
    tree = parse_template(str(root / "09-SOCIE.xlsx"))
    assert tree.shape == "matrix"
    opening = [n for n in tree.concepts if n.canonical_label.lstrip("*") == "Balance at beginning of period"]
    assert len(opening) == len(fund_components()) == 13
    accumulated = next(n for n in opening if n.render_key["matrix_col"] == "B")
    address = accumulated.render_key["semantic_address"]
    assert address["primary_concept"] == "ssmt-mfrs_FundBalance"
    assert address["dimensions"] == {"ssmt-mfrs_ComponentsOfFundAxis": "ssmt-mfrs_AccumulatedFundsDeficitMember"}
    assert {(t["period"], t["entity_scope"], t["col"]) for t in accumulated.render_key["targets"]} == {("CY", "Company", "B"), ("PY", "Company", "B")}
    assert semantic_addresses_for(str(root / "09-SOCIE.xlsx"))


def test_clbg_prompt_and_checks_use_fund_family_only():
    from prompts import render_prompt
    from cross_checks.framework import build_default_cross_checks
    from cross_checks.clbg import clbg_checks
    prompt = render_prompt(StatementType.SOCIE, "Default", filing_standard="clbg")
    assert "CLBG CHANGES IN FUND" in prompt
    assert "FILING FAMILY: CLBG" in prompt
    assert "B–X" not in prompt
    checks = clbg_checks()
    assert len(checks) == 3
    assert all(c.applies_to({"filing_standard": "clbg"}) for c in checks)
    assert not any(c.applies_to({"filing_standard": "mfrs"}) for c in checks)
    assert {c.name for c in checks} <= {c.name for c in build_default_cross_checks("clbg")}
    assert {c.name for c in checks}.isdisjoint(c.name for c in build_default_cross_checks("mfrs"))


@pytest.mark.parametrize("check_name", ["clbg_sofp_balance", "clbg_income_to_fund", "clbg_fund_to_sofp"])
def test_clbg_tieouts_reject_disagreement_in_populated_period(tmp_path, check_name):
    from cross_checks.clbg import clbg_checks
    from openpyxl import Workbook
    check = next(c for c in clbg_checks() if c.name == check_name)
    # Isolated source controls exercise public workbook checks without any
    # generated balancing formulas hiding a disagreement.
    paths = {}
    for statement in check.required_statements:
        wb = Workbook()
        ws = wb.active
        ws.title = "SOCIE" if statement == StatementType.SOCIE else "SOIE-Nature" if statement == StatementType.SOPL else "SOFP-CuNonCu"
        for side, label, matrix in (("left", check.left_label, check.left_matrix), ("right", check.right_label, check.right_matrix)):
            if getattr(check, f"{side}_statement") != statement:
                continue
            row = 6 if matrix else 3 if side == "left" else 4
            ws.cell(row, 1, label)
            ws.cell(row, 14 if matrix else 2, 100)
        path = tmp_path / f"{statement.value}.xlsx"
        wb.save(path)
        wb.close()
        paths[statement] = str(path)
    assert check.run(paths, 0.01).status == "passed"
    from openpyxl import load_workbook
    wb = load_workbook(paths[check.right_statement])
    wb.active.cell(6 if check.right_matrix else 4, 14 if check.right_matrix else 2, 101)
    wb.save(paths[check.right_statement])
    wb.close()
    assert check.run(paths, 0.01).status == "failed"
