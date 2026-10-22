"""Report filing boundaries: typed canonical values, review edits and native patching."""
import json
import sqlite3
from contextlib import closing
from dataclasses import asdict
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Protection

from db import repository as repo
from db.schema import init_db
from reports.registry import REPORTS, fields, import_templates, template_path
from reports.repository import save_field, snapshot, StaleReportEdit
from reports.exporter import resolve_reports
from mtool.offline_fill import fill_report_fields, load_workbook_entries, format_value
from mtool.notes_exporter import build_notes_fill_doc


@pytest.mark.parametrize("family", ["MFRS_Company", "MFRS_Group", "MPERS_Company"])
def test_native_report_narratives_keep_labels_and_fill_input_cells_and_auditor_details(tmp_path, family):
    """Native reports combine narrative and auditor tables with different markers.

    Synthetic single-table fixtures miss the metadata-density bug that put a
    popup into the visible label column. Read the saved native package back.
    """
    from mtool.offline_fill import (
        fill_footnotes, get_sheet_paths, get_shared_strings, read_sheet_cells,
        inspect_footnotes, resolve_sheet_name,
    )

    source = next((Path(__file__).resolve().parents[1] / "data").glob(f"mTool_{family}_*.xlsx"), None)
    if source is None:
        pytest.skip(f"Native {family} mTool fixture is not installed")
    _, before, _ = load_workbook_entries(str(source))
    paths = get_sheet_paths(before)
    shared = get_shared_strings(before)
    cells = {s: read_sheet_cells(before[p], shared) for s, p in paths.items()}
    narratives = [
        ("Report-Directors", "ssmt_DisclosureOfDirectorsReportExplanatory"),
        ("Report-StatementByDirectors", "ssmt_DisclosureOfStatementByDirectorsExplanatory"),
        ("Report-Auditors", "ssmt_DisclosureOfAuditorsReportToMembersExplanatory"),
    ]
    footnotes = [{"source_sheet": sheet, "primary_concept": concept,
                  "label": sheet, "html": f"<h3>Opinion</h3><p>Verification {sheet}.</p>"}
                 for sheet, concept in narratives]
    labels_only = [{"source_sheet": sheet,
                    "label": cells[resolve_sheet_name(sheet, cells)][15]["D"][1],
                    "html": "<p>Label placement verification.</p>"} for sheet, _ in narratives]
    preview = fill_footnotes(str(source), {"footnotes": labels_only, "strict": True},
                             dry_run=True, create_missing=True)
    assert not preview["unresolved"], preview
    assert len(preview["footnotes_written"]) == 3
    assert {w["visible_cell"].rsplit("!", 1)[1] for w in preview["footnotes_written"]} == {"E15"}
    notes_out = tmp_path / "narratives.xlsx"
    prose = fill_footnotes(str(source), {"footnotes": footnotes, "strict": True},
                          str(notes_out), create_missing=True)
    assert not prose["unresolved"], prose
    assert not prose["footnote_mismatches"], prose
    assert not prose["errors"], prose
    assert len(prose["footnotes_written"]) == 3

    scalar_values = {
        "ssmt_LicenseNumberOfAuditor": "0012",
        "ssmt_NameOfAuditorSigningReport": "Verification Auditor",
        "ssmt_RegistrationNumberOfAuditFirm": "LLP0010081-LCA & AF 0758",
        "ssmt_NameOfAuditFirm": "Verification & Co.",
        "ssmt_AddressOne": "1 Example Street",
        "ssmt_AddressTwo": "Suite 2",
        "ssmt_AddressThree": "Example Building",
        "ssmt_PostcodeOfAuditFirm": "00100",
        "ssmt_TownWhereAuditFirmIsLocated": "Kuala Lumpur",
        "ssmt_StateWhereAuditFirmIsLocated": "Wilayah Persekutuan",
        "ssmt_DateOfSigningAuditorsReport": "2026-10-17",
        "ssmt_TypeOfAuditorsOpinion": "Unqualified opinion",
    }
    report = {"sheet": "Report-Auditors", "fields": [
        {"concept": concept, "label": concept, "value": value, "value_status": "observed",
         "kind": "date" if "DateOf" in concept else "choice" if "StateWhere" in concept or "TypeOf" in concept else "text"}
        for concept, value in scalar_values.items()]}
    _, note_data, _ = load_workbook_entries(str(notes_out))
    # The unique narrative marker scopes shared address concepts to this report.
    report["fields"].append({"concept": narratives[-1][1], "kind": "html",
                             "value": footnotes[-1]["html"], "value_status": "observed"})
    resolved = resolve_reports(note_data, {"reports": [report], "digest": "verification"})
    assert {w["cell"] for w in resolved["writes"]} == {f"{col}32" for col in "EFGHIJKLM"}
    firm_number = next(w for w in resolved["writes"] if w["concept"] == "ssmt_RegistrationNumberOfAuditFirm")
    assert firm_number["value"] == "AF0758"
    assert firm_number["source_value"] == scalar_values["ssmt_RegistrationNumberOfAuditFirm"]
    assert next(f["value"] for f in report["fields"] if f["concept"] == firm_number["concept"]) == firm_number["source_value"]
    # Native picker/date cells are protected and have no verified encoding.
    assert {u["concept"] for u in resolved["unresolved"]} == {
        "ssmt_StateWhereAuditFirmIsLocated", "ssmt_DateOfSigningAuditorsReport", "ssmt_TypeOfAuditorsOpinion"}
    output = tmp_path / "filled.xlsx"
    scalars = fill_report_fields(str(notes_out), resolved, str(output))
    assert not scalars["mismatches"]
    assert len(scalars["written"]) == 9
    _, after, _ = load_workbook_entries(str(output))
    after_shared = get_shared_strings(after)
    after_cells = {s: read_sheet_cells(after[p], after_shared) for s, p in get_sheet_paths(after).items()}
    for sheet, _ in narratives:
        native = resolve_sheet_name(sheet, cells)
        assert after_cells[native][15]["D"] == cells[native][15]["D"]
        assert after_cells[native][15]["E"] == ("S", "[Text block added]")
    targets = inspect_footnotes(after)["targets"]
    for sheet, _ in narratives:
        native = resolve_sheet_name(sheet, cells)
        assert any(t["sheet"] == native and t["cell"] == "E15" and t["payload_populated"] for t in targets)
    audit_sheet = resolve_sheet_name("Report-Auditors", cells)
    assert after_cells[audit_sheet][32]["E"] == ("S", "0012")
    assert after_cells[audit_sheet][32]["G"] == ("S", "AF0758")
    assert after_cells[audit_sheet][32]["L"] == ("S", "00100")
    for row, col in ((16, "E"), (17, "E"), (32, "N")):
        assert after_cells[audit_sheet][row][col] == cells[audit_sheet][row][col]
    # Protection, formulas and every unrelated worksheet retain their original XML.
    modified = {paths[resolve_sheet_name(sheet, cells)] for sheet, _ in narratives}
    modified.add(paths["+FootnoteTexts"])
    for path in paths.values():
        if path not in modified:
            assert after[path] == before[path]


@pytest.fixture
def report_run(tmp_path):
    db = tmp_path / "audit.db"
    init_db(db)
    import_templates(db)
    with repo.db_session(db) as conn:
        rid = repo.create_run(conn, "source.pdf", status="completed", config={
            "filing_standard": "mfrs", "filing_level": "company",
            "reports_to_run": list(REPORTS), "denomination": "units"})
    return db, rid


@pytest.mark.parametrize("standard", ["mfrs", "mpers"])
@pytest.mark.parametrize("level", ["company", "group"])
def test_report_catalogue_is_writable_in_each_exact_filing_family(report_run, standard, level):
    from concept_model.filing_targets import targets_for_template
    for report, count in (("DIRECTORS_REPORT", 22), ("STATEMENT_BY_DIRECTORS", 27), ("AUDITORS_REPORT", 13)):
        _, targets = targets_for_template(template_path(report, standard, level))
        writable = {t.taxonomy_element_id: t for t in targets if t.writable}
        assert len(writable) == count
        for field in fields(report, standard):
            assert writable[field.concept].value_kind == field.kind
        assert sum(f.kind == "html" for f in fields(report, standard)) == 1


def test_report_values_round_trip_edits_clear_and_stale_saves(report_run):
    db, rid = report_run
    report = "AUDITORS_REPORT"
    identifier = next(f for f in fields(report) if f.concept.endswith("LicenseNumberOfAuditor"))
    narrative = next(f for f in fields(report) if f.kind == "html")
    with repo.db_session(db) as conn:
        save_field(conn, rid, report, identifier.concept, "0012", "observed", "Licence 0012", [3])
        save_field(conn, rid, report, narrative.concept, "<h3>Opinion</h3><p>Full report.</p>", "observed", "Opinion", [2, 3])
    with repo.db_session(db) as conn:
        before = snapshot(conn, rid)
        view = next(r for r in before["reports"] if r["type"] == report)
        assert next(f for f in view["fields"] if f["concept"] == identifier.concept)["value"] == "0012"
        assert repo.list_notes_cells_for_run(conn, rid) == []  # reports cannot become financial notes
        save_field(conn, rid, report, identifier.concept, None, "not_disclosed", "Cleared", [], expected_revision=1)
        with pytest.raises(StaleReportEdit):
            save_field(conn, rid, report, identifier.concept, "9999", "user_override", "", [], expected_revision=1)
        assert snapshot(conn, rid)["digest"] != before["digest"]
    exported = build_notes_fill_doc(db, rid, decorate=False)
    assert len(exported["footnotes"]) == 1
    assert exported["footnotes"][0]["html"] == "<h3>Opinion</h3><p>Full report.</p>"


@pytest.mark.parametrize("report", list(REPORTS))
def test_report_prose_reuses_cleanup_and_formatter_without_entering_financial_notes(report_run, report):
    from notes.auto_format import candidate_sheets
    from notes.cleanup_agent import cleanup_cells
    from notes.cleanup_patch import Patch, Removal
    from notes.cleanup_repository import save_cleanup
    from reports.repository import report_view

    db, rid = report_run
    narrative = next(f for f in fields(report) if f.kind == "html")
    sheet = REPORTS[report][2]
    html = "<h3>EXAMPLE SDN BHD</h3><h3>Report</h3><p>Complete substantive wording.</p>"
    with repo.db_session(db) as conn:
        save_field(conn, rid, report, narrative.concept, html, "observed", "Report", [1])
        assert repo.list_notes_cells_for_run(conn, rid) == []
    assert candidate_sheets(str(db), rid, [sheet]) == [sheet]
    cells = cleanup_cells(str(db), rid, [sheet])
    cid = f"{sheet}:{narrative.row}"
    patch = Patch(inspected_cells=[cid], removals=[Removal(cell=cid, block=0,
        expected_text="EXAMPLE SDN BHD", reason="page_banner", evidence_page=1,
        justification="Company banner repeated at the top of the source page")])
    with repo.db_session(db) as conn:
        saved = save_cleanup(conn, run_id=rid, cells=cells, patch=patch, viewed_pages={1}, model="test")
        assert saved["changed_rows"] == 1
        view = report_view(conn, rid, report)
        value = next(f["value"] for f in view["fields"] if f["kind"] == "html")
        assert "EXAMPLE SDN BHD" not in value
        assert "Complete substantive wording." in value
        # Cleanup changed the saved text, so an edit made against it is stale.
        assert next(f["revision"] for f in view["fields"] if f["kind"] == "html") == 2
        with pytest.raises(StaleReportEdit):
            save_field(conn, rid, report, narrative.concept, "<p>Human revision.</p>",
                       "user_override", "Correction", [1], expected_revision=1)
        save_field(conn, rid, report, narrative.concept, "<p>Human revision.</p>",
                   "user_override", "Correction", [1], expected_revision=2)
    assert cleanup_cells(str(db), rid, [sheet]) == {}
    with repo.db_session(db) as conn:
        conn.execute("UPDATE notes_cells SET concept_uuid='wrong-family' WHERE run_id=? AND sheet=?", (rid, sheet))
    assert candidate_sheets(str(db), rid, [sheet]) == []


def test_formatter_write_makes_earlier_report_edit_stale_and_unchanged_saves_keep_style(report_run):
    from reports.repository import report_view

    db, rid = report_run
    report = "STATEMENT_BY_DIRECTORS"
    narrative = next(f for f in fields(report) if f.kind == "html")
    sheet = REPORTS[report][2]
    with repo.db_session(db) as conn:
        save_field(conn, rid, report, narrative.concept, "<p>Statement.</p>", "observed", "Source", [1])
        assert repo.cas_update_notes_cell_html(conn, run_id=rid, sheet=sheet, row=narrative.row,
            expected_html="<p>Statement.</p>", new_html="<p><strong>Statement.</strong></p>",
            style_source="formatter")
        field = next(f for f in report_view(conn, rid, report)["fields"] if f["kind"] == "html")
        assert field["revision"] == 2
        with pytest.raises(StaleReportEdit):
            save_field(conn, rid, report, narrative.concept, "<p>Old text edit.</p>",
                       "user_override", "Correction", [1], expected_revision=1)
        # Re-saving identical text (new evidence only) keeps the verified style.
        save_field(conn, rid, report, narrative.concept, field["value"], "observed", "Recaptured", [1, 2])
        style = conn.execute("SELECT style_source FROM notes_cells WHERE run_id=? AND sheet=?",
                             (rid, sheet)).fetchone()[0]
        assert style == "formatter"


def test_overlay_keeps_updating_report_text_that_begins_with_equals(report_run, tmp_path):
    import shutil
    import openpyxl
    from notes.persistence import overlay_notes_cells_into_workbook

    db, rid = report_run
    report = "STATEMENT_BY_DIRECTORS"
    narrative = next(f for f in fields(report) if f.kind == "html")
    sheet = REPORTS[report][2]
    workbook = tmp_path / "merged.xlsx"
    shutil.copy(template_path(report, "mfrs", "company"), workbook)
    wb = openpyxl.load_workbook(workbook)
    wb[sheet]["C1"] = "=1+1"
    wb.save(workbook)
    for text in ("= RM1.2m restated", "= RM1.3m restated"):
        with repo.db_session(db) as conn:
            save_field(conn, rid, report, narrative.concept, f"<p>{text}</p>", "observed", "Source", [1])
        out = overlay_notes_cells_into_workbook(xlsx_path=workbook, run_id=rid, db_path=str(db))
        shutil.copy(out, workbook)
        ws = openpyxl.load_workbook(workbook)[sheet]
        assert ws.cell(narrative.row, 2).value == text
        assert ws["C1"].value == "=1+1"


@pytest.mark.parametrize("kind,value", [("date", "17/10/2026"), ("choice", "clean")])
def test_invalid_scalar_type_or_selection_never_reaches_canonical_values(report_run, kind, value):
    db, rid = report_run
    field = next(f for f in fields("AUDITORS_REPORT") if f.kind == kind)
    with repo.db_session(db) as conn:
        with pytest.raises(ValueError):
            save_field(conn, rid, "AUDITORS_REPORT", field.concept, value, "observed", "source", [1])
        assert conn.execute("SELECT count(*) FROM report_fields").fetchone()[0] == 0
    with pytest.raises(ValueError):
        format_value("0012")  # numeric patching keeps its original contract


@pytest.mark.asyncio
@pytest.mark.parametrize("report", list(REPORTS))
async def test_report_cleanup_preserves_content_and_report_edit_revisions(
    report_run, tmp_path, monkeypatch, report,
):
    from notes import cleanup_agent
    from notes.cleanup_patch import Patch, Removal, enumerate_blocks
    from reports.repository import report_view

    db, rid = report_run
    narrative = next(f for f in fields(report) if f.kind == "html")
    title = REPORTS[report][0]
    kept = (f"<h2>{title}</h2><h3>Report section</h3><p>Complete report wording.</p>"
            "<table><tr><td>2024</td><td>34,970</td></tr></table>"
            "<p>Signed by the director or auditor on 31 December 2024.</p>")
    original = "<h1>COMPANY SDN. BHD.</h1><p>Registration No. 123456</p>" + kept
    original += f"<h2>{title} (continued)</h2><p>Remaining report wording.</p>"
    with repo.db_session(db) as conn:
        save_field(conn, rid, report, narrative.concept, original, "observed", "Report pages", [1, 2])
        before = snapshot(conn, rid)["digest"]
    sheet = REPORTS[report][2]
    cid = f"{sheet}:{narrative.row}"
    pdf = tmp_path / "uploaded.pdf"
    pdf.write_bytes(b"%PDF")

    async def propose(**kwargs):
        assert set(kwargs["cells"]) == {cid}
        blocks = enumerate_blocks(kwargs["cells"][cid]["html"])[2]
        index = {b["text"]: b["block"] for b in blocks}
        removals = [Removal(
            cell=cid, block=index[text], expected_text=text, reason="page_banner",
            evidence_page=1, justification="Separate company identity page banner.",
        ) for text in ("COMPANY SDN. BHD.", "Registration No. 123456")]
        removals.append(Removal(
            cell=cid, block=index[f"{title} (continued)"], expected_text=f"{title} (continued)",
            reason="continuation_heading", retained_block=index[title], evidence_page=2,
            justification="Continuation repeats the retained report title.",
        ))
        return Patch(inspected_cells=[cid], removals=removals), {1, 2}

    monkeypatch.setattr(cleanup_agent, "propose_cleanup", propose)
    result = await cleanup_agent.run_notes_cleanup(
        run_id=rid, db_path=str(db), pdf_path=str(pdf), sheets=[sheet],
        model_name="scripted", model_factory=lambda: "scripted", output_dir=str(tmp_path),
        on_progress=lambda *args: None,
    )
    assert result["ok"] and result["removed_blocks"] == 3
    with repo.db_session(db) as conn:
        saved = next(f for f in report_view(conn, rid, report)["fields"] if f["kind"] == "html")
        assert saved["value"] == kept + "<p>Remaining report wording.</p>"
        assert saved["revision"] == 2
        assert saved["value_status"] == "observed" and saved["source_pages"] == [1, 2]
        assert snapshot(conn, rid)["digest"] != before
        receipt = conn.execute("SELECT before_html, after_html FROM notes_cleanup_receipts WHERE run_id=?", (rid,)).fetchone()
        assert receipt[:] == (original, saved["value"])
        exported = next(f for f in build_notes_fill_doc(db, rid)["footnotes"] if f["source_sheet"] == sheet)
        assert "COMPANY SDN. BHD." not in exported["html"]
        assert "Registration No." not in exported["html"]
        assert "Complete report wording." in exported["html"]
        assert "Remaining report wording." in exported["html"]
        with pytest.raises(StaleReportEdit):
            save_field(conn, rid, report, narrative.concept, original, "user_override", "Edit", [1], expected_revision=1)
        save_field(conn, rid, report, narrative.concept, original, "user_override", "Edit", [1], expected_revision=2)
    # Human report overrides use report_fields status rather than notes lineage.
    assert cleanup_agent.cleanup_cells(str(db), rid, [sheet]) == {}


def _native_workbook(path, report, report_fields, *, protected=False):
    wb = Workbook(); ws = wb.active; ws.title = REPORTS[report][2]
    for row, field in enumerate(report_fields, 1):
        ws.cell(row, 1, field.concept)
        ws.cell(row, 2, field.label)
        ws.cell(row, 3, "old")
        wb.defined_names.add(DefinedName(field.concept, attr_text=f"'{ws.title}'!$C${row}"))
        if field.kind == "choice":
            dv = DataValidation(type="list", formula1='"' + ','.join(field.choices) + '"')
            ws.add_data_validation(dv); dv.add(ws.cell(row, 3))
        if field.kind == "date":
            ws.cell(row, 3).number_format = "yyyy-mm-dd"
            dv = DataValidation(type="date", operator="between", formula1="1", formula2="100000")
            ws.add_data_validation(dv); dv.add(ws.cell(row, 3))
        ws.cell(row, 3).protection = Protection(locked=protected)
    ws["F1"] = "=1+2"
    ws.protection.sheet = protected
    wb.save(path); wb.close()


def test_typed_report_export_preserves_ids_selections_dates_and_formula(report_run, tmp_path):
    db, rid = report_run
    report = "AUDITORS_REPORT"
    selected = [next(f for f in fields(report) if f.concept.endswith("LicenseNumberOfAuditor")),
                next(f for f in fields(report) if f.kind == "date"), next(f for f in fields(report) if f.kind == "choice")]
    values = ["0012", "2026-10-17", selected[-1].choices[-1]]
    with repo.db_session(db) as conn:
        for f, value in zip(selected, values):
            save_field(conn, rid, report, f.concept, value, "observed", "Source", [1])
        snap = snapshot(conn, rid)
    src, out = tmp_path / "native.xlsx", tmp_path / "filled.xlsx"
    _native_workbook(src, report, selected)
    _, data, _ = load_workbook_entries(str(src))
    resolved = resolve_reports(data, snap)
    assert len(resolved["writes"]) == 3
    result = fill_report_fields(str(src), resolved, str(out))
    assert not result["mismatches"]
    assert any(f["reason"] == "assessment_incomplete" for f in result["unresolved"])
    wb = load_workbook(out)
    try:
        ws = wb[REPORTS[report][2]]
        assert ws["C1"].value == "0012"
        assert ws["C1"].data_type == "s"
        assert ws["C2"].value.date().isoformat() == "2026-10-17"
        assert ws["C3"].value == values[-1]
        assert ws["F1"].value == "=1+2"
    finally:
        wb.close()
    # An explicit human clear must remove a previously filled identifier.
    with repo.db_session(db) as conn:
        save_field(conn, rid, report, selected[0].concept, None, "cleared", "Cleared", [], expected_revision=1)
        cleared = snapshot(conn, rid)
    _, filled_data, _ = load_workbook_entries(str(out))
    cleared_path = tmp_path / "cleared.xlsx"
    cleared_result = fill_report_fields(str(out), resolve_reports(filled_data, cleared), str(cleared_path))
    assert not cleared_result["mismatches"]
    wb = load_workbook(cleared_path)
    try:
        assert wb[REPORTS[report][2]]["C1"].value in {None, ""}
    finally:
        wb.close()


def test_column_tagged_report_table_fills_its_single_typing_row(report_run, tmp_path):
    # Native auditor details tag each column once and take values in a typing
    # row lower in the table (observed in a real mTool AuditReport sheet).
    db, rid = report_run
    report = "AUDITORS_REPORT"
    firm, reg, state = ("ssmt_NameOfAuditFirm", "ssmt_RegistrationNumberOfAuditFirm",
                        "ssmt_StateWhereAuditFirmIsLocated")
    with repo.db_session(db) as conn:
        for concept, value in ((firm, "Example & Co."), (reg, "AF 0001"), (state, "Selangor")):
            save_field(conn, rid, report, concept, value, "observed", "Source", [1])
        snap = snapshot(conn, rid)

    def build(path, typing_rows):
        wb = Workbook()
        ws = wb.active
        ws.title = REPORTS[report][2]
        for col, concept in zip("EFG", (firm, reg, state)):
            ws[f"{col}3"] = f"ssmt-cor_2022-12-31.xsd#{concept}"
        for row in typing_rows:
            for col in "EF":  # State stays locked, as mTool's picker cells do.
                ws[f"{col}{row}"].protection = Protection(locked=False)
            ws[f"G{row}"] = None
        ws["C7"] = "#LAYOUTSCER#"
        ws["E9"].protection = Protection(locked=False)  # outside the table
        ws.protection.sheet = True
        wb.save(path); wb.close()

    src, out = tmp_path / "table.xlsx", tmp_path / "filled.xlsx"
    build(src, [6])
    _, data, _ = load_workbook_entries(str(src))
    resolved = resolve_reports(data, snap)
    assert {(w["concept"], w["cell"]) for w in resolved["writes"]} == {(firm, "E6"), (reg, "F6")}
    assert any(u.get("concept") == state and u["reason"] == "missing_or_ambiguous_report_destination"
               for u in resolved["unresolved"])
    result = fill_report_fields(str(src), resolved, str(out))
    assert not result["mismatches"]
    wb = load_workbook(out)
    try:
        ws = wb[REPORTS[report][2]]
        assert (ws["E6"].value, ws["F6"].value, ws["E9"].value) == ("Example & Co.", "AF0001", None)
    finally:
        wb.close()

    # Two typing rows (for example two auditors) never pick one.
    two = tmp_path / "two.xlsx"
    build(two, [5, 6])
    _, data, _ = load_workbook_entries(str(two))
    assert not resolve_reports(data, snap)["writes"]


def test_protected_or_unknown_report_destination_is_never_guessed(report_run, tmp_path):
    db, rid = report_run
    report = "AUDITORS_REPORT"
    field = next(f for f in fields(report) if f.concept == "ssmt_NameOfAuditorSigningReport")
    with repo.db_session(db) as conn:
        save_field(conn, rid, report, field.concept, "Source auditor", "observed", "Source", [1])
        snap = snapshot(conn, rid)
    src, out = tmp_path / "locked.xlsx", tmp_path / "filled.xlsx"
    _native_workbook(src, report, [field], protected=True)
    _, data, _ = load_workbook_entries(str(src))
    doc = resolve_reports(data, snap)
    result = fill_report_fields(str(src), doc, str(out))
    assert not result["written"]
    assert any(e["reason"] == "protected_or_missing_cell" for e in result["unresolved"])
    with pytest.raises(ValueError, match="overlaps"):
        resolve_reports(data, snap, occupied=[(REPORTS[report][2], "C1")])
    # Local and workbook names may share a taxonomy name. Both destinations
    # must remain visible to ambiguity detection rather than last-name wins.
    import xml.etree.ElementTree as ET
    root = ET.fromstring(data["xl/workbook.xml"])
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    defined = root.find(f"{ns}definedNames")
    ET.SubElement(defined, f"{ns}definedName", name=field.concept, localSheetId="0").text = f"'{REPORTS[report][2]}'!$D$1"
    data["xl/workbook.xml"] = ET.tostring(root)
    ambiguous = resolve_reports(data, snap)
    assert not ambiguous["writes"]
    assert any(e["reason"] == "missing_or_ambiguous_report_destination" for e in ambiguous["unresolved"])
    # Generic taxonomy concepts in filing information do not identify an
    # auditor-report destination just because a workbook name points at them.
    _, data, _ = load_workbook_entries(str(src))
    data["xl/workbook.xml"] = data["xl/workbook.xml"].replace(b"Report-Auditors", b"Filing-Information")
    unrelated = resolve_reports(data, snap)
    assert not unrelated["writes"]


@pytest.mark.parametrize("concept,value", [
    ("ssmt_LicenseNumberOfAuditor", "03339/12/2025 J"),
    ("ssmt_RegistrationNumberOfAuditFirm", "AF0001 & AF0002"),
    ("ssmt_RegistrationNumberOfAuditFirm", "AF0001 & AF00022"),
    ("ssmt_RegistrationNumberOfAuditFirm", "AF07589"),
    ("ssmt_RegistrationNumberOfAuditFirm", "LLP0010081-LCA"),
])
def test_report_identifiers_without_one_verified_native_encoding_stay_unwritten(report_run, tmp_path, concept, value):
    db, rid = report_run
    field = next(f for f in fields("AUDITORS_REPORT") if f.concept == concept)
    with repo.db_session(db) as conn:
        save_field(conn, rid, "AUDITORS_REPORT", concept, value, "observed", "Source identifier", [1])
        snap = snapshot(conn, rid)
    source, output = tmp_path / "identifiers.xlsx", tmp_path / "filled.xlsx"
    _native_workbook(source, "AUDITORS_REPORT", [field])
    _, data, _ = load_workbook_entries(str(source))
    resolved = resolve_reports(data, snap)
    assert not resolved["writes"]
    assert any(u.get("concept") == concept and u["reason"] == "native_identifier_encoding_unverified"
               for u in resolved["unresolved"])
    result = fill_report_fields(str(source), resolved, str(output))
    assert not result["written"]
    wb = load_workbook(output)
    try:
        assert wb["Report-Auditors"]["C1"].value == "old"
    finally:
        wb.close()
    with repo.db_session(db) as conn:
        saved = next(r for r in snapshot(conn, rid)["reports"] if r["type"] == "AUDITORS_REPORT")
        assert next(f["value"] for f in saved["fields"] if f["concept"] == concept) == value


def test_report_review_api_rejects_stale_edits_and_edits_during_extraction(report_run, monkeypatch):
    import server
    db, rid = report_run
    monkeypatch.setattr(server, "AUDIT_DB_PATH", db)
    client = TestClient(server.app)
    field = next(f for f in fields("AUDITORS_REPORT") if f.kind == "text")
    url = f"/api/runs/{rid}/reports/AUDITORS_REPORT/{field.concept}"
    response = client.patch(url, json={"value": "0012", "revision": 0})
    assert response.status_code == 200, response.text
    assert response.json()["value"] == "0012"
    assert client.patch(url, json={"value": "0001", "revision": 0}).status_code == 409
    assert client.patch(url, json={"value": None, "revision": 1}).json()["value"] is None
    with repo.db_session(db) as conn:
        repo.upsert_notes_format_task(conn, rid, REPORTS["AUDITORS_REPORT"][2], "running", model="test")
    assert client.patch(url, json={"value": "0012", "revision": 2}).status_code == 409
    with repo.db_session(db) as conn:
        repo.upsert_notes_format_task(conn, rid, REPORTS["AUDITORS_REPORT"][2], "done", model="test")
    with repo.db_session(db) as conn:
        conn.execute("UPDATE runs SET status='running' WHERE id=?", (rid,))
    assert client.patch(url, json={"value": "0012", "revision": 2}).status_code == 409
    assert client.get("/api/runs/99999/reports").status_code == 404


def test_report_only_workbook_preparation_exposes_unresolved_fields_and_stale_download(report_run, monkeypatch, tmp_path):
    import server
    db, rid = report_run
    monkeypatch.setattr(server, "AUDIT_DB_PATH", db)
    monkeypatch.setattr(server, "OUTPUT_DIR", tmp_path)
    report = "AUDITORS_REPORT"
    field = next(f for f in fields(report) if f.kind == "text")
    with repo.db_session(db) as conn:
        save_field(conn, rid, report, field.concept, "0012", "observed", "Source", [1])
    src = tmp_path / "native.xlsx"
    _native_workbook(src, report, [field])
    client = TestClient(server.app)
    response = client.post(f"/api/runs/{rid}/mtool-fill/patch", files={"template": (src.name, src.read_bytes())})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["status"] == "degraded"
    assert result["reports"]["written"][0]["value"] == "0012"
    assert result["reports"]["unresolved"]
    with repo.db_session(db) as conn:
        save_field(conn, rid, report, field.concept, "0013", "user_override", "Correction", [], expected_revision=1)
    download = client.get(result["download_url"], params={"acknowledge_degraded": "Reviewed"})
    assert download.status_code == 409
    assert "Reports changed" in download.text


@pytest.mark.parametrize("cleared", [False, True])
def test_reports_export_narratives_independently_of_financial_notes(report_run, monkeypatch, tmp_path, cleared):
    from test_mtool_offline_fill import footnote_template
    from zipfile import ZipFile
    import server
    db, rid = report_run
    monkeypatch.setattr(server, "AUDIT_DB_PATH", db)
    monkeypatch.setattr(server, "OUTPUT_DIR", tmp_path)
    report = "AUDITORS_REPORT"
    narrative = next(f for f in fields(report) if f.kind == "html")
    with repo.db_session(db) as conn:
        save_field(conn, rid, report, narrative.concept, "<p>Opinion and complete report.</p>", "observed", "Opinion", [1])
        if cleared:
            save_field(conn, rid, report, narrative.concept, None, "cleared", "Cleared", [], expected_revision=1)
    source = footnote_template.__wrapped__(tmp_path)
    renamed = tmp_path / "report-popup.xlsx"
    with ZipFile(source) as original, ZipFile(renamed, "w") as out:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename.endswith(".xml"):
                data = data.replace(b"Notes-Listofnotes", b"Report-Auditors")
                data = data.replace(b"Property, plant and equipment", narrative.label.encode())
                if item.filename == "xl/worksheets/sheet1.xml":
                    marker = f'<c r="A132" t="inlineStr"><is><t>fields.xsd#{narrative.concept}</t></is></c>'.encode()
                    data = data.replace(b'<row r="132">', b'<row r="132">' + marker)
            out.writestr(item, data)
    client = TestClient(server.app)
    response = client.post(f"/api/runs/{rid}/mtool-fill/patch", files={"template": (renamed.name, renamed.read_bytes())}, data={"fill_notes": "false"})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["notes"]["counts"]["written"] == 1, result["notes"]
    download = client.get(result["download_url"], params={"acknowledge_degraded": "Reviewed"})
    assert download.status_code == 200
    filled = tmp_path / "downloaded.xlsx"; filled.write_bytes(download.content)
    from mtool.offline_fill import get_shared_strings
    _, data, _ = load_workbook_entries(str(filled))
    texts = get_shared_strings(data)
    assert not any("Existing PPE note" in text for text in texts)
    if cleared:
        assert not any("Opinion and complete report." in text for text in texts)
        assert not any(e["reason"] == "report_narrative_missing" and e.get("sheet") == REPORTS[report][2]
                       for e in result["reports"]["unresolved"])
    else:
        assert any("Opinion and complete report." in text for text in texts)


@pytest.mark.parametrize("kind,value", [("text", "0012"), ("text", "=1+1"), ("html", "=1+1"),
                                       ("text", "#REF!"), ("html", "#REF!")])
def test_report_only_diagnostic_download_rebuilds_current_canonical_edits(report_run, monkeypatch, tmp_path, kind, value):
    import server
    db, rid = report_run
    monkeypatch.setattr(server, "AUDIT_DB_PATH", db)
    field = next(f for f in fields("AUDITORS_REPORT") if f.kind == kind)
    with repo.db_session(db) as conn:
        conn.execute("UPDATE runs SET merged_workbook_path=? WHERE id=?", (str(tmp_path / "merged.xlsx"), rid))
        save_field(conn, rid, "AUDITORS_REPORT", field.concept, f"<p>{value}</p>" if kind == "html" else value,
                   "user_override", "Correction", [])
    rebuilt = server._reexport_and_remerge_from_facts(rid)
    assert rebuilt is not None
    try:
        wb = load_workbook(rebuilt)
        try:
            cell = wb[REPORTS["AUDITORS_REPORT"][2]].cell(field.row, 2)
            assert cell.value == value
            assert cell.data_type == "s"
        finally:
            wb.close()
    finally:
        rebuilt.unlink(missing_ok=True)


@pytest.mark.asyncio
@pytest.mark.parametrize("crashed", [False, True])
async def test_report_worker_distinguishes_incomplete_assessment_from_crash(report_run, monkeypatch, tmp_path, crashed):
    import asyncio
    from types import SimpleNamespace
    from pydantic_ai.models.test import TestModel
    from reports import agent

    db, rid = report_run
    report = "AUDITORS_REPORT"
    with repo.db_session(db) as conn:
        rai = repo.create_run_agent(conn, rid, statement_type=report)

    def page_count(_path):
        if crashed:
            raise RuntimeError("PDF read failed")
        return 1

    monkeypatch.setattr(agent, "count_pdf_pages", page_count)
    config = SimpleNamespace(db_path=str(db), run_id=rid, pdf_path="unused.pdf", output_dir=str(tmp_path),
                             model=TestModel(call_tools=[], custom_output_text="Assessment ended."), filing_standard="mfrs")
    results = await agent.run_reports(config, {report}, {}, {report: rai}, asyncio.Queue(), "report-outcome-test")
    assert results[0].status == "failed"
    with repo.db_session(db) as conn:
        saved = conn.execute("SELECT status,error_type,error_message,workbook_path FROM run_agents WHERE id=?", (rai,)).fetchone()
    assert saved[0] == "failed"
    assert saved[1] == (None if crashed else "report_incomplete")
    assert saved[2] == ("Report extraction failed. Captured values are retained for review." if crashed
                        else "Report assessment incomplete. Review missing or uncertain fields.")
    assert bool(saved[3]) is not crashed


@pytest.mark.asyncio
async def test_report_tasks_cancelled_before_start_still_finish_audit_rows(report_run, monkeypatch, tmp_path):
    import asyncio
    from types import SimpleNamespace
    from reports.agent import run_reports
    import task_registry
    db, rid = report_run
    with repo.db_session(db) as conn:
        ids = {r: repo.create_run_agent(conn, rid, statement_type=r) for r in REPORTS}
    monkeypatch.setattr(task_registry, "register", lambda session, role, task: task.cancel(task_registry.USER_ABORT_REASON))
    config = SimpleNamespace(db_path=str(db), run_id=rid, pdf_path="unused.pdf", output_dir=str(tmp_path), model="unused", filing_standard="mfrs")
    results = await run_reports(config, set(REPORTS), {}, ids, asyncio.Queue(), "cancelled-report-test")
    assert [r.status for r in results] == ["cancelled"] * len(REPORTS)
    with repo.db_session(db) as conn:
        assert {r[0] for r in conn.execute("SELECT status FROM run_agents WHERE run_id=?", (rid,))} == {"cancelled"}


@pytest.mark.asyncio
@pytest.mark.parametrize("outcome", ["recovered", "exhausted", "cancelled"])
async def test_report_provider_retry_retains_values_usage_and_cancellation(report_run, monkeypatch, tmp_path, outcome):
    import asyncio
    import httpx
    from contextlib import asynccontextmanager
    from types import SimpleNamespace
    from openai import APIError
    from reports import agent
    from reports.repository import report_view
    import task_registry

    db, rid = report_run
    report = "AUDITORS_REPORT"
    identifier = next(f for f in fields(report) if f.kind == "text")
    with repo.db_session(db) as conn:
        rai = repo.create_run_agent(conn, rid, statement_type=report)
    deps_seen, events = [], []

    def create(_model, deps):
        deps_seen.append(deps)
        @asynccontextmanager
        async def iterate(*args, **kwargs):
            yield SimpleNamespace(ctx=SimpleNamespace(state=SimpleNamespace(message_history=[])),
                usage=SimpleNamespace(input_tokens=7, output_tokens=5, total_tokens=12,
                    cache_read_tokens=2, cache_write_tokens=0, details={"reasoning_tokens": 1}))
        return SimpleNamespace(iter=iterate)

    async def loop(_run, deps, _spec, _emit, turns):
        assert deps.viewed == set()  # A retry must inspect its own evidence.
        deps.viewed.add(1)
        turns.append({"prompt_tokens": 3, "completion_tokens": 2, "total_tokens": 5,
                      "_n_tool_calls": 1})
        with repo.db_session(db) as conn:
            if len(deps_seen) == 1:
                save_field(conn, rid, report, identifier.concept, "0012", "observed", "Licence", [1])
            else:
                view = report_view(conn, rid, report)
                assert next(f["value"] for f in view["fields"] if f["concept"] == identifier.concept) == "0012"
                if outcome == "recovered":
                    for f in fields(report):
                        if f.concept == identifier.concept:
                            continue
                        save_field(conn, rid, report, f.concept, "<p>Complete report.</p>" if f.kind == "html" else None,
                                   "observed" if f.kind == "html" else "not_disclosed", "Source assessed", [1])
        if len(deps_seen) == 1 or outcome == "exhausted":
            raise APIError("Our servers are currently overloaded. Please try again later.",
                           request=httpx.Request("POST", "https://provider.invalid/v1/responses"), body=None)

    async def emit(kind, payload):
        events.append((kind, payload))
        if outcome == "cancelled" and payload.get("phase") == "retrying":
            raise asyncio.CancelledError(task_registry.USER_ABORT_REASON)

    monkeypatch.setattr(agent, "create_report_agent", create)
    monkeypatch.setattr(agent, "run_agent_loop", loop)
    monkeypatch.setattr(agent, "count_pdf_pages", lambda _: 1)
    monkeypatch.setattr(agent, "make_emitter", lambda *args: (emit, None))
    monkeypatch.setattr("notes._rate_limit.compute_backoff_delay", lambda *_: 0)
    config = SimpleNamespace(db_path=str(db), run_id=rid, pdf_path="unused.pdf", output_dir=str(tmp_path),
                             model="test-model", filing_standard="mfrs")
    results = await agent.run_reports(config, {report}, {}, {report: rai}, asyncio.Queue(), "report-retry-test")
    attempts = 1 if outcome == "cancelled" else 2
    expected_status = {"recovered": "succeeded", "exhausted": "failed", "cancelled": "cancelled"}[outcome]
    assert results[0].status == expected_status
    assert len(deps_seen) == attempts
    assert sum(payload.get("phase") == "retrying" for _, payload in events) == 1
    assert sum(kind == "complete" for kind, _ in events) == 1
    with repo.db_session(db) as conn:
        saved = conn.execute("SELECT status,error_type,total_tokens,prompt_tokens,completion_tokens,reasoning_tokens,turn_count,cache_read_tokens FROM run_agents WHERE id=?", (rai,)).fetchone()
        assert tuple(saved) == (expected_status, {"recovered": None, "exhausted": "transient_exhausted", "cancelled": "cancelled"}[outcome],
                                12 * attempts, 7 * attempts, 4 * attempts, attempts, attempts, 2 * attempts)
        assert conn.execute("SELECT count(*) FROM run_agent_turns WHERE run_agent_id=?", (rai,)).fetchone()[0] == attempts
        view = report_view(conn, rid, report)
        assert next(f["value"] for f in view["fields"] if f["concept"] == identifier.concept) == "0012"


@pytest.mark.asyncio
async def test_report_tool_rejects_unknown_batch_with_exact_catalogue_and_recovers(report_run):
    from pydantic_ai.models.function import FunctionModel
    from pydantic_ai.messages import ModelResponse, RetryPromptPart, TextPart, ToolCallPart
    from reports.agent import ReportDeps, create_report_agent
    from reports.repository import report_view

    db, rid = report_run
    report = "STATEMENT_BY_DIRECTORS"
    narrative = "ssmt_DisclosureOfStatementByDirectorsExplanatory"
    prefix = "ssmt_DisclosureWhether"
    correct = [f"{prefix}{slot}DirectorIsAlsoPrimarilyResponsibleForFinancialManagementOfCompany"
               for slot in ("Fourth", "Fifth")]
    invalid = [concept.replace("OfCompany", "OfTheCompany") for concept in correct]
    permitted = [f.concept for f in fields(report)]
    step = 0

    def respond(messages, info):
        nonlocal step
        tool = next(t for t in info.function_tools if t.name == "save_report_fields")
        enum = tool.parameters_json_schema["$defs"]["CapturedField"]["properties"]["concept"]["enum"]
        assert enum == permitted
        if step in {1, 2}:
            retry = next(p for p in messages[-1].parts if isinstance(p, RetryPromptPart))
            feedback = str(retry.content)
            for concept in invalid[step - 1:]:
                assert concept in feedback
            for concept in permitted:
                assert concept in feedback
            assert "Nothing in this batch was saved" in feedback
            with repo.db_session(db) as conn:
                assert conn.execute("SELECT count(*) FROM report_fields WHERE run_id=?", (rid,)).fetchone()[0] == 0
                assert conn.execute("SELECT count(*) FROM notes_cells WHERE run_id=?", (rid,)).fetchone()[0] == 0
        if step == 3:
            return ModelResponse(parts=[TextPart("Report fields saved after correcting taxonomy IDs.")])
        concepts = invalid if step == 0 else [correct[0], invalid[1]] if step == 1 else correct
        step += 1
        return ModelResponse(parts=[ToolCallPart("save_report_fields", {"captured": [
            {"concept": narrative, "value": "<p>Complete statement.</p>", "status": "observed",
             "evidence": "Statement source", "pages": [1]},
            *[{"concept": c, "status": "not_disclosed", "evidence": "No additional signer", "pages": []}
              for c in concepts],
        ]})])

    deps = ReportDeps("unused.pdf", str(db), rid, report, "mfrs", 1, viewed={1})
    agent = create_report_agent(FunctionModel(respond), deps)
    await agent.run("Capture the statement.", deps=deps)
    assert step == 3
    with repo.db_session(db) as conn:
        saved = report_view(conn, rid, report)
        by_concept = {f["concept"]: f for f in saved["fields"]}
        assert by_concept[narrative]["value"] == "<p>Complete statement.</p>"
        assert all(by_concept[c]["value_status"] == "not_disclosed" for c in correct)
        assert conn.execute("SELECT count(*) FROM report_fields WHERE run_id=?", (rid,)).fetchone()[0] == 3
