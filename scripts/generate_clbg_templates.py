"""Generate CLBG Company templates from the committed SSM presentation/calculation linkbases."""
from __future__ import annotations

import argparse
from functools import lru_cache
from pathlib import Path
import shutil
import sys
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts import generate_mpers_templates as common

TAXONOMY = ROOT / "SSMxT_2022v1.0/rep/ssm/ca-2016/fs/clbg"
TEMPLATES = {
    "01-SOFP-CuNonCu.xlsx": (("SOFP-CuNonCu", "210000", "210000"), ("SOFP-Sub-CuNonCu", "210100", "210100")),
    "03-SOPL-Function.xlsx": (("SOIE-Function", "310000", "300100"), ("SOIE-Analysis-Function", "310100", "300200")),
    "04-SOPL-Nature.xlsx": (("SOIE-Nature", "320000", "300100"), ("SOIE-Analysis-Nature", "320100", "300200")),
    "07-SOCF-Indirect.xlsx": (("SOCF-Indirect", "510000", "510000"),),
    "09-SOCIE.xlsx": (("SOCIE", "410000", "410000"),),
    "10-Notes-CorporateInfo.xlsx": (("Notes-CI", "610000", None),),
    "11-Notes-AccountingPolicies.xlsx": (("Notes-SummaryofAccPol", "620000", None),),
    "12-Notes-ListOfNotes.xlsx": (("Notes-Listofnotes", "630000", None),),
    "14-Notes-RelatedParty.xlsx": (("Notes-RelatedPartytran", "640000", None),),
}
FUND_AXIS = "ssmt-mfrs_ComponentsOfFundAxis"


@lru_cache(maxsize=1)
def labels():
    from scripts.generate_concept_units import load_role_labels, _parse_label_linkbase
    result = load_role_labels()
    for identifier, roles in _parse_label_linkbase(TAXONOMY / "lab_en-ssmt-fs-clbg_2022-12-31.xml").items():
        result.setdefault(identifier, {}).update(roles)
    return result


def display_label(identifier):
    roles = labels().get(identifier, {})
    value = roles.get("ReportingLabel") or roles.get(common._STANDARD_LABEL_ROLE)
    if not value:
        raise ValueError(f"Missing taxonomy label for {identifier}")
    return common._strip_display_suffix(value)


@lru_cache(maxsize=16)
def role_rows(role):
    from concept_model.taxonomy_semantics import taxonomy_registry
    registry = taxonomy_registry()
    out = []
    xlink = "{http://www.w3.org/1999/xlink}"
    tree = ET.parse(TAXONOMY / f"pre_ssmt-fs-clbg_2022-12-31_role-{role}.xml")
    def visit(locator, depth, preferred, locs, edges):
        identifier = locs[locator]
        concept = registry.get(identifier)
        if concept is None:
            raise ValueError(f"Missing taxonomy declaration for {identifier}")
        if concept.concept_role in {"PRIMARY_ITEM", "ABSTRACT"} and not identifier.endswith(("Table", "Axis", "Member", "LineItems")):
            roles = labels().get(identifier, {})
            label = roles.get(preferred) or roles.get(common._normalise_label_role(preferred or "")) or display_label(identifier)
            out.append((depth, identifier, common._strip_display_suffix(label), concept.abstract))
        for _, target, child_preferred in sorted(edges.get(locator, []), key=lambda item: item[0]):
            visit(target, depth+1, child_preferred, locs, edges)
    for link in tree.getroot().findall("{http://www.xbrl.org/2003/linkbase}presentationLink"):
        locs = {n.get(xlink+"label"): n.get(xlink+"href").split("#")[-1]
                for n in link if n.tag.endswith("}loc")}
        edges, children = {}, set()
        for arc in link:
            if arc.tag.endswith("}presentationArc"):
                source, target = (arc.get(xlink+k) for k in ("from", "to"))
                edges.setdefault(source, []).append((float(arc.get("order", "0")), target, arc.get("preferredLabel")))
                children.add(target)
        for root in locs:
            if root not in children:
                visit(root, 0, None, locs, edges)
    return tuple(out)


@lru_cache(maxsize=1)
def fund_components():
    """Presentation-order leaves and subtotals, each with its immediate children."""
    xlink = "{http://www.w3.org/1999/xlink}"
    tree = ET.parse(TAXONOMY / "pre_ssmt-fs-clbg_2022-12-31_role-410000.xml")
    link, = tree.getroot().findall("{http://www.xbrl.org/2003/linkbase}presentationLink")
    locs = {n.get(xlink+"label"): n.get(xlink+"href").split("#")[-1]
            for n in link if n.tag.endswith("}loc")}
    edges = {}
    for arc in link:
        if arc.tag.endswith("}presentationArc"):
            source, target = (locs[arc.get(xlink+k)] for k in ("from", "to"))
            edges.setdefault(source, []).append((float(arc.get("order", "0")), target))
    out = []
    def visit(member):
        children = [v for _, v in sorted(edges.get(member, []))]
        for child in children:
            visit(child)
        out.append((member, display_label(member), tuple(children)))
    for _, root in sorted(edges[FUND_AXIS]):
        visit(root)
    return tuple(out)


def fund_line_items():
    rows = list(role_rows("410000"))
    start = next(i for i, r in enumerate(rows) if r[1] == "ssmt-mfrs_FundBalance")
    return rows[start:]


def build_template(filename, out_dir):
    import openpyxl
    from openpyxl.styles import Font
    from utils.workbook_io import atomic_save_workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    emitted = []
    for sheet, role, calc_role in TEMPLATES[filename]:
        ws = wb.create_sheet(sheet)
        rows = list(role_rows(role))
        calc = common.parse_calc_linkbase_grouped(TAXONOMY / f"cal_ssmt-fs-clbg_2022-12-31_role-{calc_role}.xml") if calc_role else []
        if filename == "09-SOCIE.xlsx":
            rows = fund_line_items()
            components = fund_components()
            columns = tuple(openpyxl.utils.get_column_letter(c) for c in range(2, len(components)+2))
            ws.cell(1, len(components)+2, "Source")
            for c, (_, label, _) in enumerate(components, 2):
                ws.cell(2, c, label)
            for block, period in enumerate(("CY", "PY")):
                base = 6 + block*(len(rows)+3)
                ws.cell(base-1, 1, f"Company {period}")
                for i, (_, _, label, abstract) in enumerate(rows):
                    cell = ws.cell(base+i, 1, label)
                    if abstract:
                        common._apply_abstract_row_styling(cell)
                common._inject_sum_formulas(ws, rows, calc, columns, base_row=base)
                member_cols = {member: col for (member, _, _), col in zip(components, columns)}
                for i, (_, _, _, abstract) in enumerate(rows):
                    if abstract:
                        continue
                    for (member, _, children), col in zip(components, columns):
                        if children:
                            ws[f"{col}{base+i}"] = "="+"+".join(f"{member_cols[child]}{base+i}" for child in children)
            ws.freeze_panes = "B6"
            ws.column_dimensions["A"].width = 55
            for col in columns:
                ws.column_dimensions[col].width = 18
        else:
            common._apply_company_sheet_layout(ws, rows)
            common._inject_sum_formulas(ws, rows, calc)
        emitted.append((ws, rows))
    if len(emitted) == 2:
        common._inject_face_to_sub_rollups(emitted[0][0], emitted[0][1], emitted[1][0], emitted[1][1], value_columns=("B", "C"))
    path = Path(out_dir)/filename
    path.parent.mkdir(parents=True, exist_ok=True)
    atomic_save_workbook(wb, str(path))
    wb.close()
    return path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--snapshot", type=Path, required=True, help="Before/after audit snapshot directory; existing files are never overwritten")
    args = parser.parse_args()
    out = ROOT/"XBRL-template-CLBG/Company"
    for name in TEMPLATES:
        before = args.snapshot/"before"/name
        after = args.snapshot/"after"/name
        if before.exists() or after.exists():
            raise FileExistsError(f"Snapshot already exists for {name}")
        if (out/name).exists():
            before.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(out/name, before)
        built = build_template(name, out)
        after.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(built, after)
        print(built.relative_to(ROOT))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
