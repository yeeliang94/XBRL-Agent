"""Generate report templates from SSM presentation roles; no formulas."""
from __future__ import annotations

import argparse
from pathlib import Path
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from scripts.generate_mpers_templates import walk_role_for_taxonomy, _apply_company_sheet_layout
from reports.registry import REPORTS, template_path
from utils.workbook_io import atomic_save_workbook


def generate(snapshot: Path) -> None:
    snapshot.mkdir(parents=True, exist_ok=False)
    for standard in ("mfrs", "mpers"):
        for report, (_, role, sheet) in REPORTS.items():
            rows = walk_role_for_taxonomy(ROOT / f"SSMxT_2022v1.0/rep/ssm/ca-2016/fs/{standard}", standard, role)
            for level in ("company", "group"):
                path = template_path(report, standard, level)
                relative = path.relative_to(ROOT)
                if path.exists():
                    before = snapshot / "before" / relative
                    before.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copyfile(path, before)
                wb = Workbook()
                ws = wb.active
                ws.title = sheet
                _apply_company_sheet_layout(ws, rows)
                ws["B1"] = "Report value"
                ws["C1"] = ""
                ws["D1"] = "Source"
                atomic_save_workbook(wb, str(path))
                wb.close()
                after = snapshot / "after" / relative
                after.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(path, after)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--snapshot", type=Path, required=True)
    generate(parser.parse_args().snapshot)
