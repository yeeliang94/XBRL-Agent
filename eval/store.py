"""Benchmark library store — CRUD + the gold-fact grid.

The eval subsystem owns its own DB writes (the repository layer only persists
the per-run scorecard). This module is the single place that knows the
``eval_benchmarks`` / ``eval_benchmark_templates`` / ``gold_concept_facts``
shape, so the API router (``api/eval.py``) stays a thin HTTP shell.

The grid builder (:func:`benchmark_concepts`) returns the SAME view-row shape
as ``concept_model/concepts_routes.py``'s ``/api/runs/{id}/concepts`` so the
frontend ``ConceptsPage`` can render gold facts with zero new grid code (just a
``source`` prop — see docs/PLAN-eval-benchmark.md Step 9).
"""
from __future__ import annotations

import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import openpyxl

from statement_types import VARIANTS, template_path
from concept_model.parser import _derive_template_id
from eval.ingest import ingest_workbook, IngestResult


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds").replace(
        "+00:00", "Z"
    )


def _candidate_templates(standard: str, level: str) -> list[tuple[str, str]]:
    """Every ``(template_id, statement_value)`` registered for a
    ``(standard, level)`` whose file exists on disk.

    Derived from the variant registry so the statement_type is authoritative
    (a slug like ``sore`` maps to statement ``SOCIE``). Used to auto-detect a
    benchmark's template set from an uploaded workbook.
    """
    out: list[tuple[str, str]] = []
    for (statement, variant_name) in VARIANTS:
        try:
            path = template_path(statement, variant_name, level, standard)
        except ValueError:
            continue  # NotPrepared / standard mismatch — no template
        if not path.exists():
            continue
        out.append((_derive_template_id(path), statement.value))
    return out


def _template_sheets(conn: sqlite3.Connection, template_id: str) -> set[str]:
    """Distinct render + alias sheet names a template occupies."""
    sheets = {
        r[0]
        for r in conn.execute(
            "SELECT DISTINCT render_sheet FROM concept_nodes WHERE template_id = ?",
            (template_id,),
        ).fetchall()
    }
    sheets |= {
        r[0]
        for r in conn.execute(
            "SELECT DISTINCT a.alias_sheet FROM concept_render_aliases a "
            "JOIN concept_nodes n ON n.concept_uuid = a.concept_uuid "
            "WHERE n.template_id = ?",
            (template_id,),
        ).fetchall()
    }
    return sheets


def resolve_template_set(
    conn: sqlite3.Connection,
    standard: str,
    level: str,
    xlsx_path: str | Path,
) -> list[tuple[str, str]]:
    """Auto-detect which ``(standard, level)`` templates a workbook covers.

    Returns the ``(template_id, statement_value)`` pairs whose sheets appear in
    the uploaded workbook — the benchmark's explicit template set. Empty when
    no template matches (the caller rejects that loudly).
    """
    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True)
    try:
        wb_sheets = set(wb.sheetnames)
    finally:
        wb.close()
    matched: list[tuple[str, str]] = []
    for template_id, statement in _candidate_templates(standard, level):
        if _template_sheets(conn, template_id) & wb_sheets:
            matched.append((template_id, statement))
    return matched


def create_benchmark_from_workbook(
    conn: sqlite3.Connection,
    *,
    name: str,
    document: Optional[str],
    filing_standard: str,
    filing_level: str,
    xlsx_path: str | Path,
) -> dict:
    """Create a benchmark, record its template set, and ingest gold facts.

    Auto-detects the template set from the workbook's sheets. Raises
    ``ValueError`` when no template matches (wrong file / wrong standard or
    level). The caller owns the commit so the whole create+ingest is one
    transaction — a failed ingest rolls back the half-made benchmark.
    """
    template_set = resolve_template_set(
        conn, filing_standard, filing_level, xlsx_path
    )
    if not template_set:
        raise ValueError(
            f"No {filing_standard.upper()} {filing_level} template matched the "
            "uploaded workbook's sheets. Check the filing standard / level and "
            "that this is an MBRS template workbook."
        )

    cur = conn.execute(
        "INSERT INTO eval_benchmarks(name, document, filing_standard, "
        "filing_level, created_at) VALUES (?, ?, ?, ?, ?)",
        (name, document, filing_standard, filing_level, _now()),
    )
    benchmark_id = int(cur.lastrowid)
    for template_id, statement in template_set:
        conn.execute(
            "INSERT INTO eval_benchmark_templates(benchmark_id, template_id, "
            "statement_type) VALUES (?, ?, ?)",
            (benchmark_id, template_id, statement),
        )

    ingest: IngestResult = ingest_workbook(
        conn, benchmark_id, xlsx_path, [t for t, _ in template_set]
    )
    # Reject a benchmark that matched sheets but yielded no gold cells — it's a
    # useless, selectable artifact that would later grade to a meaningless 0/0.
    # The caller rolls back on this ValueError, discarding the half-made rows.
    if ingest.ingested == 0:
        raise ValueError(
            "No numeric gold cells found in the matched worksheets "
            f"({', '.join(ingest.matched_sheets) or 'none'}). The workbook's "
            "value columns appear empty — fill the gold figures and re-upload."
        )
    return {
        "id": benchmark_id,
        "ingested": ingest.ingested,
        "skipped": ingest.skipped,
        "matched_sheets": ingest.matched_sheets,
        "statements": sorted({s for _, s in template_set}),
        "template_ids": [t for t, _ in template_set],
    }


def benchmark_template_ids(conn: sqlite3.Connection, benchmark_id: int) -> list[str]:
    return [
        r[0]
        for r in conn.execute(
            "SELECT template_id FROM eval_benchmark_templates WHERE benchmark_id = ?",
            (benchmark_id,),
        ).fetchall()
    ]


def list_benchmarks(conn: sqlite3.Connection) -> list[dict]:
    """Every benchmark with its template/statement set + gold cell count."""
    prior = conn.row_factory
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            "SELECT id, name, document, filing_standard, filing_level, "
            "created_at FROM eval_benchmarks ORDER BY created_at DESC, id DESC"
        ).fetchall()
        out: list[dict] = []
        for r in rows:
            bid = r["id"]
            statements = [
                t[0]
                for t in conn.execute(
                    "SELECT DISTINCT statement_type FROM eval_benchmark_templates "
                    "WHERE benchmark_id = ? ORDER BY statement_type",
                    (bid,),
                ).fetchall()
            ]
            gold_count = conn.execute(
                "SELECT COUNT(*) FROM gold_concept_facts WHERE benchmark_id = ?",
                (bid,),
            ).fetchone()[0]
            out.append({
                "id": bid,
                "name": r["name"],
                "document": r["document"],
                "filing_standard": r["filing_standard"],
                "filing_level": r["filing_level"],
                "created_at": r["created_at"],
                "statements": statements,
                "gold_cell_count": int(gold_count),
            })
        return out
    finally:
        conn.row_factory = prior


def get_benchmark(conn: sqlite3.Connection, benchmark_id: int) -> Optional[dict]:
    prior = conn.row_factory
    conn.row_factory = sqlite3.Row
    try:
        r = conn.execute(
            "SELECT id, name, document, filing_standard, filing_level, "
            "created_at FROM eval_benchmarks WHERE id = ?",
            (benchmark_id,),
        ).fetchone()
        if r is None:
            return None
        templates = [
            {"template_id": t["template_id"], "statement_type": t["statement_type"]}
            for t in conn.execute(
                "SELECT template_id, statement_type FROM eval_benchmark_templates "
                "WHERE benchmark_id = ? ORDER BY statement_type",
                (benchmark_id,),
            ).fetchall()
        ]
        gold_count = conn.execute(
            "SELECT COUNT(*) FROM gold_concept_facts WHERE benchmark_id = ?",
            (benchmark_id,),
        ).fetchone()[0]
        return {
            "id": r["id"],
            "name": r["name"],
            "document": r["document"],
            "filing_standard": r["filing_standard"],
            "filing_level": r["filing_level"],
            "created_at": r["created_at"],
            "templates": templates,
            "statements": sorted({t["statement_type"] for t in templates}),
            "gold_cell_count": int(gold_count),
        }
    finally:
        conn.row_factory = prior


def delete_benchmark(conn: sqlite3.Connection, benchmark_id: int) -> bool:
    """Hard-delete a benchmark; CASCADE sweeps its templates + gold + scores.
    Also nulls ``runs.benchmark_id`` on any run that graded against it (the
    runs FK is ON DELETE SET NULL)."""
    conn.execute("PRAGMA foreign_keys = ON")
    cur = conn.execute("DELETE FROM eval_benchmarks WHERE id = ?", (benchmark_id,))
    return cur.rowcount > 0


def benchmark_concepts(conn: sqlite3.Connection, benchmark_id: int) -> list[dict]:
    """Concept grid for a benchmark, with gold facts embedded.

    Mirrors ``/api/runs/{id}/concepts``'s view-row shape (so ``ConceptsPage``
    renders it unchanged) but joins ``gold_concept_facts`` instead of
    ``run_concept_facts``. Only LEAF / MATRIX_CELL leaves are editable; alias
    rows are read-only (the workbook's cross-sheet formula owns the value).
    """
    template_ids = benchmark_template_ids(conn, benchmark_id)
    if not template_ids:
        return []
    placeholders = ",".join("?" for _ in template_ids)
    prior = conn.row_factory
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            f"""
            SELECT n.concept_uuid, n.parent_uuid, n.kind,
                   n.canonical_label, n.display_label,
                   n.render_sheet, n.render_row, n.render_col,
                   n.matrix_col, n.matrix_col_label, n.template_id, tpl.shape,
                   g.value, g.value_status, g.source,
                   (SELECT COUNT(*) FROM concept_edges e
                      WHERE e.parent_uuid = n.concept_uuid) AS edge_count
            FROM concept_nodes n
            JOIN concept_templates tpl ON tpl.template_id = n.template_id
            LEFT JOIN gold_concept_facts g
              ON g.concept_uuid = n.concept_uuid
              AND g.benchmark_id = ?
              AND g.period = 'CY'
              AND g.entity_scope = 'Company'
            WHERE n.template_id IN ({placeholders})
            ORDER BY n.template_id, n.render_sheet, n.render_row, n.render_col
            """,
            (benchmark_id, *template_ids),
        ).fetchall()

        all_facts = conn.execute(
            "SELECT concept_uuid, period, entity_scope, value "
            "FROM gold_concept_facts WHERE benchmark_id = ?",
            (benchmark_id,),
        ).fetchall()
        scope_facts: dict[str, dict] = {}
        for f in all_facts:
            bucket = scope_facts.setdefault(f["concept_uuid"], {})
            bucket.setdefault(f["entity_scope"], {})[f["period"]] = f["value"]

        out: list[dict] = []
        for r in rows:
            out.append({
                "concept_uuid": r["concept_uuid"],
                "parent_uuid": r["parent_uuid"],
                "kind": r["kind"],
                "canonical_label": r["canonical_label"],
                "display_label": r["display_label"],
                "render_sheet": r["render_sheet"],
                "render_row": r["render_row"],
                "render_col": r["render_col"],
                "matrix_col": r["matrix_col"],
                "matrix_col_label": r["matrix_col_label"],
                "shape": r["shape"],
                "template_id": r["template_id"],
                "value": r["value"],
                "value_status": r["value_status"],
                "children_status": None,
                "source": r["source"],
                "evidence": None,
                "editable": (
                    r["kind"] in ("LEAF", "MATRIX_CELL")
                    and (r["edge_count"] or 0) == 0
                ),
                "is_alias": False,
                "scope_facts": scope_facts.get(r["concept_uuid"], {}),
            })
        return out
    finally:
        conn.row_factory = prior


def patch_gold_fact(
    conn: sqlite3.Connection,
    benchmark_id: int,
    concept_uuid: str,
    *,
    period: str,
    entity_scope: str,
    value: Optional[float],
) -> dict:
    """Upsert a single gold value (the spot-edit path after import).

    Validates the concept exists, belongs to the benchmark's template set, and
    is a gradeable leaf (LEAF / MATRIX_CELL). Returns the stored fact dict.
    Raises ``ValueError`` on a bad target so the route can map it to a 4xx.
    """
    if period not in ("CY", "PY"):
        raise ValueError(f"period must be CY or PY, got {period!r}")
    if entity_scope not in ("Company", "Group"):
        raise ValueError(
            f"entity_scope must be Company or Group, got {entity_scope!r}"
        )
    template_ids = set(benchmark_template_ids(conn, benchmark_id))
    node = conn.execute(
        "SELECT template_id, kind FROM concept_nodes WHERE concept_uuid = ?",
        (concept_uuid,),
    ).fetchone()
    if node is None:
        raise ValueError("Concept not found")
    if node[0] not in template_ids:
        raise ValueError("Concept is not part of this benchmark's templates")
    if node[1] not in ("LEAF", "MATRIX_CELL"):
        raise ValueError("Only LEAF / MATRIX_CELL gold cells are editable")

    # An empty value clears the gold cell entirely (the human marks it blank).
    if value is None:
        conn.execute(
            "DELETE FROM gold_concept_facts WHERE benchmark_id = ? "
            "AND concept_uuid = ? AND period = ? AND entity_scope = ?",
            (benchmark_id, concept_uuid, period, entity_scope),
        )
        return {"benchmark_id": benchmark_id, "concept_uuid": concept_uuid,
                "period": period, "entity_scope": entity_scope, "value": None}

    conn.execute(
        "INSERT INTO gold_concept_facts(benchmark_id, concept_uuid, period, "
        "entity_scope, value, value_status, source, updated_at) "
        "VALUES (?, ?, ?, ?, ?, 'observed', 'user_edit', ?) "
        "ON CONFLICT(benchmark_id, concept_uuid, period, entity_scope) "
        "DO UPDATE SET value = excluded.value, "
        "value_status = excluded.value_status, "
        "source = excluded.source, updated_at = excluded.updated_at",
        (benchmark_id, concept_uuid, period, entity_scope, float(value), _now()),
    )
    return {"benchmark_id": benchmark_id, "concept_uuid": concept_uuid,
            "period": period, "entity_scope": entity_scope, "value": float(value)}
